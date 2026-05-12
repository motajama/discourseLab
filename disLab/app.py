import csv
from datetime import datetime
import html
import io
import json
from pathlib import Path
import re
import sqlite3
import zipfile
from uuid import uuid4

from docx import Document as DocxDocument
from flask import (
    Flask,
    Response,
    abort,
    flash,
    g,
    has_request_context,
    jsonify,
    redirect,
    render_template,
    request,
    session,
    url_for,
)
from markupsafe import Markup, escape
from werkzeug.utils import secure_filename

try:
    from version import APP_NAME, APP_PHASE, APP_RELEASE_LABEL, APP_VERSION
except ImportError:
    from disLab.version import APP_NAME, APP_PHASE, APP_RELEASE_LABEL, APP_VERSION


CURRENT_PHASE_LABEL = f"{APP_VERSION} - {APP_RELEASE_LABEL}"
DEFAULT_PROJECT_NAME = "Demo Project"
DEFAULT_PROJECT_DESCRIPTION = "Initial local discourseLab project."
DEFAULT_CODE_COLOR = "#f4c542"
DEFAULT_CDA_MARKER_COLOR = "#7c9a45"
METHODOLOGY_MODES = {
    "generic": "Generic",
    "gt": "Grounded Theory",
    "cda": "Critical Discourse Analysis",
    "mixed": "Mixed",
}
METHODOLOGY_MODE_EXPLANATIONS = {
    "generic": "Generic qualitative coding: documents, segments, open codes, memos, and exports.",
    "gt": "Grounded Theory: open coding, axial coding, categories, constant comparison, and theoretical integration.",
    "cda": "Critical Discourse Analysis: textual, discursive-practice, and social-practice analysis with attention to power, ideology, actors, voice, and discourse features.",
    "mixed": "Mixed GT + CDA: combines grounded theory development with critical discourse analysis; requires explicit protocol decisions.",
}
ALLOWED_EXTENSIONS = {"txt", "docx"}
MAX_UPLOAD_SIZE = 16 * 1024 * 1024
MEMO_TYPES = {
    "project": "Project memo",
    "document": "Document memo",
    "segment": "Segment memo",
    "code": "Code memo",
    "methodological": "Methodological memo",
    "theoretical": "Theoretical memo",
    "reflexive": "Reflexive memo",
    "comparison": "Comparison memo",
    "negative_case": "Negative case memo",
}
MEMO_STATUSES = {
    "draft": "Draft",
    "important": "Important",
    "use_in_article": "Use in article",
    "archived": "Archived",
}
GT_COLUMNS = [
    "gt_conditions",
    "gt_context",
    "gt_actions_interactions",
    "gt_consequences",
    "gt_properties",
    "gt_dimensions",
    "gt_theoretical_note",
]
CDA_MARKER_TYPES = {
    "textual": "Textual",
    "discursive_practice": "Discursive practice",
    "social_practice": "Social practice",
    "actor": "Actor",
    "agency": "Agency",
    "voice": "Voice",
    "silence": "Silence",
    "modality": "Modality",
    "evaluation": "Evaluation",
    "metaphor": "Metaphor",
    "presupposition": "Presupposition",
    "nominalization": "Nominalization",
    "passivization": "Passivization",
    "intertextuality": "Intertextuality",
    "legitimation": "Legitimation",
    "framing": "Framing",
    "ideology": "Ideology",
    "power_relation": "Power relation",
    "other": "Other",
}
CDA_MARKER_TYPE_EXPLANATIONS = {
    "textual": "Textual-level features such as vocabulary, grammar, modality, metaphor, nominalization, or passivization.",
    "discursive_practice": "Features concerning production, circulation, genre, intertextuality, voice, quotation, and reception.",
    "social_practice": "Features connecting discourse to institutions, ideology, power, hegemony, and broader social relations.",
    "actor": "A discourse participant or represented social actor.",
    "agency": "How action, responsibility, and capacity to act are assigned or removed.",
    "voice": "Who is allowed to speak directly or indirectly.",
    "silence": "Who is absent, backgrounded, or denied voice.",
    "modality": "Expressions of necessity, possibility, certainty, obligation, or probability.",
    "evaluation": "Positive or negative valuation, judgment, affect, or appraisal.",
    "metaphor": "A meaning transfer where one domain is understood through another, such as migration as a wave.",
    "presupposition": "Something treated as already given or taken for granted by the text.",
    "nominalization": "Turning processes or actions into nouns, often hiding agency.",
    "passivization": "Using passive constructions that may background or omit the actor.",
    "intertextuality": "Links to other texts, voices, genres, discourses, or quoted authorities.",
    "legitimation": "Discursive justification of actions, policies, hierarchies, or institutions.",
    "framing": "Selection and organization of meaning that defines what the issue is about.",
    "ideology": "A structured system of meanings that supports or contests social power.",
    "power_relation": "A relation of domination, authority, access, dependence, or control.",
    "other": "A project-specific discourse feature not covered by existing types.",
}
ACTOR_TYPES = {
    "individual": "Individual",
    "group": "Group",
    "institution": "Institution",
    "state_actor": "State actor",
    "expert": "Expert",
    "journalist": "Journalist",
    "politician": "Politician",
    "public": "Public",
    "vulnerable_group": "Vulnerable group",
    "abstract_actor": "Abstract actor",
    "other": "Other",
}
ACTOR_TYPE_EXPLANATIONS = {
    "individual": "A named or individualized person.",
    "group": "A collective actor such as migrants, citizens, workers, or audiences.",
    "institution": "An organization or institutional body.",
    "state_actor": "A state body, authority, police, military, court, ministry, or government agency.",
    "expert": "A speaker positioned as specialist, analyst, scientist, professional, or authority.",
    "journalist": "A journalist, presenter, reporter, editor, or media worker.",
    "politician": "An elected official, party representative, candidate, or political spokesperson.",
    "public": "The general public, citizens, audiences, or public opinion.",
    "vulnerable_group": "A socially vulnerable or marginalized group represented in discourse.",
    "abstract_actor": "An abstract entity represented as acting, such as the market, Europe, the state, or history.",
    "other": "A project-specific actor type.",
}
ACTOR_RELATION_TYPES = {
    "speaks": "Speaks",
    "is_quoted": "Is quoted",
    "is_spoken_about": "Is spoken about",
    "is_evaluated": "Is evaluated",
    "acts": "Acts",
    "is_acted_upon": "Is acted upon",
    "is_silenced": "Is silenced",
    "is_backgrounded": "Is backgrounded",
    "is_aggregated": "Is aggregated",
    "is_individualized": "Is individualized",
}
ACTOR_RELATION_TYPE_EXPLANATIONS = {
    "speaks": "The actor is given direct or indirect speech.",
    "is_quoted": "The actor's words are quoted or reported as a source.",
    "is_spoken_about": "The actor is represented as an object of discussion.",
    "is_evaluated": "The actor is judged, appraised, praised, blamed, or otherwise valued.",
    "acts": "The actor is represented as doing something or causing an action.",
    "is_acted_upon": "The actor is represented as receiving, suffering, or being affected by action.",
    "is_silenced": "The actor is absent as a speaker or denied voice.",
    "is_backgrounded": "The actor is present but made less visible or less central.",
    "is_aggregated": "The actor is represented as part of a collective or mass.",
    "is_individualized": "The actor is represented as a specific person or individualized case.",
}
DISCOURSE_FEATURE_TYPES = {
    "metaphor": "Metaphor",
    "modality": "Modality",
    "evaluation": "Evaluation",
    "presupposition": "Presupposition",
    "legitimation": "Legitimation",
    "intertextuality": "Intertextuality",
    "framing": "Framing",
    "nominalization": "Nominalization",
    "passivization": "Passivization",
    "agency": "Agency",
    "ideology": "Ideology",
    "power_relation": "Power relation",
    "other": "Other",
}
DISCOURSE_FEATURE_TYPE_EXPLANATIONS = {
    key: CDA_MARKER_TYPE_EXPLANATIONS[key]
    for key in DISCOURSE_FEATURE_TYPES
    if key in CDA_MARKER_TYPE_EXPLANATIONS
}
RELATION_ENTITY_TYPES = {
    "document": "Document",
    "segment": "Segment",
    "code": "Code",
    "memo": "Memo",
    "research_question": "Research question",
    "discourse_marker": "CDA marker",
    "actor": "Actor",
    "discourse_feature": "Discourse feature",
}
RELATION_TYPES = {
    "supports": "Supports",
    "contradicts": "Contradicts",
    "elaborates": "Elaborates",
    "explains": "Explains",
    "is_evidence_for": "Is evidence for",
    "is_example_of": "Is example of",
    "is_negative_case_for": "Is negative case for",
    "is_part_of": "Is part of",
    "contrasts_with": "Contrasts with",
    "leads_to": "Leads to",
    "conditions": "Conditions",
    "causes": "Causes",
    "enables": "Enables",
    "limits": "Limits",
    "transforms_into": "Transforms into",
    "open_code_supports_axial_code": "Open code supports axial code",
    "axial_code_supports_category": "Axial code supports category",
    "category_integrates": "Category integrates",
    "property_of": "Property of",
    "dimension_of": "Dimension of",
    "consequence_of": "Consequence of",
    "condition_for": "Condition for",
    "frames": "Frames",
    "legitimizes": "Legitimizes",
    "delegitimizes": "Delegitimizes",
    "naturalizes": "Naturalizes",
    "silences": "Silences",
    "foregrounds": "Foregrounds",
    "backgrounds": "Backgrounds",
    "individualizes": "Individualizes",
    "aggregates": "Aggregates",
    "constructs_actor_as": "Constructs actor as",
    "reproduces_power_relation": "Reproduces power relation",
    "challenges_power_relation": "Challenges power relation",
    "presupposes": "Presupposes",
    "metaphorizes": "Metaphorizes",
}
RELATION_TYPE_EXPLANATIONS = {
    "supports": "The source strengthens, corroborates, or gives backing to the target.",
    "contradicts": "The source conflicts with, challenges, or weakens the target.",
    "elaborates": "The source adds detail, nuance, or specification to the target.",
    "explains": "The source helps account for why the target occurs or matters.",
    "is_evidence_for": "The source is empirical evidence for the target claim, code, or relation.",
    "is_example_of": "The source is a concrete instance of the target concept.",
    "is_negative_case_for": "The source is a case that complicates or challenges the target pattern.",
    "is_part_of": "The source belongs within the target as a component or subpart.",
    "contrasts_with": "The source is analytically different from the target in a meaningful way.",
    "leads_to": "The source precedes or contributes to the target as a processual outcome.",
    "conditions": "The source shapes the circumstances under which the target becomes possible.",
    "causes": "The source is interpreted as a cause of the target.",
    "enables": "The source makes the target possible or easier to occur.",
    "limits": "The source constrains, blocks, or narrows the target.",
    "transforms_into": "The source changes into or is reworked as the target.",
    "open_code_supports_axial_code": "An open code provides grounded evidence for an axial code.",
    "axial_code_supports_category": "An axial code contributes to a broader category.",
    "category_integrates": "A category organizes or integrates other analytical elements.",
    "property_of": "The source names a property or attribute of the target.",
    "dimension_of": "The source names a dimension or range of variation of the target.",
    "consequence_of": "The source is an outcome or consequence of the target.",
    "condition_for": "The source is a condition that supports or shapes the target.",
    "frames": "The source defines how the target should be understood.",
    "legitimizes": "The source justifies or normalizes the target.",
    "delegitimizes": "The source undermines or contests the target's legitimacy.",
    "naturalizes": "The source makes the target appear normal, inevitable, or common sense.",
    "silences": "The source suppresses, omits, or denies voice to the target.",
    "foregrounds": "The source makes the target more visible or central.",
    "backgrounds": "The source makes the target less visible or less central.",
    "individualizes": "The source represents the target as a specific person or individualized case.",
    "aggregates": "The source represents the target as a collective, mass, or category.",
    "constructs_actor_as": "The source represents an actor through a specific identity, role, or attribute.",
    "reproduces_power_relation": "The source sustains a relation of authority, domination, dependence, or control.",
    "challenges_power_relation": "The source contests a relation of authority, domination, dependence, or control.",
    "presupposes": "The source treats the target as already given or taken for granted.",
    "metaphorizes": "The source understands the target through a metaphorical domain.",
}
RELATION_STRENGTHS = {
    "weak": "Weak",
    "moderate": "Moderate",
    "strong": "Strong",
    "uncertain": "Uncertain",
}
RELATION_STRENGTH_EXPLANATIONS = {
    "weak": "An exploratory or minor relation with limited evidence.",
    "moderate": "A plausible relation supported by some evidence.",
    "strong": "A central or well-supported analytical relation.",
    "uncertain": "A tentative relation that requires further checking.",
}
GT_TERM_EXPLANATIONS = {
    "open_code": "A close-to-data code that names what is happening in a segment.",
    "axial_code": "A higher-level code that groups open codes around conditions, actions, consequences, properties, or dimensions.",
    "category": "An integrating analytical concept with explanatory weight across several codes.",
    "conditions": "Circumstances that shape or enable the phenomenon.",
    "context": "The setting, situation, or background that gives the phenomenon meaning.",
    "actions_interactions": "Actions, strategies, responses, or interactions connected to the phenomenon.",
    "consequences": "Outcomes or effects that follow from actions or conditions.",
    "properties": "Attributes or characteristics that define a code or category.",
    "dimensions": "Variations, ranges, or degrees along which a property changes.",
}
VISUAL_MODEL_MODES = {
    "simplified": "Simplified",
    "argument": "Argument",
    "evidence": "Evidence",
    "gt": "GT",
    "cda": "CDA",
    "full": "Full",
}
VISUAL_MODEL_MODE_ORDER = ["simplified", "argument", "evidence", "gt", "cda", "full"]
METHODOLOGY_NOTE_TYPES = {
    "protocol": "Protocol",
    "coding_rule": "Coding rule",
    "sampling_rule": "Sampling rule",
    "interpretation_rule": "Interpretation rule",
    "reflexive_note": "Reflexive note",
    "source_note": "Source note",
    "decision_log": "Decision log",
    "warning": "Warning",
    "other": "Other",
}
METHODOLOGY_AREAS = {
    "generic": "Generic",
    "gt": "Grounded Theory",
    "cda": "CDA",
    "mixed": "Mixed",
    "export": "Export",
    "ethics": "Ethics",
    "other": "Other",
}
METHODOLOGY_NOTE_STATUSES = {
    "draft": "Draft",
    "active": "Active",
    "needs_review": "Needs review",
    "archived": "Archived",
}
METHODOLOGY_LINKED_ENTITY_TYPES = {
    "project": "Project",
    "document": "Document",
    "segment": "Segment",
    "code": "Code",
    "memo": "Memo",
    "research_question": "Research question",
    "discourse_marker": "CDA marker",
    "actor": "Actor",
    "discourse_feature": "Discourse feature",
    "relation": "Relation",
}

BASE_DIR = Path(__file__).resolve().parent
INSTANCE_DIR = BASE_DIR / "instance"
UPLOAD_DIR = BASE_DIR / "uploads"
EXPORT_DIR = BASE_DIR / "exports"
METHODOLOGY_DIR = BASE_DIR / "methodology"
DATABASE = INSTANCE_DIR / "disLab.sqlite"
SCHEMA = BASE_DIR / "schema.sql"


def create_app() -> Flask:
    app = Flask(__name__, instance_path=str(INSTANCE_DIR))
    app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_SIZE
    app.config["SECRET_KEY"] = "discourseLab-local-development-key"

    INSTANCE_DIR.mkdir(exist_ok=True)
    UPLOAD_DIR.mkdir(exist_ok=True)
    EXPORT_DIR.mkdir(exist_ok=True)

    with app.app_context():
        init_db_if_needed()
        get_active_project()

    app.teardown_appcontext(close_db)

    @app.context_processor
    def inject_project_context():
        active_project = get_active_project()
        return {
            "active_project": active_project,
            "methodology_mode_labels": METHODOLOGY_MODES,
            "methodology_mode_explanations": METHODOLOGY_MODE_EXPLANATIONS,
            "marker_type_explanations": CDA_MARKER_TYPE_EXPLANATIONS,
            "actor_type_explanations": ACTOR_TYPE_EXPLANATIONS,
            "actor_relation_type_explanations": ACTOR_RELATION_TYPE_EXPLANATIONS,
            "feature_type_explanations": DISCOURSE_FEATURE_TYPE_EXPLANATIONS,
            "relation_type_explanations": RELATION_TYPE_EXPLANATIONS,
            "relation_strength_explanations": RELATION_STRENGTH_EXPLANATIONS,
            "gt_term_explanations": GT_TERM_EXPLANATIONS,
            "supports_gt": project_supports_gt(active_project),
            "supports_cda": project_supports_cda(active_project),
            "visual_model_modes": VISUAL_MODEL_MODES,
            "visual_model_mode_order": VISUAL_MODEL_MODE_ORDER,
            "methodology_note_type_labels": METHODOLOGY_NOTE_TYPES,
            "methodology_area_labels": METHODOLOGY_AREAS,
            "methodology_note_status_labels": METHODOLOGY_NOTE_STATUSES,
        }

    @app.route("/")
    def dashboard():
        active_project = get_active_project()
        counts = get_dashboard_counts(active_project["id"])
        audit_entries = get_latest_audit_entries(active_project["id"])
        methodology_counts = get_methodology_note_counts(active_project["id"])
        progress_barometers = get_dashboard_progress_barometers(active_project, counts)
        first_document_id = get_first_document_id(active_project["id"])
        return render_template(
            "dashboard.html",
            title="Dashboard",
            active_page="dashboard",
            active_project=active_project,
            counts=counts,
            audit_entries=audit_entries,
            methodology_counts=methodology_counts,
            progress_barometers=progress_barometers,
            suggested_actions=get_dashboard_suggestions(
                active_project, counts, methodology_counts, progress_barometers
            ),
            first_document_id=first_document_id,
            current_phase=CURRENT_PHASE_LABEL,
            memo_type_labels=MEMO_TYPES,
            memo_status_labels=MEMO_STATUSES,
        )

    @app.route("/projects")
    def projects():
        active_project = get_active_project()
        return render_template(
            "projects.html",
            title="Projects",
            active_page="projects",
            active_project=active_project,
            projects=get_projects(),
            methodology_mode_labels=METHODOLOGY_MODES,
        )

    @app.route("/projects/new")
    def new_project():
        return render_template(
            "project_form.html",
            title="New Project",
            active_page="projects",
            active_project=get_active_project(),
            project=None,
            methodology_mode_labels=METHODOLOGY_MODES,
            form_action=url_for("create_project"),
        )

    @app.route("/projects/new", methods=["POST"])
    def create_project():
        name = request.form.get("name", "").strip()
        description = request.form.get("description", "").strip()
        methodology_mode = request.form.get("methodology_mode", "mixed").strip()
        research_goal = request.form.get("research_goal", "").strip()
        principal_investigator = request.form.get("principal_investigator", "").strip()
        if not name:
            flash("Project name is required.", "error")
            return redirect(url_for("new_project"))
        if methodology_mode not in METHODOLOGY_MODES:
            flash("Invalid methodology mode.", "error")
            return redirect(url_for("new_project"))
        project_id = execute_write(
            """
            INSERT INTO projects (
                name, description, methodology_mode, status, last_opened_at,
                research_goal, principal_investigator
            )
            VALUES (?, ?, ?, 'active', CURRENT_TIMESTAMP, ?, ?)
            """,
            (name, description, methodology_mode, research_goal, principal_investigator),
        )
        set_active_project(project_id)
        log_action(project_id, "project", project_id, "create_project", f"Created project: {name}")
        flash(f"Created project: {name}", "success")
        return redirect(url_for("dashboard"))

    @app.route("/projects/<int:project_id>/open", methods=["POST"])
    def open_project(project_id: int):
        project = get_project(project_id)
        if project is None:
            flash("Project not found.", "error")
            return redirect(url_for("projects"))
        set_active_project(project_id)
        log_action(project_id, "project", project_id, "open_project", f"Opened project: {project['name']}")
        flash(f"Opened project: {project['name']}", "success")
        return redirect(url_for("dashboard"))

    @app.route("/projects/<int:project_id>/edit")
    def edit_project(project_id: int):
        project = get_project(project_id)
        if project is None:
            flash("Project not found.", "error")
            abort(404)
        return render_template(
            "project_form.html",
            title=f"Edit Project: {project['name']}",
            active_page="projects",
            active_project=get_active_project(),
            project=project,
            methodology_mode_labels=METHODOLOGY_MODES,
            form_action=url_for("update_project", project_id=project_id),
        )

    @app.route("/projects/<int:project_id>/edit", methods=["POST"])
    def update_project(project_id: int):
        project = get_project(project_id)
        if project is None:
            flash("Project not found.", "error")
            abort(404)
        name = request.form.get("name", "").strip()
        description = request.form.get("description", "").strip()
        methodology_mode = request.form.get("methodology_mode", "mixed").strip()
        research_goal = request.form.get("research_goal", "").strip()
        principal_investigator = request.form.get("principal_investigator", "").strip()
        if not name:
            flash("Project name is required.", "error")
            return redirect(url_for("edit_project", project_id=project_id))
        if methodology_mode not in METHODOLOGY_MODES:
            flash("Invalid methodology mode.", "error")
            return redirect(url_for("edit_project", project_id=project_id))
        execute_write(
            """
            UPDATE projects
            SET name = ?, description = ?, methodology_mode = ?,
                research_goal = ?, principal_investigator = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND status != 'deleted'
            """,
            (
                name,
                description,
                methodology_mode,
                research_goal,
                principal_investigator,
                project_id,
            ),
        )
        log_action(project_id, "project", project_id, "update_project", f"Updated project: {name}")
        flash(f"Updated project: {name}", "success")
        return redirect(url_for("projects"))

    @app.route("/projects/<int:project_id>/delete", methods=["POST"])
    def soft_delete_project(project_id: int):
        project = get_project(project_id)
        if project is None:
            flash("Project not found.", "error")
            return redirect(url_for("projects"))
        execute_write(
            """
            UPDATE projects
            SET status = 'deleted', updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (project_id,),
        )
        log_action(
            project_id,
            "project",
            project_id,
            "soft_delete_project",
            f"Soft-deleted project: {project['name']}",
        )
        if session.get("active_project_id") == project_id:
            session.pop("active_project_id", None)
            get_active_project()
        flash(f"Deleted project: {project['name']}", "success")
        return redirect(url_for("projects"))

    @app.route("/documents")
    def documents():
        active_project = get_active_project()
        document_rows = get_documents_for_project(active_project["id"])
        return render_template(
            "documents.html",
            title="Documents",
            active_page="documents",
            active_project=active_project,
            documents=document_rows,
        )

    @app.route("/documents/import", methods=["POST"])
    def import_document():
        active_project = get_active_project()
        uploaded_file = request.files.get("document_file")
        title = request.form.get("title", "").strip()
        note = request.form.get("note", "").strip()

        if uploaded_file is None or uploaded_file.filename == "":
            flash("Choose a TXT or DOCX file to import.", "error")
            return redirect(url_for("documents"))

        original_filename = uploaded_file.filename
        extension = get_file_extension(original_filename)
        if extension not in ALLOWED_EXTENSIONS:
            flash("Unsupported file type. Import only TXT or DOCX files.", "error")
            return redirect(url_for("documents"))

        safe_name = build_unique_upload_name(original_filename)
        saved_path = UPLOAD_DIR / safe_name
        uploaded_file.save(saved_path)

        try:
            text_content = extract_text(saved_path, extension)
        except Exception as error:
            saved_path.unlink(missing_ok=True)
            flash(f"Could not extract text from {original_filename}: {error}", "error")
            return redirect(url_for("documents"))

        text_content = normalize_line_endings(text_content).strip()
        if not text_content:
            saved_path.unlink(missing_ok=True)
            flash("The uploaded file did not contain extractable text.", "error")
            return redirect(url_for("documents"))

        if not title:
            title = Path(original_filename).stem

        document_id = execute_write(
            """
            INSERT INTO documents (
                project_id, title, original_filename, file_type, text_content, note
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (active_project["id"], title, original_filename, extension, text_content, note),
        )
        log_action(
            project_id=active_project["id"],
            entity_type="document",
            entity_id=document_id,
            action="import_document",
            details=f"Imported document: {title}",
        )
        flash(f"Imported document: {title}", "success")
        return redirect(url_for("documents"))

    @app.route("/documents/<int:document_id>")
    def document_view(document_id: int):
        active_project = get_active_project()
        document = get_document_for_project(document_id, active_project["id"])
        if document is None:
            flash("Document not found.", "error")
            abort(404)

        segments = get_segments_for_document(document_id, active_project["id"])
        open_codes = get_open_codes_for_project(active_project["id"])
        discourse_markers = get_discourse_markers_for_project(active_project["id"])
        actors = get_actors_for_project(active_project["id"])
        document_memos = get_memos_for_entity(active_project["id"], "document", document_id)
        highlighted_text = build_highlighted_document_text(
            document["text_content"] or "", segments, project_supports_cda(active_project)
        )
        return render_template(
            "document_view.html",
            title=document["title"],
            active_page="documents",
            active_project=active_project,
            document=document,
            text_length=len(document["text_content"] or ""),
            segment_count=len(segments),
            coded_segment_count=sum(1 for segment in segments if segment["codes"]),
            document_memo_count=len(document_memos),
            cda_annotation_count=sum(
                len(segment["discourse_markers"])
                + len(segment["actors"])
                + len(segment["features"])
                for segment in segments
            ),
            segments=segments,
            open_codes=open_codes,
            discourse_markers=discourse_markers,
            actors=actors,
            marker_type_labels=CDA_MARKER_TYPES,
            actor_type_labels=ACTOR_TYPES,
            actor_relation_type_labels=ACTOR_RELATION_TYPES,
            feature_type_labels=DISCOURSE_FEATURE_TYPES,
            supports_gt=project_supports_gt(active_project),
            supports_cda=project_supports_cda(active_project),
            document_memos=document_memos,
            memo_type_labels=MEMO_TYPES,
            memo_status_labels=MEMO_STATUSES,
            highlighted_text=highlighted_text,
            methodology_helper=get_methodology_helper(active_project, "document"),
        )

    @app.route("/documents/<int:document_id>/segments", methods=["POST"])
    def create_segment(document_id: int):
        active_project = get_active_project()
        document = get_document_for_project(document_id, active_project["id"])
        if document is None:
            flash("Document not found.", "error")
            abort(404)

        segment_name = request.form.get("name", "").strip()
        selected_text = request.form.get("selected_text", "")
        note = request.form.get("note", "").strip()
        try:
            start_offset = int(request.form.get("start_offset", ""))
            end_offset = int(request.form.get("end_offset", ""))
        except ValueError:
            flash("Invalid selection. Select text inside the document panel.", "error")
            return redirect_to_document(document_id)

        document_text = document["text_content"] or ""
        if not is_valid_segment_selection(
            document_text, selected_text, start_offset, end_offset
        ):
            flash("Invalid selection. Select a passage inside the document text.", "error")
            return redirect_to_document(document_id)

        selected_text = document_text[start_offset:end_offset]
        segment_id = execute_write(
            """
            INSERT INTO segments (
                document_id, name, selected_text, start_offset, end_offset, note
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (document_id, segment_name, selected_text, start_offset, end_offset, note),
        )
        log_action(
            project_id=active_project["id"],
            entity_type="segment",
            entity_id=segment_id,
            action="create_segment",
            details=f"Created segment in document: {document['title']}",
        )
        flash("Segment created.", "success")
        return redirect_to_document(document_id)

    @app.route("/documents/<int:document_id>/delete", methods=["POST"])
    def delete_document(document_id: int):
        active_project = get_active_project()
        document = get_document_for_project(document_id, active_project["id"])
        if document is None:
            flash("Document not found.", "error")
            abort(404)

        delete_document_data(document_id)
        log_action(
            project_id=active_project["id"],
            entity_type="document",
            entity_id=document_id,
            action="delete_document",
            details=f"Deleted document: {document['title']}",
        )
        flash(f"Deleted document: {document['title']}", "success")
        return redirect(url_for("documents"))

    @app.route("/segments/<int:segment_id>/delete", methods=["POST"])
    def delete_segment(segment_id: int):
        active_project = get_active_project()
        segment = get_segment_for_project(segment_id, active_project["id"])
        if segment is None:
            flash("Segment not found.", "error")
            abort(404)

        db = get_db()
        db.execute("DELETE FROM segment_codes WHERE segment_id = ?", (segment_id,))
        db.execute("DELETE FROM segment_discourse_markers WHERE segment_id = ?", (segment_id,))
        db.execute("DELETE FROM segment_actors WHERE segment_id = ?", (segment_id,))
        db.execute("DELETE FROM discourse_features WHERE segment_id = ?", (segment_id,))
        db.execute("DELETE FROM segments WHERE id = ?", (segment_id,))
        db.commit()
        log_action(
            project_id=active_project["id"],
            entity_type="segment",
            entity_id=segment_id,
            action="delete_segment",
            details=f"Deleted segment from document: {segment['document_title']}",
        )
        flash("Segment deleted.", "success")
        return redirect_to_document(segment["document_id"])

    @app.route("/codes")
    def codes():
        active_project = get_active_project()
        code_type_filter = request.args.get("code_type", "all").strip()
        if not project_supports_gt(active_project):
            code_type_filter = "open"
        codes_rows = get_codes_for_project(active_project["id"], code_type_filter)
        return render_template(
            "codes.html",
            title="Codes",
            active_page="codes",
            active_project=active_project,
            codes=codes_rows,
            code_type_filter=code_type_filter,
            default_code_color=DEFAULT_CODE_COLOR,
            supports_gt=project_supports_gt(active_project),
        )

    @app.route("/codes/create", methods=["POST"])
    def create_code():
        active_project = get_active_project()
        name = request.form.get("name", "").strip()
        description = request.form.get("description", "").strip()
        color = normalize_code_color(request.form.get("color", ""))
        document_id = request.form.get("document_id", "").strip()

        if not name:
            flash("Code name is required.", "error")
            return redirect_after_code_change(document_id)

        code_id = execute_write(
            """
            INSERT INTO codes (project_id, name, description, code_type, color)
            VALUES (?, ?, ?, 'open', ?)
            """,
            (active_project["id"], name, description, color),
        )
        log_action(
            project_id=active_project["id"],
            entity_type="code",
            entity_id=code_id,
            action="create_open_code",
            details=f"Created open code: {name}",
        )
        flash(f"Created open code: {name}", "success")
        return redirect_after_code_change(document_id)

    @app.route("/codes/<int:code_id>")
    def code_detail(code_id: int):
        active_project = get_active_project()
        code = get_code_for_project(code_id, active_project["id"])
        if code is None:
            flash("Code not found.", "error")
            abort(404)

        return render_template(
            "code_detail.html",
            title=f"Code: {code['name']}",
            active_page="codes",
            active_project=active_project,
            code=code,
            usage_count=get_code_usage_count(code_id),
            code_memos=get_memos_for_entity(active_project["id"], "code", code_id),
            coded_segments=get_segments_for_code(code_id, active_project["id"]),
            parent_code=get_parent_code(code["parent_id"], active_project["id"]) if code["parent_id"] else None,
            child_open_codes=get_child_codes(code_id, active_project["id"], "open"),
            child_axial_codes=get_child_codes(code_id, active_project["id"], "axial"),
            category_open_codes=get_open_codes_under_category(code_id, active_project["id"]),
            hierarchy_segments=get_hierarchy_segments_for_code(code, active_project["id"]),
            model_relations=get_relations_for_entity(active_project["id"], "code", code_id, limit=5),
            model_relation_count=get_relation_count_for_entity(active_project["id"], "code", code_id),
            memo_type_labels=MEMO_TYPES,
            memo_status_labels=MEMO_STATUSES,
        )

    @app.route("/codes/<int:code_id>/edit")
    def edit_code(code_id: int):
        active_project = get_active_project()
        code = get_code_for_project(code_id, active_project["id"])
        if code is None:
            flash("Invalid code.", "error")
            abort(404)

        return render_template(
            "code_edit.html",
            title=f"Edit Code: {code['name']}",
            active_page="codes",
            active_project=active_project,
            code=code,
            axial_codes=get_gt_axial_codes(active_project["id"]),
            categories=get_gt_categories(active_project["id"]),
            supports_gt=project_supports_gt(active_project),
        )

    @app.route("/codes/<int:code_id>/edit", methods=["POST"])
    def update_code(code_id: int):
        active_project = get_active_project()
        code = get_code_for_project(code_id, active_project["id"])
        if code is None:
            flash("Invalid code.", "error")
            abort(404)

        name = request.form.get("name", "").strip()
        description = request.form.get("description", "").strip()
        color = normalize_code_color(request.form.get("color", ""))
        parent_id_raw = request.form.get("parent_id", "").strip()
        definition = request.form.get("definition", "").strip()
        include_when = request.form.get("include_when", "").strip()
        exclude_when = request.form.get("exclude_when", "").strip()
        example = request.form.get("example", "").strip()
        analytical_note = request.form.get("analytical_note", "").strip()
        if not name:
            flash("Code name is required.", "error")
            return redirect(url_for("edit_code", code_id=code_id))
        parent_id = code["parent_id"]
        if project_supports_gt(active_project):
            parent_id = None
        if project_supports_gt(active_project) and parent_id_raw:
            if not parent_id_raw.isdigit():
                flash("Invalid GT hierarchy operation.", "error")
                return redirect(url_for("edit_code", code_id=code_id))
            parent_id = int(parent_id_raw)
            valid, message = validate_parent_assignment(
                code_id, code["code_type"], parent_id, active_project["id"]
            )
            if not valid:
                flash(message, "error")
                return redirect(url_for("edit_code", code_id=code_id))

        execute_write(
            """
            UPDATE codes
            SET name = ?, description = ?, color = ?, definition = ?,
                include_when = ?, exclude_when = ?, example = ?,
                analytical_note = ?, parent_id = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND project_id = ?
            """,
            (
                name,
                description,
                color,
                definition,
                include_when,
                exclude_when,
                example,
                analytical_note,
                parent_id,
                code_id,
                active_project["id"],
            ),
        )
        log_action(
            project_id=active_project["id"],
            entity_type="code",
            entity_id=code_id,
            action="update_open_code",
            details=f"Updated open code: {name}",
        )
        flash(f"Updated codebook entry: {name}", "success")
        return redirect(url_for("codes"))

    @app.route("/codes/<int:code_id>/delete", methods=["POST"])
    def delete_code(code_id: int):
        active_project = get_active_project()
        code = get_code_for_project(code_id, active_project["id"])
        if code is None:
            flash("Invalid code.", "error")
            abort(404)

        db = get_db()
        if code["code_type"] == "open":
            db.execute("DELETE FROM segment_codes WHERE code_id = ?", (code_id,))
        if code["code_type"] == "axial":
            db.execute(
                "UPDATE codes SET parent_id = NULL WHERE parent_id = ? AND project_id = ? AND code_type = 'open'",
                (code_id, active_project["id"]),
            )
        if code["code_type"] == "category":
            db.execute(
                "UPDATE codes SET parent_id = NULL WHERE parent_id = ? AND project_id = ? AND code_type = 'axial'",
                (code_id, active_project["id"]),
            )
        db.execute(
            "DELETE FROM codes WHERE id = ? AND project_id = ?",
            (code_id, active_project["id"]),
        )
        db.commit()
        log_action(
            project_id=active_project["id"],
            entity_type="code",
            entity_id=code_id,
            action="delete_open_code",
            details=f"Deleted {code['code_type']} code: {code['name']}",
        )
        flash(f"Deleted code: {code['name']}", "success")
        return redirect(url_for("codes"))

    @app.route("/segments/<int:segment_id>/codes", methods=["POST"])
    def assign_code_to_segment(segment_id: int):
        active_project = get_active_project()
        segment = get_segment_for_project(segment_id, active_project["id"])
        if segment is None:
            flash("Invalid segment.", "error")
            abort(404)

        try:
            code_id = int(request.form.get("code_id", ""))
        except ValueError:
            flash("Invalid code.", "error")
            return redirect_to_document(segment["document_id"])

        code = get_code_for_project(code_id, active_project["id"])
        if code is None:
            flash("Invalid code.", "error")
            return redirect_to_document(segment["document_id"])

        if segment_has_code(segment_id, code_id):
            flash("Code already assigned to this segment.", "error")
            return redirect_to_document(segment["document_id"])

        execute_write(
            "INSERT INTO segment_codes (segment_id, code_id) VALUES (?, ?)",
            (segment_id, code_id),
        )
        segment_label = segment["name"] or f"segment {segment_id}"
        log_action(
            project_id=active_project["id"],
            entity_type="segment",
            entity_id=segment_id,
            action="assign_open_code",
            details=f"Assigned code {code['name']} to {segment_label}",
        )
        flash(f"Assigned code: {code['name']}", "success")
        return redirect_to_document(segment["document_id"])

    @app.route("/segments/<int:segment_id>/codes/<int:code_id>/remove", methods=["POST"])
    def remove_code_from_segment(segment_id: int, code_id: int):
        active_project = get_active_project()
        segment = get_segment_for_project(segment_id, active_project["id"])
        if segment is None:
            flash("Invalid segment.", "error")
            abort(404)

        code = get_code_for_project(code_id, active_project["id"])
        if code is None:
            flash("Invalid code.", "error")
            return redirect_to_document(segment["document_id"])

        execute_write(
            "DELETE FROM segment_codes WHERE segment_id = ? AND code_id = ?",
            (segment_id, code_id),
        )
        segment_label = segment["name"] or f"segment {segment_id}"
        log_action(
            project_id=active_project["id"],
            entity_type="segment",
            entity_id=segment_id,
            action="remove_open_code",
            details=f"Removed code {code['name']} from {segment_label}",
        )
        flash(f"Removed code: {code['name']}", "success")
        return redirect_to_document(segment["document_id"])

    @app.route("/memos")
    def memos():
        active_project = get_active_project()
        filters = {
            "memo_type": request.args.get("memo_type", "").strip(),
            "status": request.args.get("status", "").strip(),
            "linked_entity_type": request.args.get("linked_entity_type", "").strip(),
            "linked_entity_id": request.args.get("linked_entity_id", "").strip(),
        }
        memo_rows = get_memos_for_project(active_project["id"], filters)
        return render_template(
            "memos.html",
            title="Memos",
            active_page="memos",
            active_project=active_project,
            memos=memo_rows,
            filters=filters,
            memo_type_labels=MEMO_TYPES,
            memo_status_labels=MEMO_STATUSES,
            documents=get_documents_for_memo_links(active_project["id"]),
            segments=get_segments_for_memo_links(active_project["id"]),
            codes=get_open_codes_for_project(active_project["id"]),
        )

    @app.route("/memos/create", methods=["POST"])
    def create_memo():
        active_project = get_active_project()
        data, error = validate_memo_form(active_project["id"])
        next_url = safe_next_url(request.form.get("next_url", ""))
        if error:
            flash(error, "error")
            return redirect(next_url or url_for("memos"))

        memo_id = execute_write(
            """
            INSERT INTO memos (
                project_id, title, body, memo_type, linked_entity_type,
                linked_entity_id, status
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                active_project["id"],
                data["title"],
                data["body"],
                data["memo_type"],
                data["linked_entity_type"],
                data["linked_entity_id"],
                data["status"],
            ),
        )
        log_action(
            project_id=active_project["id"],
            entity_type="memo",
            entity_id=memo_id,
            action="create_memo",
            details=f"Created memo: {data['title']}",
        )
        flash(f"Created memo: {data['title']}", "success")
        return redirect(next_url or url_for("memos"))

    @app.route("/memos/<int:memo_id>/edit")
    def edit_memo(memo_id: int):
        active_project = get_active_project()
        memo = get_memo_for_project(memo_id, active_project["id"])
        if memo is None:
            flash("Memo not found.", "error")
            abort(404)

        return render_template(
            "memo_edit.html",
            title=f"Edit Memo: {memo['title']}",
            active_page="memos",
            active_project=active_project,
            memo=memo,
            memo_type_labels=MEMO_TYPES,
            memo_status_labels=MEMO_STATUSES,
            documents=get_documents_for_memo_links(active_project["id"]),
            segments=get_segments_for_memo_links(active_project["id"]),
            codes=get_open_codes_for_project(active_project["id"]),
        )

    @app.route("/memos/<int:memo_id>/edit", methods=["POST"])
    def update_memo(memo_id: int):
        active_project = get_active_project()
        memo = get_memo_for_project(memo_id, active_project["id"])
        if memo is None:
            flash("Memo not found.", "error")
            abort(404)

        data, error = validate_memo_form(active_project["id"])
        if error:
            flash(error, "error")
            return redirect(url_for("edit_memo", memo_id=memo_id))

        execute_write(
            """
            UPDATE memos
            SET title = ?, body = ?, memo_type = ?, linked_entity_type = ?,
                linked_entity_id = ?, status = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND project_id = ?
            """,
            (
                data["title"],
                data["body"],
                data["memo_type"],
                data["linked_entity_type"],
                data["linked_entity_id"],
                data["status"],
                memo_id,
                active_project["id"],
            ),
        )
        log_action(
            project_id=active_project["id"],
            entity_type="memo",
            entity_id=memo_id,
            action="update_memo",
            details=f"Updated memo: {data['title']}",
        )
        flash(f"Updated memo: {data['title']}", "success")
        return redirect(url_for("memos"))

    @app.route("/memos/<int:memo_id>/delete", methods=["POST"])
    def delete_memo(memo_id: int):
        active_project = get_active_project()
        memo = get_memo_for_project(memo_id, active_project["id"])
        if memo is None:
            flash("Memo not found.", "error")
            abort(404)

        execute_write(
            "DELETE FROM memos WHERE id = ? AND project_id = ?",
            (memo_id, active_project["id"]),
        )
        log_action(
            project_id=active_project["id"],
            entity_type="memo",
            entity_id=memo_id,
            action="delete_memo",
            details=f"Deleted memo: {memo['title']}",
        )
        flash(f"Deleted memo: {memo['title']}", "success")
        return redirect(url_for("memos"))

    @app.route("/gt")
    def gt_workspace():
        active_project = get_active_project()
        if not project_supports_gt(active_project):
            return mode_notice("GT Workspace", active_project)
        return render_template(
            "gt_workspace.html",
            title="Grounded Theory Workspace",
            active_page="gt",
            active_project=active_project,
            open_codes=get_gt_open_codes(active_project["id"]),
            axial_codes=get_gt_axial_codes(active_project["id"]),
            categories=get_gt_categories(active_project["id"]),
            default_code_color=DEFAULT_CODE_COLOR,
            methodology_helper=get_methodology_helper(active_project, "gt_workspace"),
        )

    @app.route("/gt/axial/create", methods=["POST"])
    def create_axial_code():
        if not project_supports_gt(get_active_project()):
            flash("GT Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        return create_gt_code("axial", "create_axial_code", "Created axial code")

    @app.route("/gt/category/create", methods=["POST"])
    def create_category_code():
        if not project_supports_gt(get_active_project()):
            flash("GT Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        return create_gt_code("category", "create_category", "Created category")

    @app.route("/gt/open/<int:open_code_id>/assign-axial", methods=["POST"])
    def assign_open_to_axial(open_code_id: int):
        if not project_supports_gt(get_active_project()):
            flash("GT Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        return assign_code_parent(
            child_id=open_code_id,
            child_type="open",
            parent_type="axial",
            form_field="axial_code_id",
            action="assign_open_to_axial",
            success_template="Assigned open code {child} to axial code {parent}",
        )

    @app.route("/gt/open/<int:open_code_id>/unassign", methods=["POST"])
    def unassign_open_from_axial(open_code_id: int):
        if not project_supports_gt(get_active_project()):
            flash("GT Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        return unassign_code_parent(open_code_id, "open", "unassign_open_from_axial")

    @app.route("/gt/axial/<int:axial_code_id>/assign-category", methods=["POST"])
    def assign_axial_to_category(axial_code_id: int):
        if not project_supports_gt(get_active_project()):
            flash("GT Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        return assign_code_parent(
            child_id=axial_code_id,
            child_type="axial",
            parent_type="category",
            form_field="category_code_id",
            action="assign_axial_to_category",
            success_template="Assigned axial code {child} to category {parent}",
        )

    @app.route("/gt/axial/<int:axial_code_id>/unassign", methods=["POST"])
    def unassign_axial_from_category(axial_code_id: int):
        if not project_supports_gt(get_active_project()):
            flash("GT Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        return unassign_code_parent(axial_code_id, "axial", "unassign_axial_from_category")

    @app.route("/gt/codes/<int:code_id>/edit")
    def edit_gt_code(code_id: int):
        active_project = get_active_project()
        if not project_supports_gt(active_project):
            return mode_notice("GT Workspace", active_project)
        code = get_code_for_project(code_id, active_project["id"])
        if code is None:
            flash("Code not found.", "error")
            abort(404)
        if code["code_type"] == "open":
            return redirect(url_for("edit_code", code_id=code_id))
        return render_template(
            "gt_code_edit.html",
            title=f"Edit GT Code: {code['name']}",
            active_page="gt",
            active_project=active_project,
            code=code,
        )

    @app.route("/gt/codes/<int:code_id>/edit", methods=["POST"])
    def update_gt_code(code_id: int):
        active_project = get_active_project()
        if not project_supports_gt(active_project):
            flash("GT Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        code = get_code_for_project(code_id, active_project["id"])
        if code is None or code["code_type"] not in {"axial", "category"}:
            flash("Code not found.", "error")
            abort(404)

        name = request.form.get("name", "").strip()
        if not name:
            flash("Code name is required.", "error")
            return redirect(url_for("edit_gt_code", code_id=code_id))
        values = (
            name,
            request.form.get("description", "").strip(),
            normalize_code_color(request.form.get("color", "")),
            request.form.get("definition", "").strip(),
            request.form.get("analytical_note", "").strip(),
            request.form.get("gt_conditions", "").strip(),
            request.form.get("gt_context", "").strip(),
            request.form.get("gt_actions_interactions", "").strip(),
            request.form.get("gt_consequences", "").strip(),
            request.form.get("gt_properties", "").strip(),
            request.form.get("gt_dimensions", "").strip(),
            request.form.get("gt_theoretical_note", "").strip(),
            code_id,
            active_project["id"],
        )
        execute_write(
            """
            UPDATE codes
            SET name = ?, description = ?, color = ?, definition = ?,
                analytical_note = ?, gt_conditions = ?, gt_context = ?,
                gt_actions_interactions = ?, gt_consequences = ?,
                gt_properties = ?, gt_dimensions = ?, gt_theoretical_note = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND project_id = ?
            """,
            values,
        )
        log_action(
            active_project["id"],
            "code",
            code_id,
            "update_gt_code",
            f"Updated GT code: {name}",
        )
        flash(f"Updated GT code: {name}", "success")
        return redirect(url_for("gt_workspace"))

    @app.route("/gt/compare")
    def gt_compare():
        active_project = get_active_project()
        if not project_supports_gt(active_project):
            return mode_notice("GT Workspace", active_project)
        open_codes = get_open_codes_for_project(active_project["id"])
        code_a = get_compare_code(request.args.get("code_a"), active_project["id"])
        code_b = get_compare_code(request.args.get("code_b"), active_project["id"])
        return render_template(
            "gt_compare.html",
            title="Constant Comparison",
            active_page="gt",
            active_project=active_project,
            open_codes=open_codes,
            code_a=code_a,
            code_b=code_b,
            memo_status_labels=MEMO_STATUSES,
            methodology_helper=get_methodology_helper(active_project, "gt_compare"),
        )

    @app.route("/cda")
    def cda_workspace():
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            return mode_notice("CDA Workspace", active_project)
        return render_template(
            "cda_workspace.html",
            title="CDA Workspace",
            active_page="cda",
            active_project=active_project,
            counts=get_cda_counts(active_project["id"]),
            markers=get_discourse_markers_for_project(active_project["id"]),
            actors=get_actors_for_project(active_project["id"]),
            marker_type_labels=CDA_MARKER_TYPES,
            actor_type_labels=ACTOR_TYPES,
            default_marker_color=DEFAULT_CDA_MARKER_COLOR,
            methodology_helper=get_methodology_helper(active_project, "cda_workspace"),
        )

    @app.route("/cda/markers/create", methods=["POST"])
    def create_discourse_marker():
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        name = request.form.get("name", "").strip()
        marker_type = request.form.get("marker_type", "").strip()
        description = request.form.get("description", "").strip()
        color = normalize_cda_color(request.form.get("color", ""))
        if not name:
            flash("CDA marker name is required.", "error")
            return redirect(url_for("cda_workspace"))
        if marker_type not in CDA_MARKER_TYPES:
            flash("Invalid marker.", "error")
            return redirect(url_for("cda_workspace"))
        marker_id = execute_write(
            """
            INSERT INTO discourse_markers (project_id, name, marker_type, description, color)
            VALUES (?, ?, ?, ?, ?)
            """,
            (active_project["id"], name, marker_type, description, color),
        )
        log_action(
            active_project["id"],
            "discourse_marker",
            marker_id,
            "create_discourse_marker",
            f"Created CDA marker: {name}",
        )
        flash(f"CDA marker created: {name}", "success")
        return redirect(url_for("cda_workspace"))

    @app.route("/cda/markers/<int:marker_id>/edit")
    def edit_discourse_marker(marker_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        marker = get_discourse_marker_for_project(marker_id, active_project["id"])
        if marker is None:
            flash("Invalid marker.", "error")
            return redirect(url_for("cda_workspace"))
        return render_template(
            "cda_marker_edit.html",
            title=f"Edit CDA Marker: {marker['name']}",
            active_page="cda",
            active_project=active_project,
            marker=marker,
            marker_type_labels=CDA_MARKER_TYPES,
        )

    @app.route("/cda/markers/<int:marker_id>/edit", methods=["POST"])
    def update_discourse_marker(marker_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        marker = get_discourse_marker_for_project(marker_id, active_project["id"])
        if marker is None:
            flash("Invalid marker.", "error")
            return redirect(url_for("cda_workspace"))

        name = request.form.get("name", "").strip()
        marker_type = request.form.get("marker_type", "").strip()
        description = request.form.get("description", "").strip()
        color = normalize_cda_color(request.form.get("color", ""))
        if not name:
            flash("CDA marker name is required.", "error")
            return redirect(url_for("edit_discourse_marker", marker_id=marker_id))
        if marker_type not in CDA_MARKER_TYPES:
            flash("Invalid marker.", "error")
            return redirect(url_for("edit_discourse_marker", marker_id=marker_id))
        execute_write(
            """
            UPDATE discourse_markers
            SET name = ?, marker_type = ?, description = ?, color = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND project_id = ?
            """,
            (name, marker_type, description, color, marker_id, active_project["id"]),
        )
        log_action(
            active_project["id"],
            "discourse_marker",
            marker_id,
            "update_discourse_marker",
            f"Updated CDA marker: {name}",
        )
        flash(f"CDA marker updated: {name}", "success")
        return redirect(url_for("cda_workspace"))

    @app.route("/cda/markers/<int:marker_id>/delete", methods=["POST"])
    def delete_discourse_marker(marker_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        marker = get_discourse_marker_for_project(marker_id, active_project["id"])
        if marker is None:
            flash("Invalid marker.", "error")
            return redirect(url_for("cda_workspace"))
        db = get_db()
        db.execute("DELETE FROM segment_discourse_markers WHERE marker_id = ?", (marker_id,))
        db.execute(
            "DELETE FROM discourse_markers WHERE id = ? AND project_id = ?",
            (marker_id, active_project["id"]),
        )
        db.commit()
        log_action(
            active_project["id"],
            "discourse_marker",
            marker_id,
            "delete_discourse_marker",
            f"Deleted CDA marker: {marker['name']}",
        )
        flash(f"CDA marker deleted: {marker['name']}", "success")
        return redirect(url_for("cda_workspace"))

    @app.route("/cda/actors/create", methods=["POST"])
    def create_actor():
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        name = request.form.get("name", "").strip()
        actor_type = request.form.get("actor_type", "").strip()
        description = request.form.get("description", "").strip()
        if not name:
            flash("Actor name is required.", "error")
            return redirect(url_for("cda_workspace"))
        if actor_type not in ACTOR_TYPES:
            flash("Invalid actor.", "error")
            return redirect(url_for("cda_workspace"))
        actor_id = execute_write(
            """
            INSERT INTO actors (project_id, name, actor_type, description)
            VALUES (?, ?, ?, ?)
            """,
            (active_project["id"], name, actor_type, description),
        )
        log_action(
            active_project["id"],
            "actor",
            actor_id,
            "create_actor",
            f"Created actor: {name}",
        )
        flash(f"Actor created: {name}", "success")
        return redirect(url_for("cda_workspace"))

    @app.route("/cda/actors/<int:actor_id>/edit")
    def edit_actor(actor_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        actor = get_actor_for_project(actor_id, active_project["id"])
        if actor is None:
            flash("Invalid actor.", "error")
            return redirect(url_for("cda_workspace"))
        return render_template(
            "cda_actor_edit.html",
            title=f"Edit Actor: {actor['name']}",
            active_page="cda",
            active_project=active_project,
            actor=actor,
            actor_type_labels=ACTOR_TYPES,
        )

    @app.route("/cda/actors/<int:actor_id>/edit", methods=["POST"])
    def update_actor(actor_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        actor = get_actor_for_project(actor_id, active_project["id"])
        if actor is None:
            flash("Invalid actor.", "error")
            return redirect(url_for("cda_workspace"))

        name = request.form.get("name", "").strip()
        actor_type = request.form.get("actor_type", "").strip()
        description = request.form.get("description", "").strip()
        if not name:
            flash("Actor name is required.", "error")
            return redirect(url_for("edit_actor", actor_id=actor_id))
        if actor_type not in ACTOR_TYPES:
            flash("Invalid actor.", "error")
            return redirect(url_for("edit_actor", actor_id=actor_id))
        execute_write(
            """
            UPDATE actors
            SET name = ?, actor_type = ?, description = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND project_id = ?
            """,
            (name, actor_type, description, actor_id, active_project["id"]),
        )
        log_action(
            active_project["id"],
            "actor",
            actor_id,
            "update_actor",
            f"Updated actor: {name}",
        )
        flash(f"Actor updated: {name}", "success")
        return redirect(url_for("cda_workspace"))

    @app.route("/cda/actors/<int:actor_id>/delete", methods=["POST"])
    def delete_actor(actor_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        actor = get_actor_for_project(actor_id, active_project["id"])
        if actor is None:
            flash("Invalid actor.", "error")
            return redirect(url_for("cda_workspace"))
        db = get_db()
        db.execute("DELETE FROM segment_actors WHERE actor_id = ?", (actor_id,))
        db.execute(
            "DELETE FROM actors WHERE id = ? AND project_id = ?",
            (actor_id, active_project["id"]),
        )
        db.commit()
        log_action(
            active_project["id"],
            "actor",
            actor_id,
            "delete_actor",
            f"Deleted actor: {actor['name']}",
        )
        flash(f"Actor deleted: {actor['name']}", "success")
        return redirect(url_for("cda_workspace"))

    @app.route("/segments/<int:segment_id>/discourse-markers", methods=["POST"])
    def assign_discourse_marker_to_segment(segment_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        segment = get_segment_for_project(segment_id, active_project["id"])
        if segment is None:
            flash("Invalid segment.", "error")
            abort(404)
        try:
            marker_id = int(request.form.get("marker_id", ""))
        except ValueError:
            flash("Invalid marker.", "error")
            return redirect_to_document(segment["document_id"])
        marker = get_discourse_marker_for_project(marker_id, active_project["id"])
        if marker is None:
            flash("Invalid marker.", "error")
            return redirect_to_document(segment["document_id"])
        if segment_has_discourse_marker(segment_id, marker_id):
            flash("CDA marker already assigned to this segment.", "error")
            return redirect_to_document(segment["document_id"])
        execute_write(
            """
            INSERT INTO segment_discourse_markers (segment_id, marker_id, note)
            VALUES (?, ?, ?)
            """,
            (segment_id, marker_id, request.form.get("note", "").strip()),
        )
        segment_label = segment["name"] or f"segment {segment_id}"
        log_action(
            active_project["id"],
            "segment",
            segment_id,
            "assign_discourse_marker",
            f"Assigned CDA marker {marker['name']} to {segment_label}",
        )
        flash(f"Marker assigned to segment: {marker['name']}", "success")
        return redirect_to_document(segment["document_id"])

    @app.route("/segments/<int:segment_id>/discourse-markers/<int:marker_id>/remove", methods=["POST"])
    def remove_discourse_marker_from_segment(segment_id: int, marker_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        segment = get_segment_for_project(segment_id, active_project["id"])
        marker = get_discourse_marker_for_project(marker_id, active_project["id"])
        if segment is None:
            flash("Invalid segment.", "error")
            abort(404)
        if marker is None:
            flash("Invalid marker.", "error")
            return redirect_to_document(segment["document_id"])
        execute_write(
            "DELETE FROM segment_discourse_markers WHERE segment_id = ? AND marker_id = ?",
            (segment_id, marker_id),
        )
        segment_label = segment["name"] or f"segment {segment_id}"
        log_action(
            active_project["id"],
            "segment",
            segment_id,
            "remove_discourse_marker",
            f"Removed CDA marker {marker['name']} from {segment_label}",
        )
        flash(f"Marker removed from segment: {marker['name']}", "success")
        return redirect_to_document(segment["document_id"])

    @app.route("/segments/<int:segment_id>/actors", methods=["POST"])
    def assign_actor_to_segment(segment_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        segment = get_segment_for_project(segment_id, active_project["id"])
        if segment is None:
            flash("Invalid segment.", "error")
            abort(404)
        try:
            actor_id = int(request.form.get("actor_id", ""))
        except ValueError:
            flash("Invalid actor.", "error")
            return redirect_to_document(segment["document_id"])
        relation_type = request.form.get("relation_type", "").strip()
        actor = get_actor_for_project(actor_id, active_project["id"])
        if actor is None:
            flash("Invalid actor.", "error")
            return redirect_to_document(segment["document_id"])
        if relation_type not in ACTOR_RELATION_TYPES:
            flash("Invalid actor relation.", "error")
            return redirect_to_document(segment["document_id"])
        execute_write(
            """
            INSERT INTO segment_actors (segment_id, actor_id, relation_type, note)
            VALUES (?, ?, ?, ?)
            """,
            (segment_id, actor_id, relation_type, request.form.get("note", "").strip()),
        )
        segment_label = segment["name"] or f"segment {segment_id}"
        log_action(
            active_project["id"],
            "segment",
            segment_id,
            "assign_actor_to_segment",
            f"Assigned actor {actor['name']} as {relation_type} to {segment_label}",
        )
        flash(f"Actor assigned to segment: {actor['name']}", "success")
        return redirect_to_document(segment["document_id"])

    @app.route("/segments/<int:segment_id>/actors/<int:segment_actor_id>/remove", methods=["POST"])
    def remove_actor_from_segment(segment_id: int, segment_actor_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        segment_actor = get_segment_actor_for_project(segment_actor_id, segment_id, active_project["id"])
        if segment_actor is None:
            flash("Invalid actor.", "error")
            abort(404)
        segment_label = segment_actor["segment_name"] or f"segment {segment_id}"
        execute_write(
            "DELETE FROM segment_actors WHERE id = ? AND segment_id = ?",
            (segment_actor_id, segment_id),
        )
        log_action(
            active_project["id"],
            "segment",
            segment_id,
            "remove_actor_from_segment",
            f"Removed actor annotation from {segment_label}",
        )
        flash("Actor annotation removed from segment.", "success")
        return redirect_to_document(segment_actor["document_id"])

    @app.route("/segments/<int:segment_id>/features", methods=["POST"])
    def create_discourse_feature(segment_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        segment = get_segment_for_project(segment_id, active_project["id"])
        if segment is None:
            flash("Invalid segment.", "error")
            abort(404)
        feature_type = request.form.get("feature_type", "").strip()
        value = request.form.get("value", "").strip()
        interpretation = request.form.get("interpretation", "").strip()
        if feature_type not in DISCOURSE_FEATURE_TYPES:
            flash("Invalid feature.", "error")
            return redirect_to_document(segment["document_id"])
        if not value:
            flash("Invalid feature.", "error")
            return redirect_to_document(segment["document_id"])
        feature_id = execute_write(
            """
            INSERT INTO discourse_features (segment_id, feature_type, value, interpretation)
            VALUES (?, ?, ?, ?)
            """,
            (segment_id, feature_type, value, interpretation),
        )
        segment_label = segment["name"] or f"segment {segment_id}"
        log_action(
            active_project["id"],
            "segment",
            segment_id,
            "create_discourse_feature",
            f"Created discourse feature {feature_type} for {segment_label}",
        )
        flash(f"Discourse feature created: {feature_type}", "success")
        return redirect_to_document(segment["document_id"])

    @app.route("/segments/<int:segment_id>/features/<int:feature_id>/delete", methods=["POST"])
    def delete_discourse_feature(segment_id: int, feature_id: int):
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            flash("CDA Workspace is disabled for this project's methodology mode.", "error")
            return redirect(url_for("dashboard"))
        feature = get_discourse_feature_for_project(feature_id, segment_id, active_project["id"])
        if feature is None:
            flash("Invalid feature.", "error")
            abort(404)
        segment_label = feature["segment_name"] or f"segment {segment_id}"
        execute_write(
            "DELETE FROM discourse_features WHERE id = ? AND segment_id = ?",
            (feature_id, segment_id),
        )
        log_action(
            active_project["id"],
            "segment",
            segment_id,
            "delete_discourse_feature",
            f"Deleted discourse feature {feature['feature_type']} from {segment_label}",
        )
        flash(f"Discourse feature deleted: {feature['feature_type']}", "success")
        return redirect_to_document(feature["document_id"])

    @app.route("/cda/features")
    def cda_features():
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            return mode_notice("CDA Workspace", active_project)
        filters = {
            "feature_type": request.args.get("feature_type", "").strip(),
            "document_id": request.args.get("document_id", "").strip(),
        }
        return render_template(
            "cda_features.html",
            title="CDA Feature Overview",
            active_page="cda",
            active_project=active_project,
            features=get_discourse_features_for_project(active_project["id"], filters),
            counts=get_discourse_feature_counts(active_project["id"]),
            documents=get_documents_for_memo_links(active_project["id"]),
            filters=filters,
            feature_type_labels=DISCOURSE_FEATURE_TYPES,
            methodology_helper=get_methodology_helper(active_project, "cda_features"),
        )

    @app.route("/cda/voice-silence")
    def cda_voice_silence():
        active_project = get_active_project()
        if not project_supports_cda(active_project):
            return mode_notice("CDA Workspace", active_project)
        return render_template(
            "cda_voice_silence.html",
            title="Voice and Silence Report",
            active_page="cda",
            active_project=active_project,
            rows=get_voice_silence_report(active_project["id"]),
            relation_type_labels=ACTOR_RELATION_TYPES,
            methodology_helper=get_methodology_helper(active_project, "cda_voice_silence"),
        )

    @app.route("/model")
    def model_builder():
        active_project = get_active_project()
        filters = {
            "relation_type": request.args.get("relation_type", "").strip(),
            "strength": request.args.get("strength", "").strip(),
            "source_type": request.args.get("source_type", "").strip(),
            "target_type": request.args.get("target_type", "").strip(),
            "involves_type": request.args.get("involves_type", "").strip(),
            "q": request.args.get("q", "").strip(),
        }
        entities = get_relation_entity_options(active_project["id"])
        return render_template(
            "model.html",
            title="Analytical Model",
            active_page="model",
            active_project=active_project,
            entities=entities,
            relations=get_relations_for_project(active_project["id"], filters),
            relation_counts=get_model_counts(active_project["id"]),
            filters=filters,
            relation_type_labels=RELATION_TYPES,
            relation_strength_labels=RELATION_STRENGTHS,
            entity_type_labels=RELATION_ENTITY_TYPES,
            research_questions=get_research_questions_for_project(active_project["id"]),
            mode_prompts=get_model_mode_prompts(active_project),
            methodology_helper=get_methodology_helper(active_project, "model"),
        )

    @app.route("/model/relations/create", methods=["POST"])
    def create_relation():
        active_project = get_active_project()
        data, error = validate_relation_form(active_project["id"])
        if error:
            flash(error, "error")
            return redirect(url_for("model_builder"))
        relation_id = execute_write(
            """
            INSERT INTO relations (
                project_id, source_type, source_id, target_type, target_id,
                relation_type, memo, title, strength, evidence_note, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            """,
            (
                active_project["id"],
                data["source_type"],
                data["source_id"],
                data["target_type"],
                data["target_id"],
                data["relation_type"],
                data["memo"],
                data["title"],
                data["strength"],
                data["evidence_note"],
            ),
        )
        log_action(
            active_project["id"],
            "relation",
            relation_id,
            "create_relation",
            f"Created relation: {data['source_label']} {data['relation_type']} {data['target_label']}",
        )
        flash("Relation created.", "success")
        return redirect(url_for("model_builder"))

    @app.route("/model/relations/<int:relation_id>/edit")
    def edit_relation(relation_id: int):
        active_project = get_active_project()
        relation = get_relation_for_project(relation_id, active_project["id"])
        if relation is None:
            flash("Relation not found.", "error")
            abort(404)
        return render_template(
            "relation_edit.html",
            title="Edit Relation",
            active_page="model",
            active_project=active_project,
            relation=relation,
            entities=get_relation_entity_options(active_project["id"]),
            relation_type_labels=RELATION_TYPES,
            relation_strength_labels=RELATION_STRENGTHS,
        )

    @app.route("/model/relations/<int:relation_id>/edit", methods=["POST"])
    def update_relation(relation_id: int):
        active_project = get_active_project()
        relation = get_relation_for_project(relation_id, active_project["id"])
        if relation is None:
            flash("Relation not found.", "error")
            abort(404)
        data, error = validate_relation_form(active_project["id"])
        if error:
            flash(error, "error")
            return redirect(url_for("edit_relation", relation_id=relation_id))
        execute_write(
            """
            UPDATE relations
            SET source_type = ?, source_id = ?, target_type = ?, target_id = ?,
                relation_type = ?, memo = ?, title = ?, strength = ?,
                evidence_note = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND project_id = ?
            """,
            (
                data["source_type"],
                data["source_id"],
                data["target_type"],
                data["target_id"],
                data["relation_type"],
                data["memo"],
                data["title"],
                data["strength"],
                data["evidence_note"],
                relation_id,
                active_project["id"],
            ),
        )
        log_action(
            active_project["id"],
            "relation",
            relation_id,
            "update_relation",
            f"Updated relation: {data['title'] or relation_id}",
        )
        flash("Relation updated.", "success")
        return redirect(url_for("model_builder"))

    @app.route("/model/relations/<int:relation_id>/delete", methods=["POST"])
    def delete_relation(relation_id: int):
        active_project = get_active_project()
        relation = get_relation_for_project(relation_id, active_project["id"])
        if relation is None:
            flash("Relation not found.", "error")
            abort(404)
        execute_write(
            "DELETE FROM relations WHERE id = ? AND project_id = ?",
            (relation_id, active_project["id"]),
        )
        log_action(
            active_project["id"],
            "relation",
            relation_id,
            "delete_relation",
            f"Deleted relation: {relation['title'] or relation_id}",
        )
        flash("Relation deleted.", "success")
        return redirect(url_for("model_builder"))

    @app.route("/model/entity/<entity_type>/<int:entity_id>")
    def model_entity(entity_type: str, entity_id: int):
        active_project = get_active_project()
        entity = get_entity_reference(entity_type, entity_id, active_project["id"])
        if entity is None:
            flash("Invalid model entity.", "error")
            abort(404)
        return render_template(
            "model_entity.html",
            title=f"Model Entity: {entity['label']}",
            active_page="model",
            active_project=active_project,
            entity=entity,
            outgoing_relations=get_entity_relations(active_project["id"], entity_type, entity_id, "outgoing"),
            incoming_relations=get_entity_relations(active_project["id"], entity_type, entity_id, "incoming"),
            related_memos=get_related_memos_for_model_entity(active_project["id"], entity_type, entity_id),
            mode_prompts=get_model_mode_prompts(active_project),
        )

    @app.route("/network")
    def network_explorer():
        active_project = get_active_project()
        filters = get_network_filters()
        return render_template(
            "network.html",
            title="Co-occurrence Network",
            active_page="network",
            active_project=active_project,
            filters=filters,
            documents=get_documents_for_project(active_project["id"]),
            code_type_labels={"": "All code types", **{value: value.title() for value in ["open", "axial", "category"]}},
            marker_type_labels={"": "All marker types", **CDA_MARKER_TYPES},
            actor_type_labels={"": "All actor types", **ACTOR_TYPES},
            feature_type_labels={"": "All feature types", **DISCOURSE_FEATURE_TYPES},
            layout_labels={"columns": "Columns", "force": "Force", "circle": "Circle"},
        )

    @app.route("/network/data")
    def network_data():
        active_project = get_active_project()
        return jsonify(build_cooccurrence_network(active_project, get_network_filters()))

    @app.route("/research-questions/create", methods=["POST"])
    def create_research_question():
        active_project = get_active_project()
        question = request.form.get("question", "").strip()
        note = request.form.get("note", "").strip()
        if not question:
            flash("Research question is required.", "error")
            return redirect(url_for("model_builder"))
        question_id = execute_write(
            """
            INSERT INTO research_questions (project_id, question, note)
            VALUES (?, ?, ?)
            """,
            (active_project["id"], question, note),
        )
        log_action(
            active_project["id"],
            "research_question",
            question_id,
            "create_research_question",
            f"Created research question: {truncate_text(question, 80)}",
        )
        flash("Research question created.", "success")
        return redirect(url_for("model_builder"))

    @app.route("/methodology")
    def methodology():
        active_project = get_active_project()
        filters = {
            "note_type": request.args.get("note_type", "").strip(),
            "methodology_area": request.args.get("methodology_area", "").strip(),
            "status": request.args.get("status", "").strip(),
            "linked_entity_type": request.args.get("linked_entity_type", "").strip(),
            "q": request.args.get("q", "").strip(),
        }
        return render_template(
            "methodology.html",
            title="Methodology",
            active_page="methodology",
            active_project=active_project,
            libraries=get_relevant_methodology_libraries(active_project),
            overview=get_methodology_overview(active_project),
            notes=get_methodology_notes_for_project(active_project["id"], filters),
            note_counts=get_methodology_note_counts(active_project["id"]),
            filters=filters,
            entity_options=get_methodology_entity_options(active_project["id"]),
            linked_entity_type_labels=METHODOLOGY_LINKED_ENTITY_TYPES,
        )

    @app.route("/methodology/library/<library_id>")
    def methodology_library(library_id: str):
        active_project = get_active_project()
        library = load_methodology_library(library_id)
        if library is None:
            flash("Methodology library not found.", "error")
            abort(404)
        return render_template(
            "methodology_library.html",
            title=library["title"],
            active_page="methodology",
            active_project=active_project,
            library=library,
        )

    @app.route("/methodology/notes/create", methods=["POST"])
    def create_methodology_note():
        active_project = get_active_project()
        data, error = validate_methodology_note_form(active_project["id"])
        if error:
            flash(error, "error")
            return redirect(safe_next_url(request.form.get("next_url", "")) or url_for("methodology"))
        note_id = execute_write(
            """
            INSERT INTO methodology_notes (
                project_id, title, body, note_type, linked_entity_type,
                linked_entity_id, methodology_area, status, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            """,
            (
                active_project["id"],
                data["title"],
                data["body"],
                data["note_type"],
                data["linked_entity_type"],
                data["linked_entity_id"],
                data["methodology_area"],
                data["status"],
            ),
        )
        log_action(
            active_project["id"],
            "methodology_note",
            note_id,
            "create_methodology_note",
            f"Created methodology note: {data['title']}",
        )
        flash("Methodology note created.", "success")
        return redirect(safe_next_url(request.form.get("next_url", "")) or url_for("methodology"))

    @app.route("/methodology/notes/<int:note_id>/edit")
    def edit_methodology_note(note_id: int):
        active_project = get_active_project()
        note = get_methodology_note_for_project(note_id, active_project["id"])
        if note is None:
            flash("Methodology note not found.", "error")
            abort(404)
        return render_template(
            "methodology_note_edit.html",
            title="Edit Methodology Note",
            active_page="methodology",
            active_project=active_project,
            note=note,
            entity_options=get_methodology_entity_options(active_project["id"]),
            linked_entity_type_labels=METHODOLOGY_LINKED_ENTITY_TYPES,
        )

    @app.route("/methodology/notes/<int:note_id>/edit", methods=["POST"])
    def update_methodology_note(note_id: int):
        active_project = get_active_project()
        note = get_methodology_note_for_project(note_id, active_project["id"])
        if note is None:
            flash("Methodology note not found.", "error")
            abort(404)
        data, error = validate_methodology_note_form(active_project["id"])
        if error:
            flash(error, "error")
            return redirect(url_for("edit_methodology_note", note_id=note_id))
        execute_write(
            """
            UPDATE methodology_notes
            SET title = ?, body = ?, note_type = ?, linked_entity_type = ?,
                linked_entity_id = ?, methodology_area = ?, status = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ? AND project_id = ?
            """,
            (
                data["title"],
                data["body"],
                data["note_type"],
                data["linked_entity_type"],
                data["linked_entity_id"],
                data["methodology_area"],
                data["status"],
                note_id,
                active_project["id"],
            ),
        )
        log_action(
            active_project["id"],
            "methodology_note",
            note_id,
            "update_methodology_note",
            f"Updated methodology note: {data['title']}",
        )
        flash("Methodology note updated.", "success")
        return redirect(url_for("methodology"))

    @app.route("/methodology/notes/<int:note_id>/delete", methods=["POST"])
    def delete_methodology_note(note_id: int):
        active_project = get_active_project()
        note = get_methodology_note_for_project(note_id, active_project["id"])
        if note is None:
            flash("Methodology note not found.", "error")
            abort(404)
        execute_write("DELETE FROM methodology_notes WHERE id = ? AND project_id = ?", (note_id, active_project["id"]))
        log_action(
            active_project["id"],
            "methodology_note",
            note_id,
            "delete_methodology_note",
            f"Deleted methodology note: {note['title']}",
        )
        flash("Methodology note deleted.", "success")
        return redirect(url_for("methodology"))

    @app.route("/exports")
    def exports():
        active_project = get_active_project()
        export_sections = [
            {
                "title": "Codebook exports",
                "cards": [
                    {
                        "title": "Codebook Markdown",
                        "description": "Complete codebook with hierarchy, GT fields, CDA markers, actors, usage counts, and memo counts.",
                        "format": "Markdown",
                        "endpoint": "export_codebook_markdown",
                        "button": "Download codebook",
                    },
                ],
            },
            {
                "title": "Segment exports",
                "cards": [
                    {
                        "title": "Coded segments CSV",
                        "description": "One CSV row per segment-code assignment.",
                        "format": "CSV",
                        "endpoint": "export_coded_segments_csv",
                        "button": "Download CSV",
                    },
                    {
                        "title": "Coded segments Markdown",
                        "description": "Readable coded segment report grouped by document.",
                        "format": "Markdown",
                        "endpoint": "export_coded_segments_markdown",
                        "button": "Download Markdown",
                    },
                ],
            },
            {
                "title": "Memo exports",
                "cards": [
                    {
                        "title": "Memos Markdown",
                        "description": "All project memos grouped by memo type and status.",
                        "format": "Markdown",
                        "endpoint": "export_memos_markdown",
                        "button": "Download memos",
                    },
                ],
            },
            {
                "title": "Methodology exports",
                "cards": [
                    {
                        "title": "Methodological protocol Markdown",
                        "description": "Project-specific methodology notes, sources, helper prompts, and protocol decisions.",
                        "format": "Markdown",
                        "endpoint": "export_methodology_protocol_markdown",
                        "button": "Download protocol",
                    },
                ],
            },
            {
                "title": "Grounded Theory exports",
                "cards": [
                    {
                        "title": "GT hierarchy Markdown",
                        "description": "Categories, axial codes, open codes, and representative segments.",
                        "format": "Markdown",
                        "endpoint": "export_gt_hierarchy_markdown",
                        "button": "Download GT hierarchy",
                    },
                ],
            },
            {
                "title": "CDA exports",
                "cards": [
                    {
                        "title": "CDA features CSV",
                        "description": "All discourse features with document and segment context.",
                        "format": "CSV",
                        "endpoint": "export_cda_features_csv",
                        "button": "Download features",
                    },
                    {
                        "title": "Actor voice/silence CSV",
                        "description": "Actor-level voice, silence, evaluation, and agency counts.",
                        "format": "CSV",
                        "endpoint": "export_voice_silence_csv",
                        "button": "Download report",
                    },
                ],
            },
            {
                "title": "Visual model exports",
                "cards": [
                    {
                        "title": "Mermaid flowchart",
                        "description": "Use in Markdown documents, GitHub, Obsidian, or Mermaid-compatible editors.",
                        "format": "MMD",
                        "endpoint": "export_model_mermaid",
                        "button": "Download Mermaid",
                    },
                    {
                        "title": "Graphviz DOT",
                        "description": "Use with graphviz: dot -Tpng model.dot -o model.png.",
                        "format": "DOT",
                        "endpoint": "export_model_dot",
                        "button": "Download DOT",
                    },
                    {
                        "title": "LaTeX/TikZ snippet",
                        "description": "Use in LaTeX documents; paste into a TikZ-enabled document.",
                        "format": "TikZ",
                        "endpoint": "export_model_tikz",
                        "button": "Download TikZ",
                    },
                    {
                        "title": "SVG diagram",
                        "description": "Open directly in a browser or vector editor.",
                        "format": "SVG",
                        "endpoint": "export_model_svg",
                        "button": "Download SVG",
                    },
                ],
            },
            {
                "title": "Co-occurrence network exports",
                "cards": [
                    {
                        "title": "Co-occurrence network JSON",
                        "description": "Empirical network generated from segment-level code, marker, actor, and feature co-presence.",
                        "format": "JSON",
                        "endpoint": "export_cooccurrence_network_json",
                        "button": "Download network JSON",
                    },
                    {
                        "title": "Co-occurrence edges CSV",
                        "description": "One row per co-occurring node pair, with edge weights and segment IDs.",
                        "format": "CSV",
                        "endpoint": "export_cooccurrence_edges_csv",
                        "button": "Download edge CSV",
                    },
                ],
            },
            {
                "title": "Project exports",
                "cards": [
                    {
                        "title": "Analytical model Markdown",
                        "description": "Readable relation model with summaries and entity-centered counts.",
                        "format": "Markdown",
                        "endpoint": "export_model_markdown",
                        "button": "Download model",
                    },
                    {
                        "title": "Analytical model JSON",
                        "description": "Structured relation model with entity labels for downstream use.",
                        "format": "JSON",
                        "endpoint": "export_model_json",
                        "button": "Download model JSON",
                    },
                    {
                        "title": "Project summary Markdown",
                        "description": "Readable project overview with counts, top items, and audit summary.",
                        "format": "Markdown",
                        "endpoint": "export_project_summary_markdown",
                        "button": "Download summary",
                    },
                    {
                        "title": "Full project JSON",
                        "description": "Structured active-project data for archiving and inspection.",
                        "format": "JSON",
                        "endpoint": "export_project_json",
                        "button": "Download JSON",
                    },
                    {
                        "title": "Complete research package ZIP",
                        "description": "Methodology-aware research exports bundled into one ZIP file.",
                        "format": "ZIP",
                        "endpoint": "export_project_package",
                        "button": "Download package",
                    },
                    {
                        "title": "Project backup ZIP",
                        "description": "Basic active-project backup bundle for local archiving. Restore is not implemented yet.",
                        "format": "ZIP",
                        "endpoint": "export_project_backup",
                        "button": "Download backup",
                    },
                ],
            },
        ]
        if not project_supports_gt(active_project):
            export_sections = [
                section
                for section in export_sections
                if section["title"] != "Grounded Theory exports"
            ]
        if not project_supports_cda(active_project):
            export_sections = [
                section for section in export_sections if section["title"] != "CDA exports"
            ]
        return render_template(
            "exports.html",
            title="Exports",
            active_page="exports",
            active_project=active_project,
            export_sections=export_sections,
        )

    @app.route("/exports/codebook.md")
    def export_codebook_markdown():
        active_project = get_active_project()
        return download_text(
            build_codebook_markdown(active_project),
            "discourseLab_codebook.md",
            "text/markdown; charset=utf-8",
        )

    @app.route("/exports/coded-segments.csv")
    def export_coded_segments_csv():
        active_project = get_active_project()
        return download_text(
            generate_coded_segments_csv(active_project),
            "discourseLab_coded_segments.csv",
            "text/csv; charset=utf-8",
        )

    @app.route("/exports/coded-segments.md")
    def export_coded_segments_markdown():
        active_project = get_active_project()
        return download_text(
            generate_coded_segments_markdown(active_project),
            "discourseLab_coded_segments.md",
            "text/markdown; charset=utf-8",
        )

    @app.route("/exports/memos.md")
    def export_memos_markdown():
        active_project = get_active_project()
        return download_text(
            generate_memos_markdown(active_project),
            "discourseLab_memos.md",
            "text/markdown; charset=utf-8",
        )

    @app.route("/exports/methodology-protocol.md")
    def export_methodology_protocol_markdown():
        active_project = get_active_project()
        return download_text(
            generate_methodology_protocol_markdown(active_project),
            "discourseLab_methodology_protocol.md",
            "text/markdown; charset=utf-8",
        )

    @app.route("/exports/gt-hierarchy.md")
    def export_gt_hierarchy_markdown():
        active_project = get_active_project()
        return download_text(
            generate_gt_hierarchy_markdown(active_project),
            "discourseLab_gt_hierarchy.md",
            "text/markdown; charset=utf-8",
        )

    @app.route("/exports/cda-features.csv")
    def export_cda_features_csv():
        active_project = get_active_project()
        return download_text(
            generate_cda_features_csv(active_project),
            "discourseLab_cda_features.csv",
            "text/csv; charset=utf-8",
        )

    @app.route("/exports/voice-silence.csv")
    def export_voice_silence_csv():
        active_project = get_active_project()
        return download_text(
            generate_voice_silence_csv(active_project),
            "discourseLab_voice_silence.csv",
            "text/csv; charset=utf-8",
        )

    @app.route("/exports/project-summary.md")
    def export_project_summary_markdown():
        active_project = get_active_project()
        return download_text(
            generate_project_summary_markdown(active_project),
            "discourseLab_project_summary.md",
            "text/markdown; charset=utf-8",
        )

    @app.route("/exports/model.md")
    def export_model_markdown():
        active_project = get_active_project()
        return download_text(
            generate_model_markdown(active_project),
            "discourseLab_analytical_model.md",
            "text/markdown; charset=utf-8",
        )

    @app.route("/exports/model.json")
    def export_model_json():
        active_project = get_active_project()
        return download_text(
            generate_model_json(active_project),
            "discourseLab_analytical_model.json",
            "application/json; charset=utf-8",
        )

    @app.route("/exports/model.mmd")
    def export_model_mermaid():
        active_project = get_active_project()
        return download_text(
            generate_model_mermaid(active_project, get_visual_export_filters()),
            "discourseLab_analytical_model.mmd",
            "text/plain; charset=utf-8",
        )

    @app.route("/exports/model.dot")
    def export_model_dot():
        active_project = get_active_project()
        return download_text(
            generate_model_dot(active_project, get_visual_export_filters()),
            "discourseLab_analytical_model.dot",
            "text/vnd.graphviz; charset=utf-8",
        )

    @app.route("/exports/model.tikz")
    def export_model_tikz():
        active_project = get_active_project()
        return download_text(
            generate_model_tikz(active_project, get_visual_export_filters()),
            "discourseLab_analytical_model.tikz",
            "text/plain; charset=utf-8",
        )

    @app.route("/exports/model.svg")
    def export_model_svg():
        active_project = get_active_project()
        return download_text(
            generate_model_svg(active_project, get_visual_export_filters()),
            "discourseLab_analytical_model.svg",
            "image/svg+xml; charset=utf-8",
        )

    @app.route("/exports/cooccurrence-network.json")
    def export_cooccurrence_network_json():
        active_project = get_active_project()
        return download_text(
            generate_cooccurrence_network_json(active_project, get_network_filters()),
            "discourseLab_cooccurrence_network.json",
            "application/json; charset=utf-8",
        )

    @app.route("/exports/cooccurrence-edges.csv")
    def export_cooccurrence_edges_csv():
        active_project = get_active_project()
        return download_text(
            generate_cooccurrence_edges_csv(active_project, get_network_filters()),
            "discourseLab_cooccurrence_edges.csv",
            "text/csv; charset=utf-8",
        )

    @app.route("/exports/project.json")
    def export_project_json():
        active_project = get_active_project()
        return download_text(
            generate_project_json(active_project),
            "discourseLab_project.json",
            "application/json; charset=utf-8",
        )

    @app.route("/exports/package.zip")
    def export_project_package():
        active_project = get_active_project()
        return download_binary(
            generate_project_package_zip(active_project),
            "discourseLab_research_package.zip",
            "application/zip",
        )

    @app.route("/exports/project-backup.zip")
    def export_project_backup():
        active_project = get_active_project()
        return download_binary(
            generate_project_backup_zip(active_project),
            "discourseLab_project_backup.zip",
            "application/zip",
        )

    @app.route("/admin/integrity")
    def integrity_check():
        active_project = get_active_project()
        return render_template(
            "integrity.html",
            title="Integrity Check",
            active_page="integrity",
            active_project=active_project,
            report=build_integrity_report(active_project),
        )

    @app.route("/health")
    def health():
        active_project = get_active_project()
        return jsonify(
            {
                "status": "ok",
                "app": APP_NAME,
                "version": APP_VERSION,
                "phase": APP_PHASE,
                "release_label": APP_RELEASE_LABEL,
                "active_project_id": active_project["id"],
                "methodology_mode": active_project["methodology_mode"],
            }
        )

    @app.errorhandler(404)
    def not_found(error):
        return (
            render_template(
                "error.html",
                title="Not Found",
                active_page="",
                message="The requested discourseLab page or document was not found.",
            ),
            404,
        )

    return app


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        g.db = sqlite3.connect(DATABASE)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


def close_db(error=None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db_if_needed() -> None:
    is_new_database = not DATABASE.exists()
    db = get_db()
    if is_new_database:
        schema_sql = SCHEMA.read_text(encoding="utf-8")
        db.executescript(schema_sql)
        db.commit()

    run_migrations()


def run_migrations() -> None:
    db = get_db()
    project_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(projects)").fetchall()
    }
    project_migrations = {
        "methodology_mode": "ALTER TABLE projects ADD COLUMN methodology_mode TEXT NOT NULL DEFAULT 'mixed'",
        "status": "ALTER TABLE projects ADD COLUMN status TEXT NOT NULL DEFAULT 'active'",
        "last_opened_at": "ALTER TABLE projects ADD COLUMN last_opened_at TEXT",
        "research_goal": "ALTER TABLE projects ADD COLUMN research_goal TEXT",
        "principal_investigator": "ALTER TABLE projects ADD COLUMN principal_investigator TEXT",
    }
    for column, sql in project_migrations.items():
        if column not in project_columns:
            db.execute(sql)
    db.execute(
        """
        UPDATE projects
        SET methodology_mode = 'mixed'
        WHERE methodology_mode IS NULL
           OR TRIM(methodology_mode) = ''
           OR methodology_mode NOT IN ('generic', 'gt', 'cda', 'mixed')
        """
    )
    db.execute(
        """
        UPDATE projects
        SET status = 'active'
        WHERE status IS NULL OR TRIM(status) = '' OR status NOT IN ('active', 'deleted')
        """
    )
    segment_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(segments)").fetchall()
    }
    if "name" not in segment_columns:
        db.execute("ALTER TABLE segments ADD COLUMN name TEXT")
        db.commit()
    code_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(codes)").fetchall()
    }
    for column in GT_COLUMNS:
        if column not in code_columns:
            db.execute(f"ALTER TABLE codes ADD COLUMN {column} TEXT")
    relation_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(relations)").fetchall()
    }
    relation_migrations = {
        "title": "ALTER TABLE relations ADD COLUMN title TEXT",
        "strength": "ALTER TABLE relations ADD COLUMN strength TEXT DEFAULT 'moderate'",
        "evidence_note": "ALTER TABLE relations ADD COLUMN evidence_note TEXT",
        "updated_at": "ALTER TABLE relations ADD COLUMN updated_at TEXT",
    }
    for column, sql in relation_migrations.items():
        if column not in relation_columns:
            db.execute(sql)
    db.execute(
        """
        UPDATE relations
        SET strength = 'moderate'
        WHERE strength IS NULL OR TRIM(strength) = ''
           OR strength NOT IN ('weak', 'moderate', 'strong', 'uncertain')
        """
    )
    db.execute(
        """
        UPDATE relations
        SET updated_at = created_at
        WHERE updated_at IS NULL OR TRIM(updated_at) = ''
        """
    )
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS discourse_markers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            marker_type TEXT NOT NULL,
            description TEXT,
            color TEXT NOT NULL DEFAULT '#7c9a45',
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (project_id) REFERENCES projects (id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS segment_discourse_markers (
            segment_id INTEGER NOT NULL,
            marker_id INTEGER NOT NULL,
            note TEXT,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (segment_id, marker_id),
            FOREIGN KEY (segment_id) REFERENCES segments (id) ON DELETE CASCADE,
            FOREIGN KEY (marker_id) REFERENCES discourse_markers (id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS actors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            actor_type TEXT NOT NULL,
            description TEXT,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (project_id) REFERENCES projects (id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS segment_actors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            segment_id INTEGER NOT NULL,
            actor_id INTEGER NOT NULL,
            relation_type TEXT NOT NULL,
            note TEXT,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (segment_id) REFERENCES segments (id) ON DELETE CASCADE,
            FOREIGN KEY (actor_id) REFERENCES actors (id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS discourse_features (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            segment_id INTEGER NOT NULL,
            feature_type TEXT NOT NULL,
            value TEXT NOT NULL,
            interpretation TEXT,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (segment_id) REFERENCES segments (id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS methodology_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            body TEXT NOT NULL,
            note_type TEXT NOT NULL,
            linked_entity_type TEXT,
            linked_entity_id INTEGER,
            methodology_area TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'active',
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (project_id) REFERENCES projects (id) ON DELETE CASCADE
        );

        CREATE INDEX IF NOT EXISTS idx_discourse_markers_project_id
            ON discourse_markers (project_id);
        CREATE INDEX IF NOT EXISTS idx_segment_discourse_markers_segment_id
            ON segment_discourse_markers (segment_id);
        CREATE INDEX IF NOT EXISTS idx_segment_discourse_markers_marker_id
            ON segment_discourse_markers (marker_id);
        CREATE INDEX IF NOT EXISTS idx_actors_project_id ON actors (project_id);
        CREATE INDEX IF NOT EXISTS idx_segment_actors_segment_id
            ON segment_actors (segment_id);
        CREATE INDEX IF NOT EXISTS idx_segment_actors_actor_id
            ON segment_actors (actor_id);
        CREATE INDEX IF NOT EXISTS idx_discourse_features_segment_id
            ON discourse_features (segment_id);
        CREATE INDEX IF NOT EXISTS idx_methodology_notes_project_id
            ON methodology_notes (project_id);
        CREATE INDEX IF NOT EXISTS idx_methodology_notes_type
            ON methodology_notes (note_type);
        CREATE INDEX IF NOT EXISTS idx_methodology_notes_status
            ON methodology_notes (status);
        """
    )
    for table in ("discourse_markers", "actors"):
        columns = {row["name"] for row in db.execute(f"PRAGMA table_info({table})").fetchall()}
        if "updated_at" not in columns:
            db.execute(f"ALTER TABLE {table} ADD COLUMN updated_at TEXT")
        db.execute(
            f"""
            UPDATE {table}
            SET updated_at = created_at
            WHERE updated_at IS NULL OR TRIM(updated_at) = ''
            """
        )
    db.commit()


def query_one(sql: str, params: tuple = ()) -> sqlite3.Row | None:
    return get_db().execute(sql, params).fetchone()


def query_all(sql: str, params: tuple = ()) -> list[sqlite3.Row]:
    return get_db().execute(sql, params).fetchall()


def execute_write(sql: str, params: tuple = ()) -> int:
    cursor = get_db().execute(sql, params)
    get_db().commit()
    return cursor.lastrowid


def get_active_project() -> sqlite3.Row:
    if has_request_context():
        active_project_id = session.get("active_project_id")
        if isinstance(active_project_id, int):
            project = get_project(active_project_id)
            if project is not None:
                return project
        elif isinstance(active_project_id, str) and active_project_id.isdigit():
            project = get_project(int(active_project_id))
            if project is not None:
                session["active_project_id"] = int(active_project_id)
                return project

    project = query_one(
        """
        SELECT id, name, description, methodology_mode, status, last_opened_at,
               research_goal, principal_investigator, created_at, updated_at
        FROM projects
        WHERE status != 'deleted'
        ORDER BY datetime(COALESCE(last_opened_at, created_at)) DESC, id DESC
        LIMIT 1
        """
    )
    if project is not None:
        if has_request_context():
            session["active_project_id"] = project["id"]
        return project

    project_id = execute_write(
        """
        INSERT INTO projects (
            name, description, methodology_mode, status, last_opened_at
        )
        VALUES (?, ?, 'mixed', 'active', CURRENT_TIMESTAMP)
        """,
        (DEFAULT_PROJECT_NAME, DEFAULT_PROJECT_DESCRIPTION),
    )
    log_action(
        project_id=project_id,
        entity_type="project",
        entity_id=project_id,
        action="create_default_project",
        details="Created default discourseLab project.",
    )
    project = query_one(
        """
        SELECT id, name, description, methodology_mode, status, last_opened_at,
               research_goal, principal_investigator, created_at, updated_at
        FROM projects
        WHERE id = ?
        """,
        (project_id,),
    )
    if has_request_context():
        session["active_project_id"] = project_id
    return project


def get_project(project_id: int) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT id, name, description, methodology_mode, status, last_opened_at,
               research_goal, principal_investigator, created_at, updated_at
        FROM projects
        WHERE id = ? AND status != 'deleted'
        """,
        (project_id,),
    )


def get_projects() -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT id, name, description, methodology_mode, status, last_opened_at,
               research_goal, principal_investigator, created_at, updated_at
        FROM projects
        WHERE status != 'deleted'
        ORDER BY datetime(COALESCE(last_opened_at, created_at)) DESC, name COLLATE NOCASE
        """
    )


def set_active_project(project_id: int) -> None:
    execute_write(
        "UPDATE projects SET last_opened_at = CURRENT_TIMESTAMP WHERE id = ?",
        (project_id,),
    )
    if has_request_context():
        session["active_project_id"] = project_id


def project_supports_gt(project: sqlite3.Row | dict) -> bool:
    return project["methodology_mode"] in {"gt", "mixed"}


def project_supports_cda(project: sqlite3.Row | dict) -> bool:
    return project["methodology_mode"] in {"cda", "mixed"}


def mode_notice(feature_name: str, active_project: sqlite3.Row):
    return render_template(
        "disabled_workspace.html",
        title=f"{feature_name} Disabled",
        active_page="",
        active_project=active_project,
        feature_name=feature_name,
        mode_label=METHODOLOGY_MODES.get(
            active_project["methodology_mode"], active_project["methodology_mode"]
        ),
    )


def get_dashboard_counts(project_id: int) -> dict[str, int]:
    return {
        "documents": query_one(
            "SELECT COUNT(*) AS count FROM documents WHERE project_id = ?", (project_id,)
        )["count"],
        "open_codes": query_one(
            """
            SELECT COUNT(*) AS count
            FROM codes
            WHERE project_id = ? AND code_type = 'open'
            """,
            (project_id,),
        )["count"],
        "axial_codes": query_one(
            "SELECT COUNT(*) AS count FROM codes WHERE project_id = ? AND code_type = 'axial'",
            (project_id,),
        )["count"],
        "categories": query_one(
            "SELECT COUNT(*) AS count FROM codes WHERE project_id = ? AND code_type = 'category'",
            (project_id,),
        )["count"],
        "cda_markers": query_one(
            "SELECT COUNT(*) AS count FROM discourse_markers WHERE project_id = ?",
            (project_id,),
        )["count"],
        "actors": query_one(
            "SELECT COUNT(*) AS count FROM actors WHERE project_id = ?", (project_id,)
        )["count"],
        "discourse_features": query_one(
            """
            SELECT COUNT(*) AS count
            FROM discourse_features
            JOIN segments ON segments.id = discourse_features.segment_id
            JOIN documents ON documents.id = segments.document_id
            WHERE documents.project_id = ?
            """,
            (project_id,),
        )["count"],
        "segments": query_one(
            """
            SELECT COUNT(*) AS count
            FROM segments
            JOIN documents ON documents.id = segments.document_id
            WHERE documents.project_id = ?
            """,
            (project_id,),
        )["count"],
        "coded_segments": query_one(
            """
            SELECT COUNT(DISTINCT segments.id) AS count
            FROM segments
            JOIN documents ON documents.id = segments.document_id
            JOIN segment_codes ON segment_codes.segment_id = segments.id
            WHERE documents.project_id = ?
            """,
            (project_id,),
        )["count"],
        "memos": query_one(
            "SELECT COUNT(*) AS count FROM memos WHERE project_id = ?", (project_id,)
        )["count"],
        "research_questions": query_one(
            "SELECT COUNT(*) AS count FROM research_questions WHERE project_id = ?",
            (project_id,),
        )["count"],
        "relations": query_one(
            "SELECT COUNT(*) AS count FROM relations WHERE project_id = ?", (project_id,)
        )["count"],
    }


def percent(numerator: int, denominator: int) -> int:
    if denominator <= 0:
        return 0
    return round((numerator / denominator) * 100)


def dashboard_barometer(
    label: str,
    numerator: int,
    denominator: int,
    zero_text: str,
) -> dict[str, int | str]:
    return {
        "label": label,
        "numerator": numerator,
        "denominator": denominator,
        "percent": percent(numerator, denominator),
        "summary": zero_text if denominator == 0 else f"{numerator} of {denominator}",
    }


def get_dashboard_progress_barometers(
    active_project: sqlite3.Row,
    counts: dict[str, int],
) -> list[dict[str, int | str]]:
    project_id = active_project["id"]
    barometers = [
        dashboard_barometer(
            "Documents engaged",
            count_documents_with_segments(project_id),
            counts["documents"],
            "No documents yet.",
        ),
        dashboard_barometer(
            "Segments coded",
            count_segments_with_codes(project_id),
            counts["segments"],
            "No segments yet.",
        ),
        dashboard_barometer(
            "Codebook completeness",
            count_codes_with_definitions(project_id),
            counts["open_codes"] + counts["axial_codes"] + counts["categories"],
            "No codes yet.",
        ),
        dashboard_barometer(
            "Memo coverage",
            count_segments_with_memos(project_id),
            counts["segments"],
            "No segments yet.",
        ),
    ]

    if project_supports_gt(active_project):
        barometers.extend(
            [
                dashboard_barometer(
                    "Open codes assigned to axial codes",
                    count_open_codes_with_axial(project_id),
                    counts["open_codes"],
                    "No open codes yet.",
                ),
                dashboard_barometer(
                    "Axial codes assigned to categories",
                    count_axial_codes_with_category(project_id),
                    counts["axial_codes"],
                    "No axial codes yet.",
                ),
            ]
        )

    if project_supports_cda(active_project):
        barometers.extend(
            [
                dashboard_barometer(
                    "Segments with CDA annotations",
                    count_segments_with_cda_annotations(project_id),
                    counts["segments"],
                    "No segments yet.",
                ),
                dashboard_barometer(
                    "Actors used",
                    count_actors_used(project_id),
                    counts["actors"],
                    "No actors yet.",
                ),
            ]
        )

    return barometers


def get_dashboard_suggestions(
    active_project: sqlite3.Row,
    counts: dict[str, int],
    methodology_counts: dict[str, int],
    progress_barometers: list[dict[str, int | str]],
) -> list[str]:
    project_id = active_project["id"]
    progress = {item["label"]: item for item in progress_barometers}
    suggestions = []
    if counts["documents"] == 0:
        suggestions.append("Import your first document.")
    if counts["documents"] > 0 and counts["segments"] == 0:
        suggestions.append("Open a document and create your first segment.")
    if counts["segments"] > 0 and progress["Segments coded"]["percent"] < 70:
        suggestions.append("Assign open codes to uncoded segments.")
    if counts["open_codes"] + counts["axial_codes"] + counts["categories"] > 0 and progress["Codebook completeness"]["percent"] < 100:
        suggestions.append("Complete your codebook definitions.")
    if project_supports_gt(active_project):
        if counts["open_codes"] > 0 and count_open_codes_with_axial(project_id) < counts["open_codes"]:
            suggestions.append("Assign open codes to axial codes.")
        if counts["axial_codes"] > 0 and count_axial_codes_with_category(project_id) < counts["axial_codes"]:
            suggestions.append("Develop categories from axial codes.")
    if project_supports_cda(active_project) and counts["actors"] == 0:
        suggestions.append("Create actors and start voice/silence mapping.")
    if counts["relations"] == 0:
        suggestions.append("Create analytical relations in the model builder.")
    if methodology_counts["active_protocol"] == 0:
        suggestions.append("Add a methodological protocol note.")
    return suggestions[:6]


def get_first_document_id(project_id: int) -> int | None:
    row = query_one(
        """
        SELECT id
        FROM documents
        WHERE project_id = ?
        ORDER BY datetime(created_at), id
        LIMIT 1
        """,
        (project_id,),
    )
    return row["id"] if row else None


def count_documents_with_segments(project_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(DISTINCT documents.id) AS count
        FROM documents
        JOIN segments ON segments.document_id = documents.id
        WHERE documents.project_id = ?
        """,
        (project_id,),
    )["count"]


def count_segments_with_codes(project_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(DISTINCT segments.id) AS count
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        JOIN segment_codes ON segment_codes.segment_id = segments.id
        WHERE documents.project_id = ?
        """,
        (project_id,),
    )["count"]


def count_codes_with_definitions(project_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(*) AS count
        FROM codes
        WHERE project_id = ?
          AND definition IS NOT NULL
          AND TRIM(definition) != ''
        """,
        (project_id,),
    )["count"]


def count_segments_with_memos(project_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(DISTINCT segments.id) AS count
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        JOIN memos
            ON memos.linked_entity_type = 'segment'
           AND memos.linked_entity_id = segments.id
           AND memos.project_id = documents.project_id
        WHERE documents.project_id = ?
        """,
        (project_id,),
    )["count"]


def count_open_codes_with_axial(project_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(open_codes.id) AS count
        FROM codes AS open_codes
        JOIN codes AS axial_codes ON axial_codes.id = open_codes.parent_id
        WHERE open_codes.project_id = ?
          AND open_codes.code_type = 'open'
          AND axial_codes.project_id = open_codes.project_id
          AND axial_codes.code_type = 'axial'
        """,
        (project_id,),
    )["count"]


def count_axial_codes_with_category(project_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(axial_codes.id) AS count
        FROM codes AS axial_codes
        JOIN codes AS categories ON categories.id = axial_codes.parent_id
        WHERE axial_codes.project_id = ?
          AND axial_codes.code_type = 'axial'
          AND categories.project_id = axial_codes.project_id
          AND categories.code_type = 'category'
        """,
        (project_id,),
    )["count"]


def count_segments_with_cda_annotations(project_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(DISTINCT segments.id) AS count
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ?
          AND (
              EXISTS (
                  SELECT 1
                  FROM segment_discourse_markers
                  WHERE segment_discourse_markers.segment_id = segments.id
              )
              OR EXISTS (
                  SELECT 1
                  FROM segment_actors
                  WHERE segment_actors.segment_id = segments.id
              )
              OR EXISTS (
                  SELECT 1
                  FROM discourse_features
                  WHERE discourse_features.segment_id = segments.id
              )
          )
        """,
        (project_id,),
    )["count"]


def count_actors_used(project_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(DISTINCT actors.id) AS count
        FROM actors
        JOIN segment_actors ON segment_actors.actor_id = actors.id
        WHERE actors.project_id = ?
        """,
        (project_id,),
    )["count"]


def get_cda_dashboard_preview(project_id: int) -> dict[str, str | int]:
    relation = query_one(
        """
        SELECT segment_actors.relation_type, COUNT(*) AS count
        FROM segment_actors
        JOIN segments ON segments.id = segment_actors.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ?
        GROUP BY segment_actors.relation_type
        ORDER BY count DESC, segment_actors.relation_type
        LIMIT 1
        """,
        (project_id,),
    )
    feature = query_one(
        """
        SELECT discourse_features.feature_type, COUNT(*) AS count
        FROM discourse_features
        JOIN segments ON segments.id = discourse_features.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ?
        GROUP BY discourse_features.feature_type
        ORDER BY count DESC, discourse_features.feature_type
        LIMIT 1
        """,
        (project_id,),
    )
    return {
        "segments_with_markers": query_one(
            """
            SELECT COUNT(DISTINCT segments.id) AS count
            FROM segments
            JOIN documents ON documents.id = segments.document_id
            JOIN segment_discourse_markers
                ON segment_discourse_markers.segment_id = segments.id
            WHERE documents.project_id = ?
            """,
            (project_id,),
        )["count"],
        "actors_annotated": query_one(
            """
            SELECT COUNT(DISTINCT actors.id) AS count
            FROM actors
            JOIN segment_actors ON segment_actors.actor_id = actors.id
            WHERE actors.project_id = ?
            """,
            (project_id,),
        )["count"],
        "top_relation_type": relation["relation_type"] if relation else "None yet",
        "top_feature_type": feature["feature_type"] if feature else "None yet",
    }


def get_model_dashboard_preview(project_id: int) -> dict:
    top_relation = query_one(
        """
        SELECT relation_type, COUNT(*) AS count
        FROM relations
        WHERE project_id = ?
        GROUP BY relation_type
        ORDER BY count DESC, relation_type
        LIMIT 1
        """,
        (project_id,),
    )
    latest = get_relations_for_project(project_id, {}, limit=5)
    return {
        **get_model_counts(project_id),
        "top_relation_type": top_relation["relation_type"] if top_relation else "None yet",
        "latest_relations": latest,
    }


def get_model_counts(project_id: int) -> dict[str, int]:
    return {
        "total": query_one(
            "SELECT COUNT(*) AS count FROM relations WHERE project_id = ?", (project_id,)
        )["count"],
        "strong": query_one(
            "SELECT COUNT(*) AS count FROM relations WHERE project_id = ? AND strength = 'strong'",
            (project_id,),
        )["count"],
        "uncertain": query_one(
            "SELECT COUNT(*) AS count FROM relations WHERE project_id = ? AND strength = 'uncertain'",
            (project_id,),
        )["count"],
        "gt_categories": query_one(
            """
            SELECT COUNT(*) AS count
            FROM relations
            LEFT JOIN codes AS source_code
                ON source_code.id = relations.source_id
               AND relations.source_type = 'code'
               AND source_code.project_id = relations.project_id
            LEFT JOIN codes AS target_code
                ON target_code.id = relations.target_id
               AND relations.target_type = 'code'
               AND target_code.project_id = relations.project_id
            WHERE relations.project_id = ?
              AND (source_code.code_type = 'category' OR target_code.code_type = 'category')
            """,
            (project_id,),
        )["count"],
        "cda_entities": query_one(
            """
            SELECT COUNT(*) AS count
            FROM relations
            WHERE project_id = ?
              AND (
                source_type IN ('actor', 'discourse_marker', 'discourse_feature')
                OR target_type IN ('actor', 'discourse_marker', 'discourse_feature')
              )
            """,
            (project_id,),
        )["count"],
        "negative_cases": query_one(
            """
            SELECT COUNT(*) AS count
            FROM relations
            WHERE project_id = ? AND relation_type = 'is_negative_case_for'
            """,
            (project_id,),
        )["count"],
    }


def get_documents_for_project(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT
            documents.id,
            documents.title,
            documents.original_filename,
            documents.file_type,
            documents.created_at,
            LENGTH(COALESCE(documents.text_content, '')) AS text_length,
            COUNT(DISTINCT segments.id) AS segment_count,
            COUNT(DISTINCT coded_segments.id) AS coded_segment_count
        FROM documents
        LEFT JOIN segments ON segments.document_id = documents.id
        LEFT JOIN (
            SELECT DISTINCT segments.id, segments.document_id
            FROM segments
            JOIN segment_codes ON segment_codes.segment_id = segments.id
        ) AS coded_segments ON coded_segments.document_id = documents.id
        WHERE documents.project_id = ?
        GROUP BY documents.id
        ORDER BY datetime(documents.created_at) DESC, documents.id DESC
        """,
        (project_id,),
    )


def get_latest_documents(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT title, file_type, created_at
        FROM documents
        WHERE project_id = ?
        ORDER BY datetime(created_at) DESC, id DESC
        LIMIT 5
        """,
        (project_id,),
    )


def get_latest_segments(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT
            COALESCE(segments.name, '') AS name,
            segments.selected_text,
            segments.created_at,
            documents.title AS document_title
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ?
        ORDER BY datetime(segments.created_at) DESC, segments.id DESC
        LIMIT 5
        """,
        (project_id,),
    )


def get_latest_open_codes(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT
            codes.id,
            codes.name,
            codes.color,
            codes.created_at,
            COUNT(segment_codes.segment_id) AS assigned_segment_count
        FROM codes
        LEFT JOIN segment_codes ON segment_codes.code_id = codes.id
        WHERE codes.project_id = ? AND codes.code_type = 'open'
        GROUP BY codes.id
        ORDER BY datetime(codes.created_at) DESC, codes.id DESC
        LIMIT 5
        """,
        (project_id,),
    )


def get_latest_coded_segments(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT
            segments.id,
            COALESCE(segments.name, '') AS name,
            segments.selected_text,
            documents.title AS document_title,
            GROUP_CONCAT(codes.name, ', ') AS code_names,
            segments.created_at
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        JOIN segment_codes ON segment_codes.segment_id = segments.id
        JOIN codes ON codes.id = segment_codes.code_id
        WHERE documents.project_id = ? AND codes.code_type = 'open'
        GROUP BY segments.id
        ORDER BY datetime(segments.created_at) DESC, segments.id DESC
        LIMIT 5
        """,
        (project_id,),
    )


def get_latest_memos(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT title, memo_type, status, created_at
        FROM memos
        WHERE project_id = ?
        ORDER BY datetime(created_at) DESC, id DESC
        LIMIT 5
        """,
        (project_id,),
    )


def get_codes_missing_definitions(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT id, name
        FROM codes
        WHERE project_id = ? AND code_type = 'open'
          AND (definition IS NULL OR TRIM(definition) = '')
        ORDER BY name COLLATE NOCASE
        LIMIT 5
        """,
        (project_id,),
    )


def get_document_for_project(document_id: int, project_id: int) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT id, project_id, title, original_filename, file_type, text_content,
               note, created_at, updated_at
        FROM documents
        WHERE id = ? AND project_id = ?
        """,
        (document_id, project_id),
    )


def get_document_segment_count(document_id: int) -> int:
    return query_one(
        "SELECT COUNT(*) AS count FROM segments WHERE document_id = ?", (document_id,)
    )["count"]


def get_segments_for_document(document_id: int, project_id: int | None = None) -> list[dict]:
    segment_rows = query_all(
        """
        SELECT id, document_id, COALESCE(name, '') AS name, selected_text, start_offset, end_offset,
               note, created_at, updated_at
        FROM segments
        WHERE document_id = ?
        ORDER BY start_offset ASC, end_offset ASC, id ASC
        """,
        (document_id,),
    )
    segments = []
    for row in segment_rows:
        segment = dict(row)
        segment["codes"] = get_codes_for_segment(row["id"])
        segment["discourse_markers"] = get_discourse_markers_for_segment(row["id"])
        segment["actors"] = get_segment_actors_for_segment(row["id"])
        segment["features"] = get_discourse_features_for_segment(row["id"])
        segment["memos"] = get_memos_for_entity(project_id, "segment", row["id"])
        segment["relations"] = get_relations_for_entity(project_id, "segment", row["id"], limit=3) if project_id else []
        segment["relation_count"] = get_relation_count_for_entity(project_id, "segment", row["id"]) if project_id else 0
        segments.append(segment)
    return segments


def get_codes_for_segment(segment_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT codes.id, codes.name, codes.color
        FROM codes
        JOIN segment_codes ON segment_codes.code_id = codes.id
        WHERE segment_codes.segment_id = ? AND codes.code_type = 'open'
        ORDER BY codes.name COLLATE NOCASE
        """,
        (segment_id,),
    )


def get_discourse_markers_for_segment(segment_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT discourse_markers.id, discourse_markers.name,
               discourse_markers.marker_type, discourse_markers.color,
               segment_discourse_markers.note,
               segment_discourse_markers.created_at AS assigned_at
        FROM discourse_markers
        JOIN segment_discourse_markers
            ON segment_discourse_markers.marker_id = discourse_markers.id
        WHERE segment_discourse_markers.segment_id = ?
        ORDER BY discourse_markers.name COLLATE NOCASE
        """,
        (segment_id,),
    )


def get_segment_actors_for_segment(segment_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT segment_actors.id, segment_actors.relation_type, segment_actors.note,
               segment_actors.created_at, actors.name, actors.actor_type
        FROM segment_actors
        JOIN actors ON actors.id = segment_actors.actor_id
        WHERE segment_actors.segment_id = ?
        ORDER BY actors.name COLLATE NOCASE, segment_actors.relation_type
        """,
        (segment_id,),
    )


def get_discourse_features_for_segment(segment_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT id, feature_type, value, interpretation, created_at, updated_at
        FROM discourse_features
        WHERE segment_id = ?
        ORDER BY datetime(created_at) DESC, id DESC
        """,
        (segment_id,),
    )


def get_discourse_markers_for_project(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT discourse_markers.id, discourse_markers.name,
               discourse_markers.marker_type, discourse_markers.description,
               discourse_markers.color, discourse_markers.created_at,
               discourse_markers.updated_at,
               COUNT(segment_discourse_markers.segment_id) AS assigned_segment_count
        FROM discourse_markers
        LEFT JOIN segment_discourse_markers
            ON segment_discourse_markers.marker_id = discourse_markers.id
        WHERE discourse_markers.project_id = ?
        GROUP BY discourse_markers.id
        ORDER BY discourse_markers.name COLLATE NOCASE
        """,
        (project_id,),
    )


def get_discourse_marker_for_project(
    marker_id: int, project_id: int
) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT id, project_id, name, marker_type, description, color,
               created_at, updated_at
        FROM discourse_markers
        WHERE id = ? AND project_id = ?
        """,
        (marker_id, project_id),
    )


def get_actors_for_project(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT actors.id, actors.name, actors.actor_type, actors.description,
               actors.created_at, actors.updated_at,
               COUNT(DISTINCT segment_actors.id) AS annotation_count,
               COUNT(DISTINCT relations.id) AS relation_count
        FROM actors
        LEFT JOIN segment_actors ON segment_actors.actor_id = actors.id
        LEFT JOIN relations ON relations.project_id = actors.project_id
            AND (
                (relations.source_type = 'actor' AND relations.source_id = actors.id)
                OR (relations.target_type = 'actor' AND relations.target_id = actors.id)
            )
        WHERE actors.project_id = ?
        GROUP BY actors.id
        ORDER BY actors.name COLLATE NOCASE
        """,
        (project_id,),
    )


def get_actor_for_project(actor_id: int, project_id: int) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT id, project_id, name, actor_type, description, created_at, updated_at
        FROM actors
        WHERE id = ? AND project_id = ?
        """,
        (actor_id, project_id),
    )


def segment_has_discourse_marker(segment_id: int, marker_id: int) -> bool:
    return (
        query_one(
            """
            SELECT 1
            FROM segment_discourse_markers
            WHERE segment_id = ? AND marker_id = ?
            """,
            (segment_id, marker_id),
        )
        is not None
    )


def get_segment_actor_for_project(
    segment_actor_id: int, segment_id: int, project_id: int
) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT segment_actors.id, segment_actors.segment_id, segment_actors.actor_id,
               segment_actors.relation_type, actors.name AS actor_name,
               segments.document_id, COALESCE(segments.name, '') AS segment_name
        FROM segment_actors
        JOIN actors ON actors.id = segment_actors.actor_id
        JOIN segments ON segments.id = segment_actors.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE segment_actors.id = ? AND segment_actors.segment_id = ?
          AND documents.project_id = ?
        """,
        (segment_actor_id, segment_id, project_id),
    )


def get_discourse_feature_for_project(
    feature_id: int, segment_id: int, project_id: int
) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT discourse_features.id, discourse_features.segment_id,
               discourse_features.feature_type, discourse_features.value,
               segments.document_id, COALESCE(segments.name, '') AS segment_name
        FROM discourse_features
        JOIN segments ON segments.id = discourse_features.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE discourse_features.id = ? AND discourse_features.segment_id = ?
          AND documents.project_id = ?
        """,
        (feature_id, segment_id, project_id),
    )


def get_cda_counts(project_id: int) -> dict[str, int]:
    return {
        "markers": query_one(
            "SELECT COUNT(*) AS count FROM discourse_markers WHERE project_id = ?",
            (project_id,),
        )["count"],
        "actors": query_one(
            "SELECT COUNT(*) AS count FROM actors WHERE project_id = ?", (project_id,)
        )["count"],
        "segments_with_markers": query_one(
            """
            SELECT COUNT(DISTINCT segments.id) AS count
            FROM segments
            JOIN documents ON documents.id = segments.document_id
            JOIN segment_discourse_markers
                ON segment_discourse_markers.segment_id = segments.id
            WHERE documents.project_id = ?
            """,
            (project_id,),
        )["count"],
        "segments_with_actor_annotations": query_one(
            """
            SELECT COUNT(DISTINCT segments.id) AS count
            FROM segments
            JOIN documents ON documents.id = segments.document_id
            JOIN segment_actors ON segment_actors.segment_id = segments.id
            WHERE documents.project_id = ?
            """,
            (project_id,),
        )["count"],
        "metaphor_features": get_feature_type_count(project_id, "metaphor"),
        "presupposition_features": get_feature_type_count(project_id, "presupposition"),
        "legitimation_features": get_feature_type_count(project_id, "legitimation"),
    }


def get_feature_type_count(project_id: int, feature_type: str) -> int:
    return query_one(
        """
        SELECT COUNT(*) AS count
        FROM discourse_features
        JOIN segments ON segments.id = discourse_features.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ? AND discourse_features.feature_type = ?
        """,
        (project_id, feature_type),
    )["count"]


def get_discourse_features_for_project(project_id: int, filters: dict) -> list[sqlite3.Row]:
    sql = """
        SELECT discourse_features.id, discourse_features.feature_type,
               discourse_features.value, discourse_features.interpretation,
               discourse_features.created_at,
               COALESCE(segments.name, '') AS segment_name,
               segments.selected_text, documents.id AS document_id,
               documents.title AS document_title
        FROM discourse_features
        JOIN segments ON segments.id = discourse_features.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ?
    """
    params = [project_id]
    if filters.get("feature_type") in DISCOURSE_FEATURE_TYPES:
        sql += " AND discourse_features.feature_type = ?"
        params.append(filters["feature_type"])
    if filters.get("document_id", "").isdigit():
        sql += " AND documents.id = ?"
        params.append(int(filters["document_id"]))
    sql += " ORDER BY datetime(discourse_features.created_at) DESC, discourse_features.id DESC"
    return query_all(sql, tuple(params))


def get_discourse_feature_counts(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT discourse_features.feature_type, COUNT(*) AS count
        FROM discourse_features
        JOIN segments ON segments.id = discourse_features.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ?
        GROUP BY discourse_features.feature_type
        ORDER BY count DESC, discourse_features.feature_type
        """,
        (project_id,),
    )


def get_voice_silence_report(project_id: int) -> list[dict]:
    rows = query_all(
        """
        SELECT actors.id, actors.name, actors.actor_type,
               segment_actors.relation_type,
               COUNT(segment_actors.id) AS relation_count,
               COUNT(DISTINCT segment_actors.segment_id) AS segment_count,
               COUNT(DISTINCT documents.id) AS document_count
        FROM actors
        LEFT JOIN segment_actors ON segment_actors.actor_id = actors.id
        LEFT JOIN segments ON segments.id = segment_actors.segment_id
        LEFT JOIN documents ON documents.id = segments.document_id
        WHERE actors.project_id = ?
        GROUP BY actors.id, segment_actors.relation_type
        ORDER BY actors.name COLLATE NOCASE
        """,
        (project_id,),
    )
    actors: dict[int, dict] = {}
    for index, row in enumerate(rows):
        actor = actors.setdefault(
            row["id"],
            {
                "id": row["id"],
                "name": row["name"],
                "actor_type": row["actor_type"],
                "total_annotations": 0,
                "document_count": 0,
                "segment_count": 0,
                **{relation: 0 for relation in ACTOR_RELATION_TYPES},
            },
        )
        if row["relation_type"] in ACTOR_RELATION_TYPES:
            actor[row["relation_type"]] = row["relation_count"]
            actor["total_annotations"] += row["relation_count"]
            actor["document_count"] += row["document_count"]
            actor["segment_count"] += row["segment_count"]

    for actor in actors.values():
        totals = query_one(
            """
            SELECT COUNT(DISTINCT documents.id) AS document_count,
                   COUNT(DISTINCT segments.id) AS segment_count
            FROM actors
            LEFT JOIN segment_actors ON segment_actors.actor_id = actors.id
            LEFT JOIN segments ON segments.id = segment_actors.segment_id
            LEFT JOIN documents ON documents.id = segments.document_id
            WHERE actors.id = ? AND actors.project_id = ?
            """,
            (actor["id"], project_id),
        )
        actor["document_count"] = totals["document_count"]
        actor["segment_count"] = totals["segment_count"]
    return list(actors.values())


def get_segment_for_project(segment_id: int, project_id: int) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT
            segments.id,
            segments.document_id,
            COALESCE(segments.name, '') AS name,
            segments.selected_text,
            documents.title AS document_title
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        WHERE segments.id = ? AND documents.project_id = ?
        """,
        (segment_id, project_id),
    )


def get_relation_entity_options(project_id: int) -> list[dict]:
    entities = []
    for entity_type in RELATION_ENTITY_TYPES:
        entities.extend(get_entities_for_type(entity_type, project_id))
    return sorted(entities, key=lambda entity: (entity["type_label"], entity["label"].lower()))


def get_entities_for_type(entity_type: str, project_id: int) -> list[dict]:
    if entity_type == "document":
        rows = query_all(
            "SELECT id, title, note FROM documents WHERE project_id = ? ORDER BY title COLLATE NOCASE",
            (project_id,),
        )
        return [
            make_entity_option("document", row["id"], f"Document: {row['title']}", row["note"])
            for row in rows
        ]
    if entity_type == "segment":
        rows = query_all(
            """
            SELECT segments.id, COALESCE(segments.name, '') AS name,
                   segments.selected_text, segments.note, documents.title AS document_title
            FROM segments
            JOIN documents ON documents.id = segments.document_id
            WHERE documents.project_id = ?
            ORDER BY documents.title COLLATE NOCASE, segments.start_offset
            """,
            (project_id,),
        )
        return [
            make_entity_option(
                "segment",
                row["id"],
                f"Segment: {row['name'] or truncate_text(row['selected_text'], 70)} — {row['document_title']}",
                row["note"],
            )
            for row in rows
        ]
    if entity_type == "code":
        rows = query_all(
            "SELECT id, name, code_type, description FROM codes WHERE project_id = ? ORDER BY code_type, name COLLATE NOCASE",
            (project_id,),
        )
        return [
            make_entity_option(
                "code",
                row["id"],
                f"Code: {row['name']} ({row['code_type']})",
                row["description"],
            )
            for row in rows
        ]
    if entity_type == "memo":
        rows = query_all(
            "SELECT id, title, memo_type, body FROM memos WHERE project_id = ? ORDER BY title COLLATE NOCASE",
            (project_id,),
        )
        return [
            make_entity_option(
                "memo", row["id"], f"Memo: {row['title']} ({row['memo_type']})", row["body"]
            )
            for row in rows
        ]
    if entity_type == "research_question":
        rows = query_all(
            "SELECT id, question, note FROM research_questions WHERE project_id = ? ORDER BY id",
            (project_id,),
        )
        return [
            make_entity_option(
                "research_question",
                row["id"],
                f"RQ: {truncate_text(row['question'], 90)}",
                row["note"],
            )
            for row in rows
        ]
    if entity_type == "discourse_marker":
        rows = query_all(
            "SELECT id, name, marker_type, description FROM discourse_markers WHERE project_id = ? ORDER BY name COLLATE NOCASE",
            (project_id,),
        )
        return [
            make_entity_option(
                "discourse_marker",
                row["id"],
                f"CDA marker: {row['name']} ({row['marker_type']})",
                row["description"],
            )
            for row in rows
        ]
    if entity_type == "actor":
        rows = query_all(
            "SELECT id, name, actor_type, description FROM actors WHERE project_id = ? ORDER BY name COLLATE NOCASE",
            (project_id,),
        )
        return [
            make_entity_option(
                "actor", row["id"], f"Actor: {row['name']} ({row['actor_type']})", row["description"]
            )
            for row in rows
        ]
    if entity_type == "discourse_feature":
        rows = query_all(
            """
            SELECT discourse_features.id, discourse_features.feature_type,
                   discourse_features.value, discourse_features.interpretation,
                   documents.title AS document_title
            FROM discourse_features
            JOIN segments ON segments.id = discourse_features.segment_id
            JOIN documents ON documents.id = segments.document_id
            WHERE documents.project_id = ?
            ORDER BY discourse_features.feature_type, discourse_features.value COLLATE NOCASE
            """,
            (project_id,),
        )
        return [
            make_entity_option(
                "discourse_feature",
                row["id"],
                f"Feature: {row['feature_type']} — {truncate_text(row['value'], 70)}",
                row["interpretation"],
            )
            for row in rows
        ]
    return []


def make_entity_option(entity_type: str, entity_id: int, label: str, meta: str | None = None) -> dict:
    return {
        "type": entity_type,
        "id": entity_id,
        "value": f"{entity_type}:{entity_id}",
        "label": label,
        "type_label": RELATION_ENTITY_TYPES.get(entity_type) or METHODOLOGY_LINKED_ENTITY_TYPES.get(entity_type, entity_type),
        "meta": meta or "",
    }


def get_entity_reference(entity_type: str, entity_id: int, project_id: int) -> dict | None:
    if entity_type not in RELATION_ENTITY_TYPES:
        return None
    for entity in get_entities_for_type(entity_type, project_id):
        if entity["id"] == entity_id:
            return entity
    return None


def parse_entity_value(raw_value: str) -> tuple[str | None, int | None]:
    try:
        entity_type, entity_id_raw = raw_value.split(":", 1)
    except ValueError:
        return None, None
    if entity_type not in RELATION_ENTITY_TYPES or not entity_id_raw.isdigit():
        return None, None
    return entity_type, int(entity_id_raw)


def validate_relation_form(project_id: int) -> tuple[dict, str | None]:
    source_type, source_id = parse_entity_value(request.form.get("source_entity", "").strip())
    target_type, target_id = parse_entity_value(request.form.get("target_entity", "").strip())
    if source_type is None or source_id is None:
        return {}, "Invalid source entity."
    if target_type is None or target_id is None:
        return {}, "Invalid target entity."
    if source_type == target_type and source_id == target_id:
        return {}, "A relation cannot link an entity to itself."
    relation_type = request.form.get("relation_type", "").strip()
    if relation_type not in RELATION_TYPES:
        return {}, "Invalid relation type."
    strength = request.form.get("strength", "moderate").strip()
    if strength not in RELATION_STRENGTHS:
        return {}, "Invalid strength."
    source = get_entity_reference(source_type, source_id, project_id)
    if source is None:
        return {}, "Invalid source entity."
    target = get_entity_reference(target_type, target_id, project_id)
    if target is None:
        return {}, "Invalid target entity."
    return {
        "title": request.form.get("title", "").strip(),
        "source_type": source_type,
        "source_id": source_id,
        "source_label": source["label"],
        "target_type": target_type,
        "target_id": target_id,
        "target_label": target["label"],
        "relation_type": relation_type,
        "strength": strength,
        "memo": request.form.get("memo", "").strip(),
        "evidence_note": request.form.get("evidence_note", "").strip(),
    }, None


def hydrate_relation(row: sqlite3.Row, project_id: int) -> dict:
    relation = row_to_dict(row)
    source = get_entity_reference(row["source_type"], row["source_id"], project_id)
    target = get_entity_reference(row["target_type"], row["target_id"], project_id)
    relation["source_label"] = source["label"] if source else f"{row['source_type']} #{row['source_id']}"
    relation["target_label"] = target["label"] if target else f"{row['target_type']} #{row['target_id']}"
    relation["relation_label"] = RELATION_TYPES.get(row["relation_type"], row["relation_type"])
    relation["strength_label"] = RELATION_STRENGTHS.get(row["strength"], row["strength"])
    relation["source_value"] = f"{row['source_type']}:{row['source_id']}"
    relation["target_value"] = f"{row['target_type']}:{row['target_id']}"
    return relation


def get_relations_for_project(project_id: int, filters: dict, limit: int | None = None) -> list[dict]:
    sql = "SELECT * FROM relations WHERE project_id = ?"
    params: list = [project_id]
    if filters.get("relation_type") in RELATION_TYPES:
        sql += " AND relation_type = ?"
        params.append(filters["relation_type"])
    if filters.get("strength") in RELATION_STRENGTHS:
        sql += " AND strength = ?"
        params.append(filters["strength"])
    if filters.get("source_type") in RELATION_ENTITY_TYPES:
        sql += " AND source_type = ?"
        params.append(filters["source_type"])
    if filters.get("target_type") in RELATION_ENTITY_TYPES:
        sql += " AND target_type = ?"
        params.append(filters["target_type"])
    if filters.get("involves_type") in RELATION_ENTITY_TYPES:
        sql += " AND (source_type = ? OR target_type = ?)"
        params.extend([filters["involves_type"], filters["involves_type"]])
    q = filters.get("q", "").strip()
    sql += " ORDER BY datetime(created_at) DESC, id DESC"
    if limit is not None:
        sql += " LIMIT ?"
        params.append(limit)
    relations = [hydrate_relation(row, project_id) for row in query_all(sql, tuple(params))]
    if q:
        q_lower = q.lower()
        relations = [
            relation
            for relation in relations
            if q_lower in relation["source_label"].lower()
            or q_lower in relation["target_label"].lower()
            or q_lower in (relation["title"] or "").lower()
            or q_lower in (relation["memo"] or "").lower()
            or q_lower in (relation["evidence_note"] or "").lower()
        ]
    return relations


def get_visual_export_filters() -> dict:
    model_mode = request.args.get("model_mode", "simplified").strip()
    if model_mode not in VISUAL_MODEL_MODES:
        model_mode = "simplified"
    full_mode = model_mode == "full"
    include_uncertain = request.args.get("include_uncertain", "1" if full_mode else "0").strip()
    include_weak = request.args.get("include_weak", "1" if full_mode else "0").strip()
    max_default = 100 if full_mode else 25
    try:
        max_relations = int(request.args.get("max_relations", str(max_default)).strip())
    except ValueError:
        max_relations = max_default
    max_relations = max(1, min(max_relations, 250))
    return {
        "model_mode": model_mode,
        "relation_type": request.args.get("relation_type", "").strip(),
        "strength": request.args.get("strength", "").strip(),
        "involves_type": request.args.get("entity_type", "").strip(),
        "include_uncertain": include_uncertain == "1",
        "include_weak": include_weak == "1",
        "max_relations": max_relations,
    }


def get_visual_model_relations(project_id: int, filters: dict | None = None) -> list[dict]:
    filters = filters or {}
    model_mode = filters.get("model_mode", "simplified")
    relations = get_relations_for_project(
        project_id,
        {
            "relation_type": filters.get("relation_type", ""),
            "strength": filters.get("strength", ""),
            "involves_type": filters.get("involves_type", ""),
        },
    )
    if model_mode != "full" and not filters.get("include_weak") and filters.get("strength") != "weak":
        relations = [relation for relation in relations if relation["strength"] != "weak"]
    if model_mode != "full" and not filters.get("include_uncertain") and filters.get("strength") != "uncertain":
        relations = [relation for relation in relations if relation["strength"] != "uncertain"]
    if model_mode != "full":
        relations = [relation for relation in relations if relation_matches_visual_mode(relation, model_mode)]
    relations = sorted(relations, key=visual_relation_sort_key)
    return relations[: filters.get("max_relations", 25)]


def code_type_from_label(label: str) -> str:
    match = re.search(r"\((open|axial|category)\)\s*$", label or "")
    return match.group(1) if match else ""


def relation_code_types(relation: dict) -> set[str]:
    code_types = set()
    for side in ("source", "target"):
        if relation[f"{side}_type"] == "code":
            code_type = code_type_from_label(relation[f"{side}_label"])
            if code_type:
                code_types.add(code_type)
    return code_types


def relation_has_entity_type(relation: dict, entity_types: set[str]) -> bool:
    return relation["source_type"] in entity_types or relation["target_type"] in entity_types


def relation_matches_visual_mode(relation: dict, model_mode: str) -> bool:
    if model_mode == "simplified":
        return True
    if model_mode == "argument":
        return (
            relation_has_entity_type(relation, {"research_question", "memo"})
            or bool(relation_code_types(relation) & {"axial", "category"})
            or relation["strength"] == "strong"
            or relation["relation_type"] in {
                "is_evidence_for",
                "supports",
                "explains",
                "leads_to",
                "legitimizes",
                "naturalizes",
                "reproduces_power_relation",
                "challenges_power_relation",
            }
        )
    if model_mode == "evidence":
        return (
            relation_has_entity_type(relation, {"segment", "document", "memo"})
            or relation["relation_type"] in {
                "is_evidence_for",
                "is_example_of",
                "is_negative_case_for",
                "supports",
                "contradicts",
            }
        )
    if model_mode == "gt":
        return (
            relation["source_type"] == "code"
            and relation["target_type"] == "code"
            and relation["relation_type"] in {
                "open_code_supports_axial_code",
                "axial_code_supports_category",
                "category_integrates",
                "property_of",
                "dimension_of",
                "condition_for",
                "consequence_of",
            }
        )
    if model_mode == "cda":
        return (
            relation_has_entity_type(relation, {"actor", "discourse_marker", "discourse_feature"})
            or relation["relation_type"] in {
                "frames",
                "legitimizes",
                "delegitimizes",
                "naturalizes",
                "silences",
                "foregrounds",
                "backgrounds",
                "individualizes",
                "aggregates",
                "constructs_actor_as",
                "reproduces_power_relation",
                "challenges_power_relation",
                "presupposes",
                "metaphorizes",
            }
        )
    return True


def visual_relation_sort_key(relation: dict) -> tuple:
    strength_rank = {"strong": 0, "moderate": 1, "weak": 2, "uncertain": 3}
    family_rank = {
        "support": 0,
        "gt": 1,
        "cda_power": 2,
        "cda_representation": 3,
        "causal_process": 4,
        "contrast": 5,
        "generic": 6,
    }
    return (
        strength_rank.get(relation["strength"], 9),
        family_rank.get(relation_family(relation["relation_type"]), 9),
        relation["relation_type"],
        relation["id"],
    )


def get_relation_for_project(relation_id: int, project_id: int) -> dict | None:
    row = query_one("SELECT * FROM relations WHERE id = ? AND project_id = ?", (relation_id, project_id))
    if row is None:
        return None
    return hydrate_relation(row, project_id)


def get_entity_relations(project_id: int, entity_type: str, entity_id: int, direction: str) -> list[dict]:
    if direction == "outgoing":
        filters = {"source_type": entity_type}
        relations = get_relations_for_project(project_id, filters)
        return [relation for relation in relations if relation["source_id"] == entity_id]
    filters = {"target_type": entity_type}
    relations = get_relations_for_project(project_id, filters)
    return [relation for relation in relations if relation["target_id"] == entity_id]


def get_relations_for_entity(project_id: int, entity_type: str, entity_id: int, limit: int | None = None) -> list[dict]:
    relations = get_relations_for_project(project_id, {}, limit=None)
    result = [
        relation for relation in relations
        if (relation["source_type"] == entity_type and relation["source_id"] == entity_id)
        or (relation["target_type"] == entity_type and relation["target_id"] == entity_id)
    ]
    return result[:limit] if limit else result


def get_relation_count_for_entity(project_id: int, entity_type: str, entity_id: int) -> int:
    return query_one(
        """
        SELECT COUNT(*) AS count
        FROM relations
        WHERE project_id = ?
          AND ((source_type = ? AND source_id = ?) OR (target_type = ? AND target_id = ?))
        """,
        (project_id, entity_type, entity_id, entity_type, entity_id),
    )["count"]


def get_related_memos_for_model_entity(project_id: int, entity_type: str, entity_id: int) -> list[sqlite3.Row]:
    if entity_type not in {"document", "segment", "code"}:
        return []
    return get_memos_for_entity(project_id, entity_type, entity_id)


def get_research_questions_for_project(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT id, question, note, created_at, updated_at
        FROM research_questions
        WHERE project_id = ?
        ORDER BY datetime(created_at) DESC, id DESC
        """,
        (project_id,),
    )


def get_methodology_overview(active_project: sqlite3.Row) -> str:
    mode = active_project["methodology_mode"]
    descriptions = {
        "generic": "Generic qualitative coding supports open coding, memo-writing, comparison, synthesis, and transparent reporting.",
        "gt": "Grounded Theory mode supports iterative coding, constant comparison, axial coding, categories, and theoretical integration.",
        "cda": "CDA mode supports text analysis, discursive practice, social practice, actor voice, ideology, and power analysis.",
        "mixed": "Mixed GT + CDA mode enables both workflows and requires explicit protocol decisions about how methods are combined.",
    }
    return descriptions.get(mode, descriptions["generic"])


def get_methodology_entity_options(project_id: int) -> list[dict]:
    options = [make_entity_option("project", project_id, "Project: Active project")]
    for entity_type in RELATION_ENTITY_TYPES:
        options.extend(get_entities_for_type(entity_type, project_id))
    for relation in get_relations_for_project(project_id, {}):
        label = f"Relation: {relation['source_label']} — {relation['relation_label']} → {relation['target_label']}"
        options.append(make_entity_option("relation", relation["id"], truncate_text(label, 120)))
    return sorted(options, key=lambda entity: (entity["type_label"], entity["label"].lower()))


def get_methodology_entity_reference(entity_type: str | None, entity_id: int | None, project_id: int) -> dict | None:
    if not entity_type or entity_id is None:
        return None
    if entity_type == "project":
        project = get_project(entity_id)
        if project and project["id"] == project_id:
            return make_entity_option("project", project_id, f"Project: {project['name']}")
        return None
    if entity_type == "relation":
        relation = get_relation_for_project(entity_id, project_id)
        if relation:
            label = f"Relation: {relation['source_label']} — {relation['relation_label']} → {relation['target_label']}"
            return make_entity_option("relation", relation["id"], truncate_text(label, 120))
        return None
    return get_entity_reference(entity_type, entity_id, project_id)


def hydrate_methodology_note(row: sqlite3.Row, project_id: int) -> dict:
    note = row_to_dict(row)
    reference = get_methodology_entity_reference(note["linked_entity_type"], note["linked_entity_id"], project_id)
    note["linked_entity_label"] = reference["label"] if reference else "Project protocol"
    note["note_type_label"] = METHODOLOGY_NOTE_TYPES.get(note["note_type"], note["note_type"])
    note["methodology_area_label"] = METHODOLOGY_AREAS.get(note["methodology_area"], note["methodology_area"])
    note["status_label"] = METHODOLOGY_NOTE_STATUSES.get(note["status"], note["status"])
    note["entity_value"] = (
        f"{note['linked_entity_type']}:{note['linked_entity_id']}"
        if note["linked_entity_type"] and note["linked_entity_id"]
        else ""
    )
    return note


def get_methodology_notes_for_project(project_id: int, filters: dict | None = None, limit: int | None = None) -> list[dict]:
    filters = filters or {}
    sql = "SELECT * FROM methodology_notes WHERE project_id = ?"
    params: list = [project_id]
    if filters.get("note_type") in METHODOLOGY_NOTE_TYPES:
        sql += " AND note_type = ?"
        params.append(filters["note_type"])
    if filters.get("methodology_area") in METHODOLOGY_AREAS:
        sql += " AND methodology_area = ?"
        params.append(filters["methodology_area"])
    if filters.get("status") in METHODOLOGY_NOTE_STATUSES:
        sql += " AND status = ?"
        params.append(filters["status"])
    if filters.get("linked_entity_type") in METHODOLOGY_LINKED_ENTITY_TYPES:
        sql += " AND linked_entity_type = ?"
        params.append(filters["linked_entity_type"])
    q = filters.get("q", "").strip().lower()
    sql += " ORDER BY datetime(updated_at) DESC, datetime(created_at) DESC, id DESC"
    if limit is not None:
        sql += " LIMIT ?"
        params.append(limit)
    notes = [hydrate_methodology_note(row, project_id) for row in query_all(sql, tuple(params))]
    if q:
        notes = [note for note in notes if q in note["title"].lower() or q in note["body"].lower()]
    return notes


def get_methodology_note_for_project(note_id: int, project_id: int) -> dict | None:
    row = query_one("SELECT * FROM methodology_notes WHERE id = ? AND project_id = ?", (note_id, project_id))
    if row is None:
        return None
    return hydrate_methodology_note(row, project_id)


def get_methodology_note_counts(project_id: int) -> dict[str, int]:
    return {
        "total": query_one("SELECT COUNT(*) AS count FROM methodology_notes WHERE project_id = ?", (project_id,))["count"],
        "active_protocol": query_one(
            "SELECT COUNT(*) AS count FROM methodology_notes WHERE project_id = ? AND note_type = 'protocol' AND status = 'active'",
            (project_id,),
        )["count"],
        "coding_rules": query_one(
            "SELECT COUNT(*) AS count FROM methodology_notes WHERE project_id = ? AND note_type = 'coding_rule'",
            (project_id,),
        )["count"],
        "sampling_rules": query_one(
            "SELECT COUNT(*) AS count FROM methodology_notes WHERE project_id = ? AND note_type = 'sampling_rule'",
            (project_id,),
        )["count"],
        "needs_review": query_one(
            "SELECT COUNT(*) AS count FROM methodology_notes WHERE project_id = ? AND status = 'needs_review'",
            (project_id,),
        )["count"],
    }


def validate_methodology_note_form(project_id: int) -> tuple[dict, str | None]:
    title = request.form.get("title", "").strip()
    body = request.form.get("body", "").strip()
    note_type = request.form.get("note_type", "protocol").strip()
    methodology_area = request.form.get("methodology_area", "generic").strip()
    status = request.form.get("status", "active").strip()
    raw_entity = request.form.get("linked_entity", "").strip()
    linked_entity_type = None
    linked_entity_id = None
    if not title:
        return {}, "Title is required."
    if not body:
        return {}, "Body is required."
    if note_type not in METHODOLOGY_NOTE_TYPES:
        return {}, "Invalid methodology note type."
    if methodology_area not in METHODOLOGY_AREAS:
        return {}, "Invalid methodology area."
    if status not in METHODOLOGY_NOTE_STATUSES:
        return {}, "Invalid methodology note status."
    if raw_entity:
        linked_entity_type, linked_entity_id = parse_entity_value(raw_entity)
        if linked_entity_type is None and raw_entity.startswith("project:"):
            raw_type, raw_id = raw_entity.split(":", 1)
            linked_entity_type = raw_type if raw_id.isdigit() else None
            linked_entity_id = int(raw_id) if raw_id.isdigit() else None
        if linked_entity_type is None and raw_entity.startswith("relation:"):
            raw_type, raw_id = raw_entity.split(":", 1)
            linked_entity_type = raw_type if raw_id.isdigit() else None
            linked_entity_id = int(raw_id) if raw_id.isdigit() else None
        if linked_entity_type not in METHODOLOGY_LINKED_ENTITY_TYPES or linked_entity_id is None:
            return {}, "Invalid linked entity."
        if get_methodology_entity_reference(linked_entity_type, linked_entity_id, project_id) is None:
            return {}, "Linked entity does not belong to the active project."
    return {
        "title": title,
        "body": body,
        "note_type": note_type,
        "methodology_area": methodology_area,
        "status": status,
        "linked_entity_type": linked_entity_type,
        "linked_entity_id": linked_entity_id,
    }, None


def get_model_mode_prompts(active_project: sqlite3.Row) -> list[str]:
    generic = [
        "What does this relation explain?",
        "What is the evidence for this relation?",
        "Is the relation strong, weak, or uncertain?",
        "Is there a negative case?",
        "Does the relation help answer a research question?",
    ]
    gt = [
        "Is this a condition, action/interaction, consequence, property, or dimension?",
        "Does this open code support an axial code?",
        "Does this axial code support a category?",
        "Does this relation help integrate the emerging theory?",
    ]
    cda = [
        "Who benefits from this relation?",
        "Does this relation legitimize or naturalize a power structure?",
        "Does it foreground or background an actor?",
        "Does it show how discourse constructs an actor or social problem?",
        "Does it reveal voice, silence, access, or dominance?",
    ]
    if active_project["methodology_mode"] == "gt":
        return generic + gt
    if active_project["methodology_mode"] == "cda":
        return generic + cda
    if active_project["methodology_mode"] == "mixed":
        return generic + gt + cda
    return generic


def get_open_codes_for_project(project_id: int) -> list[sqlite3.Row]:
    return get_codes_for_project(project_id, "open")


def get_codes_for_project(project_id: int, code_type_filter: str = "all") -> list[sqlite3.Row]:
    where = "codes.project_id = ?"
    params = [project_id]
    if code_type_filter in {"open", "axial", "category"}:
        where += " AND codes.code_type = ?"
        params.append(code_type_filter)
    return query_all(
        f"""
        SELECT
            codes.id,
            codes.name,
            codes.description,
            codes.code_type,
            codes.color,
            codes.parent_id,
            parent.name AS parent_name,
            codes.definition,
            codes.include_when,
            codes.exclude_when,
            codes.example,
            codes.analytical_note,
            codes.created_at,
            codes.updated_at,
            COUNT(segment_codes.segment_id) AS assigned_segment_count,
            CASE
                WHEN codes.definition IS NULL OR TRIM(codes.definition) = '' THEN 'Missing definition'
                WHEN TRIM(COALESCE(codes.include_when, '')) != ''
                 AND TRIM(COALESCE(codes.exclude_when, '')) != ''
                 AND TRIM(COALESCE(codes.example, '')) != '' THEN 'Detailed entry'
                ELSE 'Basic entry'
            END AS completeness
        FROM codes
        LEFT JOIN segment_codes ON segment_codes.code_id = codes.id
        LEFT JOIN codes AS parent ON parent.id = codes.parent_id
        WHERE {where}
        GROUP BY codes.id
        ORDER BY codes.code_type, codes.name COLLATE NOCASE
        """,
        tuple(params),
    )


def get_code_for_project(code_id: int, project_id: int) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT id, project_id, name, description, code_type, color, parent_id,
               definition, include_when, exclude_when, example, analytical_note,
               gt_conditions, gt_context, gt_actions_interactions, gt_consequences,
               gt_properties, gt_dimensions, gt_theoretical_note,
               created_at, updated_at
        FROM codes
        WHERE id = ? AND project_id = ?
        """,
        (code_id, project_id),
    )


def segment_has_code(segment_id: int, code_id: int) -> bool:
    return (
        query_one(
            """
            SELECT 1
            FROM segment_codes
            WHERE segment_id = ? AND code_id = ?
            """,
            (segment_id, code_id),
        )
        is not None
    )


def create_gt_code(code_type: str, action: str, detail_prefix: str):
    active_project = get_active_project()
    name = request.form.get("name", "").strip()
    description = request.form.get("description", "").strip()
    color = normalize_code_color(request.form.get("color", ""))
    if not name:
        flash("Code name is required.", "error")
        return redirect(url_for("gt_workspace"))
    code_id = execute_write(
        """
        INSERT INTO codes (project_id, name, description, code_type, color)
        VALUES (?, ?, ?, ?, ?)
        """,
        (active_project["id"], name, description, code_type, color),
    )
    log_action(active_project["id"], "code", code_id, action, f"{detail_prefix}: {name}")
    flash(f"{detail_prefix}: {name}", "success")
    return redirect(url_for("gt_workspace"))


def assign_code_parent(
    child_id: int,
    child_type: str,
    parent_type: str,
    form_field: str,
    action: str,
    success_template: str,
):
    active_project = get_active_project()
    child = get_code_for_project(child_id, active_project["id"])
    parent_id_raw = request.form.get(form_field, "").strip()
    if child is None or child["code_type"] != child_type or not parent_id_raw.isdigit():
        flash("Invalid GT hierarchy operation.", "error")
        return redirect(url_for("gt_workspace"))
    parent_id = int(parent_id_raw)
    valid, message = validate_parent_assignment(
        child_id, child_type, parent_id, active_project["id"]
    )
    if not valid:
        flash(message, "error")
        return redirect(url_for("gt_workspace"))
    parent = get_code_for_project(parent_id, active_project["id"])
    execute_write(
        "UPDATE codes SET parent_id = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (parent_id, child_id),
    )
    details = success_template.format(child=child["name"], parent=parent["name"])
    log_action(active_project["id"], "code", child_id, action, details)
    flash(details, "success")
    return redirect(url_for("gt_workspace"))


def unassign_code_parent(code_id: int, expected_type: str, action: str):
    active_project = get_active_project()
    code = get_code_for_project(code_id, active_project["id"])
    if code is None or code["code_type"] != expected_type:
        flash("Invalid GT hierarchy operation.", "error")
        return redirect(url_for("gt_workspace"))
    execute_write(
        "UPDATE codes SET parent_id = NULL, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (code_id,),
    )
    log_action(
        active_project["id"],
        "code",
        code_id,
        action,
        f"Unassigned {expected_type} code: {code['name']}",
    )
    flash(f"Unassigned code: {code['name']}", "success")
    return redirect(url_for("gt_workspace"))


def validate_parent_assignment(
    code_id: int, child_type: str, parent_id: int, project_id: int
) -> tuple[bool, str]:
    if code_id == parent_id:
        return False, "Invalid GT hierarchy operation."
    parent = get_code_for_project(parent_id, project_id)
    expected_parent = {"open": "axial", "axial": "category"}.get(child_type)
    if expected_parent is None or parent is None or parent["code_type"] != expected_parent:
        return False, "Invalid GT hierarchy operation."
    return True, ""


def get_parent_code(parent_id: int, project_id: int) -> sqlite3.Row | None:
    return get_code_for_project(parent_id, project_id)


def get_child_codes(parent_id: int, project_id: int, code_type: str) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT id, name, description, color, code_type
        FROM codes
        WHERE parent_id = ? AND project_id = ? AND code_type = ?
        ORDER BY name COLLATE NOCASE
        """,
        (parent_id, project_id, code_type),
    )


def get_open_codes_under_category(category_id: int, project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT open_code.id, open_code.name, open_code.description, open_code.color,
               axial_code.name AS axial_name
        FROM codes AS axial_code
        JOIN codes AS open_code ON open_code.parent_id = axial_code.id
        WHERE axial_code.parent_id = ? AND axial_code.project_id = ?
          AND open_code.project_id = ? AND axial_code.code_type = 'axial'
          AND open_code.code_type = 'open'
        ORDER BY axial_code.name COLLATE NOCASE, open_code.name COLLATE NOCASE
        """,
        (category_id, project_id, project_id),
    )


def get_gt_open_codes(project_id: int) -> list[dict]:
    rows = query_all(
        """
        SELECT codes.*, parent.name AS parent_name,
               COUNT(segment_codes.segment_id) AS assigned_segment_count
        FROM codes
        LEFT JOIN codes AS parent ON parent.id = codes.parent_id
        LEFT JOIN segment_codes ON segment_codes.code_id = codes.id
        WHERE codes.project_id = ? AND codes.code_type = 'open'
        GROUP BY codes.id
        ORDER BY codes.name COLLATE NOCASE
        """,
        (project_id,),
    )
    return [dict(row) for row in rows]


def get_gt_axial_codes(project_id: int) -> list[dict]:
    rows = query_all(
        """
        SELECT codes.*, parent.name AS parent_name
        FROM codes
        LEFT JOIN codes AS parent ON parent.id = codes.parent_id
        WHERE codes.project_id = ? AND codes.code_type = 'axial'
        ORDER BY codes.name COLLATE NOCASE
        """,
        (project_id,),
    )
    result = []
    for index, row in enumerate(rows):
        code = dict(row)
        code["child_open_count"] = len(get_child_codes(row["id"], project_id, "open"))
        code["hierarchy_segment_count"] = len(get_hierarchy_segments_for_code(row, project_id))
        result.append(code)
    return result


def get_gt_categories(project_id: int) -> list[dict]:
    rows = query_all(
        """
        SELECT *
        FROM codes
        WHERE project_id = ? AND code_type = 'category'
        ORDER BY name COLLATE NOCASE
        """,
        (project_id,),
    )
    result = []
    for row in rows:
        code = dict(row)
        axial_children = get_child_codes(row["id"], project_id, "axial")
        code["child_axial_count"] = len(axial_children)
        code["child_open_count"] = sum(
            len(get_child_codes(axial["id"], project_id, "open"))
            for axial in axial_children
        )
        code["hierarchy_segment_count"] = len(get_hierarchy_segments_for_code(row, project_id))
        result.append(code)
    return result


def get_code_usage_count(code_id: int) -> int:
    return query_one(
        "SELECT COUNT(*) AS count FROM segment_codes WHERE code_id = ?",
        (code_id,),
    )["count"]


def get_segments_for_code(code_id: int, project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT
            segments.id,
            COALESCE(segments.name, '') AS name,
            segments.selected_text,
            segments.note,
            documents.id AS document_id,
            documents.title AS document_title
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        JOIN segment_codes ON segment_codes.segment_id = segments.id
        WHERE segment_codes.code_id = ? AND documents.project_id = ?
        ORDER BY documents.title COLLATE NOCASE, segments.start_offset
        """,
        (code_id, project_id),
    )


def get_hierarchy_segments_for_code(code: sqlite3.Row | dict, project_id: int) -> list[sqlite3.Row]:
    if code["code_type"] == "open":
        return get_segments_for_code(code["id"], project_id)
    if code["code_type"] == "axial":
        return query_all(
            """
            SELECT DISTINCT segments.id, COALESCE(segments.name, '') AS name,
                   segments.selected_text, segments.note, documents.id AS document_id,
                   documents.title AS document_title
            FROM codes AS open_code
            JOIN segment_codes ON segment_codes.code_id = open_code.id
            JOIN segments ON segments.id = segment_codes.segment_id
            JOIN documents ON documents.id = segments.document_id
            WHERE open_code.parent_id = ? AND documents.project_id = ?
            ORDER BY documents.title COLLATE NOCASE, segments.start_offset
            """,
            (code["id"], project_id),
        )
    if code["code_type"] == "category":
        return query_all(
            """
            SELECT DISTINCT segments.id, COALESCE(segments.name, '') AS name,
                   segments.selected_text, segments.note, documents.id AS document_id,
                   documents.title AS document_title
            FROM codes AS axial_code
            JOIN codes AS open_code ON open_code.parent_id = axial_code.id
            JOIN segment_codes ON segment_codes.code_id = open_code.id
            JOIN segments ON segments.id = segment_codes.segment_id
            JOIN documents ON documents.id = segments.document_id
            WHERE axial_code.parent_id = ? AND documents.project_id = ?
            ORDER BY documents.title COLLATE NOCASE, segments.start_offset
            """,
            (code["id"], project_id),
        )
    return []


def get_compare_code(code_id_raw: str | None, project_id: int) -> dict | None:
    if not code_id_raw or not code_id_raw.isdigit():
        return None
    code = get_code_for_project(int(code_id_raw), project_id)
    if code is None or code["code_type"] != "open":
        return None
    result = dict(code)
    result["segments"] = get_segments_for_code(code["id"], project_id)
    result["usage_count"] = len(result["segments"])
    return result


def get_gt_structure_preview(project_id: int) -> dict[str, int]:
    return {
        "open_without_axial": query_one(
            """
            SELECT COUNT(*) AS count FROM codes
            WHERE project_id = ? AND code_type = 'open' AND parent_id IS NULL
            """,
            (project_id,),
        )["count"],
        "axial_without_category": query_one(
            """
            SELECT COUNT(*) AS count FROM codes
            WHERE project_id = ? AND code_type = 'axial' AND parent_id IS NULL
            """,
            (project_id,),
        )["count"],
        "categories": query_one(
            """
            SELECT COUNT(*) AS count FROM codes
            WHERE project_id = ? AND code_type = 'category'
            """,
            (project_id,),
        )["count"],
    }


def get_documents_for_memo_links(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT id, title
        FROM documents
        WHERE project_id = ?
        ORDER BY title COLLATE NOCASE
        """,
        (project_id,),
    )


def get_segments_for_memo_links(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT
            segments.id,
            COALESCE(segments.name, '') AS name,
            segments.selected_text,
            documents.title AS document_title
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ?
        ORDER BY documents.title COLLATE NOCASE, segments.start_offset
        """,
        (project_id,),
    )


def get_memos_for_project(project_id: int, filters: dict) -> list[dict]:
    sql = """
        SELECT id, title, body, memo_type, linked_entity_type, linked_entity_id,
               status, created_at, updated_at
        FROM memos
        WHERE project_id = ?
    """
    params = [project_id]
    if filters.get("memo_type") in MEMO_TYPES:
        sql += " AND memo_type = ?"
        params.append(filters["memo_type"])
    if filters.get("status") in MEMO_STATUSES:
        sql += " AND status = ?"
        params.append(filters["status"])
    if filters.get("linked_entity_type") in {"project", "document", "segment", "code"}:
        sql += " AND linked_entity_type = ?"
        params.append(filters["linked_entity_type"])
    if filters.get("linked_entity_id", "").isdigit():
        sql += " AND linked_entity_id = ?"
        params.append(int(filters["linked_entity_id"]))
    sql += " ORDER BY datetime(created_at) DESC, id DESC"

    memos = []
    for row in query_all(sql, tuple(params)):
        memo = dict(row)
        memo["linked_entity_label"] = get_linked_entity_label(memo, project_id)
        memo["relation_count"] = get_relation_count_for_entity(project_id, "memo", memo["id"])
        memos.append(memo)
    return memos


def get_memo_for_project(memo_id: int, project_id: int) -> sqlite3.Row | None:
    return query_one(
        """
        SELECT id, project_id, title, body, memo_type, linked_entity_type,
               linked_entity_id, status, created_at, updated_at
        FROM memos
        WHERE id = ? AND project_id = ?
        """,
        (memo_id, project_id),
    )


def get_memos_for_entity(
    project_id: int | None, linked_entity_type: str, linked_entity_id: int
) -> list[sqlite3.Row]:
    if project_id is None:
        return query_all(
            """
            SELECT id, title, body, memo_type, status, created_at
            FROM memos
            WHERE linked_entity_type = ? AND linked_entity_id = ?
            ORDER BY datetime(created_at) DESC, id DESC
            """,
            (linked_entity_type, linked_entity_id),
        )
    return query_all(
        """
        SELECT id, title, body, memo_type, status, created_at
        FROM memos
        WHERE project_id = ? AND linked_entity_type = ? AND linked_entity_id = ?
        ORDER BY datetime(created_at) DESC, id DESC
        """,
        (project_id, linked_entity_type, linked_entity_id),
    )


def validate_memo_form(project_id: int) -> tuple[dict, str | None]:
    title = request.form.get("title", "").strip()
    body = request.form.get("body", "").strip()
    memo_type = request.form.get("memo_type", "").strip()
    status = request.form.get("status", "draft").strip()
    linked_target = request.form.get("linked_target", "").strip()
    linked_entity_type = request.form.get("linked_entity_type", "").strip() or None
    linked_entity_id_raw = request.form.get("linked_entity_id", "").strip()
    if linked_target:
        try:
            linked_entity_type, linked_entity_id_raw = linked_target.split(":", 1)
        except ValueError:
            return {}, "Invalid memo link."

    if not title or not body:
        return {}, "Memo title and body are required."
    if memo_type not in MEMO_TYPES:
        return {}, "Invalid memo type."
    if status not in MEMO_STATUSES:
        return {}, "Invalid memo status."

    linked_entity_id = None
    if linked_entity_type or linked_entity_id_raw:
        if linked_entity_type not in {"project", "document", "segment", "code"}:
            return {}, "Invalid memo link."
        if not linked_entity_id_raw.isdigit():
            return {}, "Invalid memo link."
        linked_entity_id = int(linked_entity_id_raw)
        if not linked_entity_exists(project_id, linked_entity_type, linked_entity_id):
            return {}, "Invalid memo link."

    return {
        "title": title,
        "body": body,
        "memo_type": memo_type,
        "status": status,
        "linked_entity_type": linked_entity_type,
        "linked_entity_id": linked_entity_id,
    }, None


def linked_entity_exists(project_id: int, entity_type: str, entity_id: int) -> bool:
    if entity_type == "project":
        return entity_id == project_id
    if entity_type == "document":
        return get_document_for_project(entity_id, project_id) is not None
    if entity_type == "segment":
        return get_segment_for_project(entity_id, project_id) is not None
    if entity_type == "code":
        return get_code_for_project(entity_id, project_id) is not None
    return False


def get_linked_entity_label(memo: dict, project_id: int) -> str:
    entity_type = memo.get("linked_entity_type")
    entity_id = memo.get("linked_entity_id")
    if not entity_type or entity_id is None:
        return "None"
    if entity_type == "project":
        project = get_active_project()
        return f"Project: {project['name']}"
    if entity_type == "document":
        document = get_document_for_project(entity_id, project_id)
        return f"Document: {document['title']}" if document else "Document: missing"
    if entity_type == "segment":
        segment = get_segment_for_project(entity_id, project_id)
        if segment is None:
            return "Segment: missing"
        label = segment["name"] or truncate_text(segment["selected_text"], 48)
        return f"Segment: {label}"
    if entity_type == "code":
        code = get_code_for_project(entity_id, project_id)
        return f"Code: {code['name']}" if code else "Code: missing"
    return "None"


def is_valid_segment_selection(
    document_text: str, selected_text: str, start_offset: int, end_offset: int
) -> bool:
    if not selected_text.strip():
        return False
    if start_offset < 0 or end_offset <= start_offset:
        return False
    if end_offset > len(document_text):
        return False
    return bool(document_text[start_offset:end_offset].strip())


def build_highlighted_document_text(
    document_text: str, segments: list[dict], use_cda_markers: bool = True
) -> Markup:
    pieces = []
    cursor = 0

    for segment in segments:
        start = segment["start_offset"]
        end = segment["end_offset"]
        if start < cursor or start < 0 or end > len(document_text) or end <= start:
            continue

        if segment["codes"]:
            color = segment["codes"][0]["color"]
        elif use_cda_markers and segment["discourse_markers"]:
            color = segment["discourse_markers"][0]["color"]
        else:
            color = "#fff0a8"
        style = f"--segment-color: {normalize_code_color(color)};"
        pieces.append(escape(document_text[cursor:start]))
        pieces.append(
            Markup(
                '<mark class="segment-highlight" data-segment-id="{}" style="{}">{}</mark>'
            ).format(
                segment["id"], style, escape(document_text[start:end])
            )
        )
        cursor = end

    pieces.append(escape(document_text[cursor:]))
    return Markup("").join(pieces)


def delete_document_data(document_id: int) -> None:
    db = get_db()
    segment_ids = [
        row["id"]
        for row in db.execute(
            "SELECT id FROM segments WHERE document_id = ?", (document_id,)
        ).fetchall()
    ]
    if segment_ids:
        placeholders = ",".join("?" for _ in segment_ids)
        db.execute(
            f"DELETE FROM segment_codes WHERE segment_id IN ({placeholders})",
            tuple(segment_ids),
        )
        db.execute(
            f"DELETE FROM segment_discourse_markers WHERE segment_id IN ({placeholders})",
            tuple(segment_ids),
        )
        db.execute(
            f"DELETE FROM segment_actors WHERE segment_id IN ({placeholders})",
            tuple(segment_ids),
        )
        db.execute(
            f"DELETE FROM discourse_features WHERE segment_id IN ({placeholders})",
            tuple(segment_ids),
        )
    db.execute("DELETE FROM segments WHERE document_id = ?", (document_id,))
    db.execute("DELETE FROM documents WHERE id = ?", (document_id,))
    db.commit()
    # The original uploaded source file is intentionally left in uploads/ for now.


def log_action(
    project_id: int | None,
    entity_type: str,
    entity_id: int | None,
    action: str,
    details: str,
) -> int:
    return execute_write(
        """
        INSERT INTO audit_log (project_id, entity_type, entity_id, action, details)
        VALUES (?, ?, ?, ?, ?)
        """,
        (project_id, entity_type, entity_id, action, details),
    )


def get_latest_audit_entries(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT entity_type, entity_id, action, details, created_at
        FROM audit_log
        WHERE project_id = ?
        ORDER BY datetime(created_at) DESC, id DESC
        LIMIT 5
        """,
        (project_id,),
    )


def render_placeholder(
    template_name: str,
    title: str,
    active_page: str,
    panel_title: str,
    panel_body: str,
) -> str:
    return render_template(
        template_name,
        title=title,
        active_page=active_page,
        active_project=get_active_project(),
        panel_title=panel_title,
        panel_body=panel_body,
    )


def redirect_to_document(document_id: int, scroll_y: str | None = None):
    scroll_y = scroll_y if scroll_y is not None else request.form.get("scroll_y", "")
    document_scroll_y = request.form.get("document_scroll_y", "")
    kwargs = {"document_id": document_id}
    if str(scroll_y).isdigit():
        kwargs["scroll_y"] = str(scroll_y)
    if str(document_scroll_y).isdigit():
        kwargs["document_scroll_y"] = str(document_scroll_y)
    return redirect(url_for("document_view", **kwargs))


def redirect_after_code_change(document_id: str):
    if document_id.isdigit():
        active_project = get_active_project()
        document = get_document_for_project(int(document_id), active_project["id"])
        if document is not None:
            return redirect_to_document(document["id"])
    return redirect(url_for("codes"))


def safe_next_url(next_url: str) -> str | None:
    if next_url.startswith("/") and not next_url.startswith("//"):
        return next_url
    return None


def truncate_text(text: str, length: int = 80) -> str:
    text = " ".join((text or "").split())
    if len(text) <= length:
        return text
    return text[: length - 3] + "..."


def short_label(label: str, max_length: int = 40) -> str:
    return truncate_text(label, max_length)


def graph_node_id(entity_type: str, entity_id: int) -> str:
    raw_id = f"{entity_type}_{entity_id}"
    return re.sub(r"[^A-Za-z0-9_]", "_", raw_id)


def short_relation_label(relation_type: str) -> str:
    labels = {
        "open_code_supports_axial_code": "open→axial",
        "axial_code_supports_category": "axial→category",
        "reproduces_power_relation": "reproduces power",
        "challenges_power_relation": "challenges power",
        "constructs_actor_as": "constructs as",
        "is_negative_case_for": "negative case",
        "is_evidence_for": "evidence for",
        "is_example_of": "example of",
        "contrasts_with": "contrasts",
        "transforms_into": "transforms",
        "category_integrates": "integrates",
        "condition_for": "condition for",
        "consequence_of": "consequence of",
        "property_of": "property of",
        "dimension_of": "dimension of",
    }
    return labels.get(relation_type, relation_type.replace("_", " "))


def relation_family(relation_type: str) -> str:
    if relation_type in {"supports", "is_evidence_for", "is_example_of", "elaborates", "explains"}:
        return "support"
    if relation_type in {"contradicts", "contrasts_with", "is_negative_case_for"}:
        return "contrast"
    if relation_type in {"leads_to", "causes", "conditions", "enables", "limits", "transforms_into"}:
        return "causal_process"
    if relation_type in {
        "open_code_supports_axial_code",
        "axial_code_supports_category",
        "category_integrates",
        "property_of",
        "dimension_of",
        "consequence_of",
        "condition_for",
    }:
        return "gt"
    if relation_type in {
        "legitimizes",
        "delegitimizes",
        "naturalizes",
        "silences",
        "foregrounds",
        "backgrounds",
        "reproduces_power_relation",
        "challenges_power_relation",
    }:
        return "cda_power"
    if relation_type in {
        "frames",
        "constructs_actor_as",
        "individualizes",
        "aggregates",
        "presupposes",
        "metaphorizes",
    }:
        return "cda_representation"
    return "generic"


def relation_family_styles() -> dict[str, dict[str, str]]:
    return {
        "support": {"label": "support", "color": "#2563eb", "tikz": "blue!70!black"},
        "contrast": {"label": "contrast", "color": "#b4233a", "tikz": "red!70!black"},
        "causal_process": {"label": "causal/process", "color": "#9a5a00", "tikz": "orange!80!black"},
        "gt": {"label": "GT", "color": "#6d4ba3", "tikz": "purple!75!black"},
        "cda_power": {"label": "CDA power", "color": "#c2410c", "tikz": "orange!90!black"},
        "cda_representation": {"label": "CDA representation", "color": "#0f766e", "tikz": "teal!80!black"},
        "generic": {"label": "generic", "color": "#64748b", "tikz": "gray!80!black"},
    }


def strength_style(strength: str) -> dict[str, str]:
    styles = {
        "strong": {
            "svg_width": "3.0",
            "svg_opacity": "0.95",
            "svg_dash": "",
            "dot_penwidth": "3",
            "dot_style": "solid",
            "tikz": "very thick",
            "mermaid": "==>",
        },
        "moderate": {
            "svg_width": "2.0",
            "svg_opacity": "0.8",
            "svg_dash": "",
            "dot_penwidth": "2",
            "dot_style": "solid",
            "tikz": "thick",
            "mermaid": "-->",
        },
        "weak": {
            "svg_width": "1.2",
            "svg_opacity": "0.45",
            "svg_dash": "5,4",
            "dot_penwidth": "1",
            "dot_style": "dashed",
            "tikz": "thin,dashed",
            "mermaid": "-.->",
        },
        "uncertain": {
            "svg_width": "1.2",
            "svg_opacity": "0.45",
            "svg_dash": "2,4",
            "dot_penwidth": "1",
            "dot_style": "dotted",
            "tikz": "thin,dotted",
            "mermaid": "-.->",
        },
    }
    return styles.get(strength, styles["moderate"])


def escape_mermaid_text(text: str) -> str:
    return empty(text).replace("\\", "\\\\").replace('"', '\\"')


def escape_dot_text(text: str) -> str:
    return empty(text).replace("\\", "\\\\").replace('"', '\\"')


def escape_latex_text(text: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "#": r"\#",
        "$": r"\$",
        "%": r"\%",
        "&": r"\&",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(char, char) for char in empty(text))


def escape_svg_text(text: str) -> str:
    return html.escape(empty(text), quote=True)


def get_dashboard_export_links() -> list[dict[str, str]]:
    return [
        {"label": "Codebook", "endpoint": "export_codebook_markdown"},
        {"label": "Coded segments", "endpoint": "export_coded_segments_csv"},
        {"label": "Memos", "endpoint": "export_memos_markdown"},
        {"label": "Methodology protocol", "endpoint": "export_methodology_protocol_markdown"},
        {"label": "Analytical model", "endpoint": "export_model_markdown"},
        {"label": "Project package", "endpoint": "export_project_package"},
    ]


def export_timestamp() -> str:
    return datetime.now().isoformat(timespec="seconds")


def empty(value) -> str:
    return "" if value is None else str(value)


def row_to_dict(row: sqlite3.Row) -> dict:
    return {key: row[key] for key in row.keys()}


def rows_to_dicts(rows: list[sqlite3.Row]) -> list[dict]:
    return [row_to_dict(row) for row in rows]


def load_methodology_library(library_id: str) -> dict | None:
    if not re.fullmatch(r"[a-z0-9_]+", library_id or ""):
        return None
    path = METHODOLOGY_DIR / f"{library_id}.json"
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return None


def methodology_library_ids_for_mode(methodology_mode: str) -> list[str]:
    mapping = {
        "generic": ["generic_qualitative"],
        "gt": ["generic_qualitative", "grounded_theory"],
        "cda": ["generic_qualitative", "cda_fairclough", "cda_van_dijk", "cda_wodak_dha"],
        "mixed": [
            "generic_qualitative",
            "grounded_theory",
            "cda_fairclough",
            "cda_van_dijk",
            "cda_wodak_dha",
            "mixed_gt_cda",
        ],
    }
    return mapping.get(methodology_mode, mapping["generic"])


def get_relevant_methodology_libraries(active_project: sqlite3.Row) -> list[dict]:
    libraries = []
    for library_id in methodology_library_ids_for_mode(active_project["methodology_mode"]):
        library = load_methodology_library(library_id)
        if library:
            libraries.append(library)
    return libraries


def get_methodology_phase_prompts(library_ids: list[str], phase_ids: list[str], limit: int = 8) -> list[dict]:
    prompts = []
    for library_id in library_ids:
        library = load_methodology_library(library_id)
        if not library:
            continue
        for phase in library.get("phases", []):
            if phase.get("id") in phase_ids:
                for prompt in phase.get("prompts", []):
                    prompts.append(
                        {
                            "library_title": library["title"],
                            "phase_title": phase["title"],
                            "prompt": prompt,
                        }
                    )
    return prompts[:limit]


def get_methodology_concept_prompts(library_ids: list[str], concept_ids: list[str], limit: int = 8) -> list[dict]:
    prompts = []
    for library_id in library_ids:
        library = load_methodology_library(library_id)
        if not library:
            continue
        for concept in library.get("concepts", []):
            if concept.get("id") in concept_ids:
                for prompt in concept.get("prompts", []):
                    prompts.append(
                        {
                            "library_title": library["title"],
                            "phase_title": concept["term"],
                            "prompt": prompt,
                        }
                    )
    return prompts[:limit]


def get_methodology_helper(active_project: sqlite3.Row, context: str) -> dict:
    mode = active_project["methodology_mode"]
    if context == "gt_workspace":
        prompts = get_methodology_phase_prompts(
            ["grounded_theory"],
            ["open_coding", "axial_coding", "category_development", "constant_comparison"],
            10,
        )
        return {"title": "Grounded Theory helper", "prompts": prompts}
    if context == "gt_compare":
        return {
            "title": "Constant comparison helper",
            "prompts": get_methodology_phase_prompts(["grounded_theory"], ["constant_comparison"], 8),
        }
    if context == "cda_workspace":
        return {
            "title": "CDA helper",
            "prompts": get_methodology_phase_prompts(
                ["cda_fairclough", "cda_van_dijk", "cda_wodak_dha"],
                ["text_analysis", "discursive_practice", "social_practice", "power_access", "ideology_groups", "discursive_strategies"],
                10,
            ),
        }
    if context == "cda_features":
        return {
            "title": "Discourse feature helper",
            "prompts": get_methodology_concept_prompts(
                ["cda_fairclough", "cda_wodak_dha"],
                ["metaphor", "presupposition", "modality", "legitimation", "argumentation"],
                8,
            ),
        }
    if context == "cda_voice_silence":
        return {
            "title": "Voice and silence helper",
            "prompts": get_methodology_phase_prompts(["cda_van_dijk"], ["power_access"], 5)
            + get_methodology_concept_prompts(["cda_van_dijk"], ["voice_silence", "dominance"], 5),
        }
    if context == "model":
        library_ids = ["mixed_gt_cda"] if mode == "mixed" else methodology_library_ids_for_mode(mode)
        return {
            "title": "Model-building methodology helper",
            "prompts": get_methodology_phase_prompts(
                library_ids,
                ["relation_modeling", "argument_building", "theoretical_integration", "analytical_synthesis", "social_practice"],
                10,
            ),
        }
    if context == "document":
        if mode == "gt":
            library_ids = ["grounded_theory"]
            phase_ids = ["open_coding", "memo_writing"]
        elif mode == "cda":
            library_ids = ["cda_fairclough", "cda_van_dijk", "cda_wodak_dha"]
            phase_ids = ["text_analysis", "power_access", "discursive_strategies"]
        elif mode == "mixed":
            library_ids = ["grounded_theory", "cda_fairclough", "cda_van_dijk", "mixed_gt_cda"]
            phase_ids = ["open_coding", "memo_writing", "text_analysis", "actor_voice_mapping", "discourse_feature_marking"]
        else:
            library_ids = ["generic_qualitative"]
            phase_ids = ["close_reading", "open_coding"]
        return {
            "title": "Reading and coding helper",
            "prompts": get_methodology_phase_prompts(library_ids, phase_ids, 8),
        }
    return {"title": "Methodology helper", "prompts": []}


def download_text(content: str, filename: str, mimetype: str) -> Response:
    return Response(
        content,
        content_type=mimetype,
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def download_binary(content: bytes, filename: str, mimetype: str) -> Response:
    return Response(
        content,
        mimetype=mimetype,
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def make_csv(headers: list[str], rows: list[dict]) -> str:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({header: empty(row.get(header)) for header in headers})
    return output.getvalue()


def markdown_blockquote(text: str) -> str:
    text = empty(text).strip()
    if not text:
        return "> "
    return "\n".join(f"> {line}" if line else ">" for line in text.splitlines())


def get_network_filters() -> dict:
    def bool_param(name: str, default: bool = True) -> bool:
        values = request.args.getlist(name)
        if not values:
            return default
        value = values[-1]
        return value not in {"0", "false", "False", "off", ""}

    def int_param(name: str, default: int, minimum: int, maximum: int) -> int:
        try:
            value = int(request.args.get(name, default))
        except (TypeError, ValueError):
            value = default
        return max(minimum, min(maximum, value))

    document_id_raw = request.args.get("document_id", "").strip()
    document_id = int(document_id_raw) if document_id_raw.isdigit() else None
    code_type = request.args.get("code_type", "").strip()
    marker_type = request.args.get("marker_type", "").strip()
    actor_type = request.args.get("actor_type", "").strip()
    feature_type = request.args.get("feature_type", "").strip()
    layout = request.args.get("layout", "columns").strip()
    return {
        "include_codes": bool_param("include_codes", True),
        "include_markers": bool_param("include_markers", True),
        "include_actors": bool_param("include_actors", True),
        "include_features": bool_param("include_features", True),
        "include_hierarchy": bool_param("include_hierarchy", False),
        "min_weight": int_param("min_weight", 1, 1, 999),
        "max_nodes": int_param("max_nodes", 80, 5, 300),
        "document_id": document_id,
        "code_type": code_type if code_type in {"open", "axial", "category"} else "",
        "marker_type": marker_type if marker_type in CDA_MARKER_TYPES else "",
        "actor_type": actor_type if actor_type in ACTOR_TYPES else "",
        "feature_type": feature_type if feature_type in DISCOURSE_FEATURE_TYPES else "",
        "layout": layout if layout in {"columns", "force", "circle"} else "columns",
    }


def default_network_filters() -> dict:
    return {
        "include_codes": True,
        "include_markers": True,
        "include_actors": True,
        "include_features": True,
        "include_hierarchy": False,
        "min_weight": 1,
        "max_nodes": 80,
        "document_id": None,
        "code_type": "",
        "marker_type": "",
        "actor_type": "",
        "feature_type": "",
        "layout": "columns",
    }


def normalize_network_feature_value(value: str) -> str:
    normalized = re.sub(r"\s+", " ", empty(value).strip().lower())
    normalized = re.sub(r"[^a-z0-9_ -]+", "", normalized)
    normalized = normalized.replace(" ", "_")
    return normalized[:48] or "feature"


def add_network_node(nodes: dict, node_id: str, node_type: str, subtype: str, label: str, color: str = "") -> None:
    if node_id not in nodes:
        nodes[node_id] = {
            "id": node_id,
            "type": node_type,
            "subtype": subtype,
            "label": label,
            "count": 0,
        }
        if color:
            nodes[node_id]["color"] = color


def append_segment_node(segment_nodes: dict[int, set[str]], segment_id: int, node_id: str) -> None:
    segment_nodes.setdefault(segment_id, set()).add(node_id)


def build_cooccurrence_network(active_project: sqlite3.Row, filters: dict) -> dict:
    project_id = active_project["id"]
    nodes: dict[str, dict] = {}
    segment_nodes: dict[int, set[str]] = {}
    segment_previews = get_network_segment_previews(project_id, filters["document_id"])

    if filters["include_codes"]:
        add_code_network_nodes(project_id, filters, nodes, segment_nodes)
    if filters["include_markers"]:
        add_marker_network_nodes(project_id, filters, nodes, segment_nodes)
    if filters["include_actors"]:
        add_actor_network_nodes(project_id, filters, nodes, segment_nodes)
    if filters["include_features"]:
        add_feature_network_nodes(project_id, filters, nodes, segment_nodes)

    nodes = {
        node_id: node
        for node_id, node in sorted(
            nodes.items(),
            key=lambda item: (-item[1]["count"], item[1]["type"], item[1]["label"].lower(), item[0]),
        )[: filters["max_nodes"]]
    }
    allowed_node_ids = set(nodes)
    edges: dict[tuple[str, str], dict] = {}
    for segment_id, node_ids in segment_nodes.items():
        present = sorted(node_id for node_id in node_ids if node_id in allowed_node_ids)
        for index, source in enumerate(present):
            for target in present[index + 1:]:
                edge_key = (source, target)
                if edge_key not in edges:
                    edges[edge_key] = {
                        "source": source,
                        "target": target,
                        "weight": 0,
                        "relation": "co_occurs_in_segment",
                        "segments": [],
                    }
                edges[edge_key]["weight"] += 1
                edges[edge_key]["segments"].append(segment_id)

    filtered_edges = []
    for edge in edges.values():
        if edge["weight"] < filters["min_weight"]:
            continue
        previews = [segment_previews[segment_id] for segment_id in edge["segments"][:10] if segment_id in segment_previews]
        edge["segment_previews"] = previews
        edge["segment_preview_count"] = len(previews)
        edge["segment_count"] = len(edge["segments"])
        filtered_edges.append(edge)

    connected_node_ids = {edge["source"] for edge in filtered_edges} | {edge["target"] for edge in filtered_edges}
    if filtered_edges:
        nodes = {node_id: node for node_id, node in nodes.items() if node_id in connected_node_ids}

    return {
        "nodes": list(nodes.values()),
        "edges": sorted(filtered_edges, key=lambda edge: (-edge["weight"], edge["source"], edge["target"])),
        "meta": {
            "project_id": project_id,
            "project_name": active_project["name"],
            "generated_at": export_timestamp(),
            "mode": active_project["methodology_mode"],
            "min_weight": filters["min_weight"],
            "max_nodes": filters["max_nodes"],
            "layout": filters["layout"],
            "filters": filters,
        },
    }


def add_code_network_nodes(project_id: int, filters: dict, nodes: dict, segment_nodes: dict[int, set[str]]) -> None:
    params: list = [project_id]
    where = ["documents.project_id = ?"]
    if filters["document_id"]:
        where.append("documents.id = ?")
        params.append(filters["document_id"])
    rows = query_all(
        f"""
        SELECT
            segments.id AS segment_id,
            codes.id, codes.name, codes.code_type, codes.color,
            axial.id AS axial_id, axial.name AS axial_name, axial.color AS axial_color,
            category.id AS category_id, category.name AS category_name, category.color AS category_color
        FROM segment_codes
        JOIN segments ON segments.id = segment_codes.segment_id
        JOIN documents ON documents.id = segments.document_id
        JOIN codes ON codes.id = segment_codes.code_id
        LEFT JOIN codes AS axial ON axial.id = codes.parent_id AND axial.project_id = codes.project_id
        LEFT JOIN codes AS category ON category.id = axial.parent_id AND category.project_id = codes.project_id
        WHERE {" AND ".join(where)}
        ORDER BY segments.id, codes.name COLLATE NOCASE
        """,
        tuple(params),
    )
    counted_pairs: set[tuple[str, int]] = set()
    for row in rows:
        candidates = [(row["id"], row["name"], row["code_type"], row["color"])]
        if filters["include_hierarchy"]:
            if row["axial_id"]:
                candidates.append((row["axial_id"], row["axial_name"], "axial", row["axial_color"]))
            if row["category_id"]:
                candidates.append((row["category_id"], row["category_name"], "category", row["category_color"]))
        for code_id, name, code_type, color in candidates:
            if filters["code_type"] and code_type != filters["code_type"]:
                continue
            node_id = f"code_{code_id}"
            add_network_node(nodes, node_id, "code", code_type, name, color or "")
            pair = (node_id, row["segment_id"])
            if pair not in counted_pairs:
                nodes[node_id]["count"] += 1
                counted_pairs.add(pair)
            append_segment_node(segment_nodes, row["segment_id"], node_id)


def add_marker_network_nodes(project_id: int, filters: dict, nodes: dict, segment_nodes: dict[int, set[str]]) -> None:
    params: list = [project_id]
    where = ["documents.project_id = ?"]
    if filters["document_id"]:
        where.append("documents.id = ?")
        params.append(filters["document_id"])
    if filters["marker_type"]:
        where.append("discourse_markers.marker_type = ?")
        params.append(filters["marker_type"])
    rows = query_all(
        f"""
        SELECT segments.id AS segment_id, discourse_markers.id,
               discourse_markers.name, discourse_markers.marker_type, discourse_markers.color
        FROM segment_discourse_markers
        JOIN segments ON segments.id = segment_discourse_markers.segment_id
        JOIN documents ON documents.id = segments.document_id
        JOIN discourse_markers ON discourse_markers.id = segment_discourse_markers.marker_id
        WHERE {" AND ".join(where)}
        ORDER BY segments.id, discourse_markers.name COLLATE NOCASE
        """,
        tuple(params),
    )
    for row in rows:
        node_id = f"marker_{row['id']}"
        add_network_node(nodes, node_id, "discourse_marker", row["marker_type"], row["name"], row["color"] or "")
        nodes[node_id]["count"] += 1
        append_segment_node(segment_nodes, row["segment_id"], node_id)


def add_actor_network_nodes(project_id: int, filters: dict, nodes: dict, segment_nodes: dict[int, set[str]]) -> None:
    params: list = [project_id]
    where = ["documents.project_id = ?"]
    if filters["document_id"]:
        where.append("documents.id = ?")
        params.append(filters["document_id"])
    if filters["actor_type"]:
        where.append("actors.actor_type = ?")
        params.append(filters["actor_type"])
    rows = query_all(
        f"""
        SELECT segments.id AS segment_id, actors.id, actors.name, actors.actor_type,
               segment_actors.relation_type
        FROM segment_actors
        JOIN segments ON segments.id = segment_actors.segment_id
        JOIN documents ON documents.id = segments.document_id
        JOIN actors ON actors.id = segment_actors.actor_id
        WHERE {" AND ".join(where)}
        ORDER BY segments.id, actors.name COLLATE NOCASE
        """,
        tuple(params),
    )
    for row in rows:
        node_id = f"actor_{row['id']}"
        add_network_node(nodes, node_id, "actor", row["actor_type"], row["name"])
        nodes[node_id]["count"] += 1
        relation_counts = nodes[node_id].setdefault("relation_counts", {})
        relation_counts[row["relation_type"]] = relation_counts.get(row["relation_type"], 0) + 1
        append_segment_node(segment_nodes, row["segment_id"], node_id)


def add_feature_network_nodes(project_id: int, filters: dict, nodes: dict, segment_nodes: dict[int, set[str]]) -> None:
    params: list = [project_id]
    where = ["documents.project_id = ?"]
    if filters["document_id"]:
        where.append("documents.id = ?")
        params.append(filters["document_id"])
    if filters["feature_type"]:
        where.append("discourse_features.feature_type = ?")
        params.append(filters["feature_type"])
    rows = query_all(
        f"""
        SELECT segments.id AS segment_id, discourse_features.feature_type,
               discourse_features.value
        FROM discourse_features
        JOIN segments ON segments.id = discourse_features.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE {" AND ".join(where)}
        ORDER BY segments.id, discourse_features.feature_type COLLATE NOCASE, discourse_features.value COLLATE NOCASE
        """,
        tuple(params),
    )
    for row in rows:
        normalized_value = normalize_network_feature_value(row["value"])
        node_id = f"feature_{row['feature_type']}_{normalized_value}"
        label = f"{DISCOURSE_FEATURE_TYPES.get(row['feature_type'], row['feature_type'])}: {truncate_text(row['value'], 44)}"
        add_network_node(nodes, node_id, "discourse_feature", row["feature_type"], label)
        nodes[node_id]["count"] += 1
        append_segment_node(segment_nodes, row["segment_id"], node_id)


def get_network_segment_previews(project_id: int, document_id: int | None = None) -> dict[int, dict]:
    params: list = [project_id]
    where = ["documents.project_id = ?"]
    if document_id:
        where.append("documents.id = ?")
        params.append(document_id)
    rows = query_all(
        f"""
        SELECT segments.id AS segment_id, segments.document_id,
               COALESCE(segments.name, '') AS segment_title,
               segments.selected_text, documents.title AS document_title
        FROM segments
        JOIN documents ON documents.id = segments.document_id
        WHERE {" AND ".join(where)}
        """,
        tuple(params),
    )
    return {
        row["segment_id"]: {
            "segment_id": row["segment_id"],
            "document_id": row["document_id"],
            "document_title": row["document_title"],
            "segment_title": row["segment_title"] or f"Segment {row['segment_id']}",
            "text_preview": truncate_text(row["selected_text"], 180),
        }
        for row in rows
    }


def generate_cooccurrence_network_json(active_project: sqlite3.Row, filters: dict) -> str:
    return json.dumps(build_cooccurrence_network(active_project, filters), ensure_ascii=False, indent=2)


def generate_cooccurrence_edges_csv(active_project: sqlite3.Row, filters: dict) -> str:
    graph = build_cooccurrence_network(active_project, filters)
    nodes = {node["id"]: node for node in graph["nodes"]}
    rows = []
    for edge in graph["edges"]:
        source = nodes.get(edge["source"], {})
        target = nodes.get(edge["target"], {})
        rows.append(
            {
                "source_id": edge["source"],
                "source_label": source.get("label", ""),
                "source_type": source.get("type", ""),
                "target_id": edge["target"],
                "target_label": target.get("label", ""),
                "target_type": target.get("type", ""),
                "weight": edge["weight"],
                "segment_ids": ";".join(str(segment_id) for segment_id in edge["segments"]),
            }
        )
    return make_csv(
        ["source_id", "source_label", "source_type", "target_id", "target_label", "target_type", "weight", "segment_ids"],
        rows,
    )


def build_integrity_report(active_project: sqlite3.Row) -> dict:
    db = get_db()
    checks = [
        (
            "Segments without documents",
            """
            SELECT COUNT(*) AS count
            FROM segments s
            LEFT JOIN documents d ON d.id = s.document_id
            WHERE d.id IS NULL
            """,
        ),
        (
            "Segment-code links with missing segments",
            """
            SELECT COUNT(*) AS count
            FROM segment_codes sc
            LEFT JOIN segments s ON s.id = sc.segment_id
            WHERE s.id IS NULL
            """,
        ),
        (
            "Segment-code links with missing codes",
            """
            SELECT COUNT(*) AS count
            FROM segment_codes sc
            LEFT JOIN codes c ON c.id = sc.code_id
            WHERE c.id IS NULL
            """,
        ),
        (
            "Code hierarchy links with missing parents",
            """
            SELECT COUNT(*) AS count
            FROM codes child
            LEFT JOIN codes parent ON parent.id = child.parent_id
            WHERE child.parent_id IS NOT NULL
                AND parent.id IS NULL
            """,
        ),
        (
            "Segment-marker links with missing segments",
            """
            SELECT COUNT(*) AS count
            FROM segment_discourse_markers sdm
            LEFT JOIN segments s ON s.id = sdm.segment_id
            WHERE s.id IS NULL
            """,
        ),
        (
            "Segment-marker links with missing markers",
            """
            SELECT COUNT(*) AS count
            FROM segment_discourse_markers sdm
            LEFT JOIN discourse_markers dm ON dm.id = sdm.marker_id
            WHERE dm.id IS NULL
            """,
        ),
        (
            "Segment-actor annotations with missing segments",
            """
            SELECT COUNT(*) AS count
            FROM segment_actors sa
            LEFT JOIN segments s ON s.id = sa.segment_id
            WHERE s.id IS NULL
            """,
        ),
        (
            "Segment-actor annotations with missing actors",
            """
            SELECT COUNT(*) AS count
            FROM segment_actors sa
            LEFT JOIN actors a ON a.id = sa.actor_id
            WHERE a.id IS NULL
            """,
        ),
        (
            "Discourse features with missing segments",
            """
            SELECT COUNT(*) AS count
            FROM discourse_features df
            LEFT JOIN segments s ON s.id = df.segment_id
            WHERE s.id IS NULL
            """,
        ),
        (
            "Document-tag links with missing documents",
            """
            SELECT COUNT(*) AS count
            FROM document_tags dt
            LEFT JOIN documents d ON d.id = dt.document_id
            WHERE d.id IS NULL
            """,
        ),
        (
            "Document-tag links with missing tags",
            """
            SELECT COUNT(*) AS count
            FROM document_tags dt
            LEFT JOIN tags t ON t.id = dt.tag_id
            WHERE t.id IS NULL
            """,
        ),
    ]
    check_results = []
    for label, sql in checks:
        count = db.execute(sql).fetchone()["count"]
        check_results.append(
            {
                "label": label,
                "count": count,
                "status": "ok" if count == 0 else "warning",
            }
        )

    project_id = active_project["id"]
    counts = [
        (
            "Documents",
            "SELECT COUNT(*) AS count FROM documents WHERE project_id = ?",
        ),
        (
            "Codes",
            "SELECT COUNT(*) AS count FROM codes WHERE project_id = ?",
        ),
        (
            "Memos",
            "SELECT COUNT(*) AS count FROM memos WHERE project_id = ?",
        ),
        (
            "CDA markers",
            "SELECT COUNT(*) AS count FROM discourse_markers WHERE project_id = ?",
        ),
        (
            "Actors",
            "SELECT COUNT(*) AS count FROM actors WHERE project_id = ?",
        ),
        (
            "Analytical relations",
            "SELECT COUNT(*) AS count FROM relations WHERE project_id = ?",
        ),
    ]
    count_results = [
        {
            "label": label,
            "count": db.execute(sql, (project_id,)).fetchone()["count"],
        }
        for label, sql in counts
    ]
    count_results.insert(
        1,
        {
            "label": "Segments",
            "count": db.execute(
                """
                SELECT COUNT(*) AS count
                FROM segments s
                JOIN documents d ON d.id = s.document_id
                WHERE d.project_id = ?
                """,
                (project_id,),
            ).fetchone()["count"],
        },
    )
    return {
        "generated_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "checks": check_results,
        "counts": count_results,
        "warning_count": sum(1 for check in check_results if check["count"]),
    }


def build_codebook_markdown(active_project: sqlite3.Row) -> str:
    project_id = active_project["id"]
    lines = [
        "# discourseLab Codebook",
        "",
        f"Project: {active_project['name']}",
        f"Exported: {export_timestamp()}",
        "",
    ]
    codes = query_all(
        """
        SELECT
            codes.*,
            parent.name AS parent_name,
            COUNT(DISTINCT segment_codes.segment_id) AS usage_count,
            COUNT(DISTINCT memos.id) AS linked_memo_count
        FROM codes
        LEFT JOIN segment_codes ON segment_codes.code_id = codes.id
        LEFT JOIN codes AS parent ON parent.id = codes.parent_id
        LEFT JOIN memos ON memos.project_id = codes.project_id
            AND memos.linked_entity_type = 'code'
            AND memos.linked_entity_id = codes.id
        WHERE codes.project_id = ?
        GROUP BY codes.id
        ORDER BY codes.code_type COLLATE NOCASE, codes.name COLLATE NOCASE
        """,
        (project_id,),
    )
    if not codes:
        lines.append("No codes created yet.")
        lines.append("")

    for code in codes:
        lines.extend(
            [
                f"## {code['name']}",
                "",
                f"- Type: {code['code_type']}",
                f"- Parent: {code['parent_name'] or 'None'}",
                f"- Color: {code['color'] or ''}",
                f"- Usage count: {code['usage_count']}",
                f"- Linked memo count: {code['linked_memo_count']}",
                "",
                f"**Description:** {code['description'] or ''}",
                "",
                f"**Definition:** {code['definition'] or ''}",
                "",
                f"**Include when:** {code['include_when'] or ''}",
                "",
                f"**Exclude when:** {code['exclude_when'] or ''}",
                "",
                f"**Example:** {code['example'] or ''}",
                "",
                f"**Analytical note:** {code['analytical_note'] or ''}",
                "",
            ]
        )
        if code["code_type"] in {"axial", "category"}:
            lines.extend(
                [
                    f"**GT conditions:** {code['gt_conditions'] or ''}",
                    "",
                    f"**GT context:** {code['gt_context'] or ''}",
                    "",
                    f"**GT actions/interactions:** {code['gt_actions_interactions'] or ''}",
                    "",
                    f"**GT consequences:** {code['gt_consequences'] or ''}",
                    "",
                    f"**GT properties:** {code['gt_properties'] or ''}",
                    "",
                    f"**GT dimensions:** {code['gt_dimensions'] or ''}",
                    "",
                    f"**GT theoretical note:** {code['gt_theoretical_note'] or ''}",
                    "",
                ]
            )
    marker_rows = query_all(
        """
        SELECT
            discourse_markers.id,
            discourse_markers.name,
            discourse_markers.marker_type,
            discourse_markers.description,
            discourse_markers.color,
            discourse_markers.created_at,
            discourse_markers.updated_at,
            COUNT(DISTINCT segment_discourse_markers.segment_id) AS usage_count
        FROM discourse_markers
        LEFT JOIN segment_discourse_markers
            ON segment_discourse_markers.marker_id = discourse_markers.id
        WHERE discourse_markers.project_id = ?
        GROUP BY discourse_markers.id
        ORDER BY discourse_markers.marker_type COLLATE NOCASE, discourse_markers.name COLLATE NOCASE
        """,
        (project_id,),
    )
    lines.extend(["## CDA Markers", ""])
    if marker_rows:
        for marker in marker_rows:
            lines.extend(
                [
                    f"### {marker['name']}",
                    "",
                    f"- Type: {marker['marker_type']}",
                    f"- Type explanation: {CDA_MARKER_TYPE_EXPLANATIONS.get(marker['marker_type'], '')}",
                    f"- Description: {marker['description'] or ''}",
                    f"- Color: {marker['color'] or ''}",
                    f"- Usage count: {marker['usage_count']}",
                    f"- Created at: {marker['created_at']}",
                    f"- Updated at: {marker['updated_at']}",
                    "",
                ]
            )
    else:
        lines.extend(["No CDA markers created yet.", ""])

    actor_rows = query_all(
        """
        SELECT
            actors.id,
            actors.name,
            actors.actor_type,
            actors.description,
            actors.created_at,
            actors.updated_at,
            COUNT(DISTINCT segment_actors.id) AS annotation_count
        FROM actors
        LEFT JOIN segment_actors ON segment_actors.actor_id = actors.id
        WHERE actors.project_id = ?
        GROUP BY actors.id
        ORDER BY actors.actor_type COLLATE NOCASE, actors.name COLLATE NOCASE
        """,
        (project_id,),
    )
    relation_count_rows = query_all(
        """
        SELECT actors.id AS actor_id, segment_actors.relation_type, COUNT(*) AS relation_count
        FROM actors
        JOIN segment_actors ON segment_actors.actor_id = actors.id
        WHERE actors.project_id = ?
        GROUP BY actors.id, segment_actors.relation_type
        """,
        (project_id,),
    )
    actor_relation_counts: dict[int, dict[str, int]] = {}
    for row in relation_count_rows:
        actor_relation_counts.setdefault(row["actor_id"], {})[row["relation_type"]] = row["relation_count"]

    lines.extend(["## Actors", ""])
    if actor_rows:
        for actor in actor_rows:
            counts = actor_relation_counts.get(actor["id"], {})
            lines.extend(
                [
                    f"### {actor['name']}",
                    "",
                    f"- Type: {actor['actor_type']}",
                    f"- Type explanation: {ACTOR_TYPE_EXPLANATIONS.get(actor['actor_type'], '')}",
                    f"- Description: {actor['description'] or ''}",
                    f"- Annotation count: {actor['annotation_count']}",
                    "- Relation counts:",
                ]
            )
            for relation_type in ACTOR_RELATION_TYPES:
                lines.append(f"  - {relation_type}: {counts.get(relation_type, 0)}")
            lines.extend(
                [
                    f"- Created at: {actor['created_at']}",
                    f"- Updated at: {actor['updated_at']}",
                    "",
                ]
            )
    else:
        lines.extend(["No actors created yet.", ""])
    return "\n".join(lines)


def get_coded_segment_export_rows(project_id: int) -> list[sqlite3.Row]:
    return query_all(
        """
        SELECT
            documents.id AS document_id,
            documents.title AS document_title,
            documents.original_filename AS document_original_filename,
            segments.id AS segment_id,
            COALESCE(segments.name, '') AS segment_title,
            segments.selected_text,
            segments.note AS segment_note,
            segments.start_offset,
            segments.end_offset,
            codes.id AS code_id,
            codes.name AS code_name,
            codes.code_type,
            codes.parent_id AS code_parent_id,
            CASE
                WHEN codes.code_type = 'open' THEN axial.name
                WHEN codes.code_type = 'axial' THEN codes.name
                ELSE NULL
            END AS axial_code_name,
            CASE
                WHEN codes.code_type = 'open' THEN category.name
                WHEN codes.code_type = 'category' THEN codes.name
                ELSE NULL
            END AS category_name,
            codes.description AS code_description,
            codes.definition AS code_definition,
            segments.created_at
        FROM segment_codes
        JOIN segments ON segments.id = segment_codes.segment_id
        JOIN documents ON documents.id = segments.document_id
        JOIN codes ON codes.id = segment_codes.code_id
        LEFT JOIN codes AS axial ON axial.id = codes.parent_id
            AND axial.code_type = 'axial'
        LEFT JOIN codes AS category ON category.id = axial.parent_id
            AND category.code_type = 'category'
        WHERE documents.project_id = ? AND codes.project_id = ?
        ORDER BY documents.title COLLATE NOCASE, segments.start_offset, codes.name COLLATE NOCASE
        """,
        (project_id, project_id),
    )


def generate_coded_segments_csv(active_project: sqlite3.Row) -> str:
    headers = [
        "project_name",
        "document_id",
        "document_title",
        "document_original_filename",
        "segment_id",
        "segment_title",
        "selected_text",
        "segment_note",
        "start_offset",
        "end_offset",
        "code_id",
        "code_name",
        "code_type",
        "code_parent_id",
        "axial_code_name",
        "category_name",
        "code_description",
        "code_definition",
        "created_at",
    ]
    rows = []
    for row in get_coded_segment_export_rows(active_project["id"]):
        data = row_to_dict(row)
        data["project_name"] = active_project["name"]
        rows.append(data)
    return make_csv(headers, rows)


def generate_coded_segments_markdown(active_project: sqlite3.Row) -> str:
    rows = get_coded_segment_export_rows(active_project["id"])
    lines = [
        "# Coded Segments",
        "",
        f"Project: {active_project['name']}",
        f"Exported: {export_timestamp()}",
        "",
    ]
    if not rows:
        lines.extend(["No coded segments found.", ""])
        return "\n".join(lines)

    current_document = None
    current_segment = None
    for index, row in enumerate(rows):
        if row["document_id"] != current_document:
            current_document = row["document_id"]
            current_segment = None
            lines.extend([f"## Document: {row['document_title']}", ""])
        if row["segment_id"] != current_segment:
            current_segment = row["segment_id"]
            segment_label = row["segment_title"] or f"Segment {row['segment_id']}"
            lines.extend(
                [
                    f"### Segment: {segment_label}",
                    "",
                    "Selected text:",
                    "",
                    markdown_blockquote(row["selected_text"]),
                    "",
                    "Segment note:",
                    "",
                    row["segment_note"] or "",
                    "",
                    "Codes:",
                ]
            )
        lines.append(f"- {row['code_name']}")
        lines.append(f"  - Type: {row['code_type']}")
        if row["axial_code_name"]:
            lines.append(f"  - Axial parent: {row['axial_code_name']}")
        if row["category_name"]:
            lines.append(f"  - Category parent: {row['category_name']}")

        next_is_new_segment = True
        next_index = index + 1
        if next_index < len(rows):
            next_is_new_segment = rows[next_index]["segment_id"] != current_segment
        if next_is_new_segment:
            memos = get_memos_for_entity(active_project["id"], "segment", row["segment_id"])
            lines.extend(["", "Memos linked to this segment:"])
            if memos:
                for memo in memos:
                    lines.append(
                        f"- {memo['title']}, {memo['status']}, {memo['memo_type']}"
                    )
                    lines.append(f"  - {empty(memo['body']).replace(chr(10), ' ')}")
            else:
                lines.append("- None")
            lines.append("")
    return "\n".join(lines)


def generate_memos_markdown(active_project: sqlite3.Row) -> str:
    memos = get_memos_for_project(
        active_project["id"],
        {"memo_type": "", "status": "", "linked_entity_type": "", "linked_entity_id": ""},
    )
    lines = [
        "# Memos",
        "",
        f"Project: {active_project['name']}",
        f"Exported: {export_timestamp()}",
        "",
    ]
    if not memos:
        lines.extend(["No memos found.", ""])
        return "\n".join(lines)
    current_type = None
    current_status = None
    for memo in sorted(memos, key=lambda item: (item["memo_type"], item["status"], item["title"].lower())):
        if memo["memo_type"] != current_type:
            current_type = memo["memo_type"]
            current_status = None
            lines.extend([f"## {MEMO_TYPES.get(current_type, current_type)}", ""])
        if memo["status"] != current_status:
            current_status = memo["status"]
            lines.extend([f"### {MEMO_STATUSES.get(current_status, current_status)}", ""])
        lines.extend(
            [
                f"#### {memo['title']}",
                "",
                f"Status: {MEMO_STATUSES.get(memo['status'], memo['status'])}",
                f"Linked to: {memo['linked_entity_label']}",
                f"Created: {memo['created_at']}",
                f"Updated: {memo['updated_at']}",
                "",
                memo["body"] or "",
                "",
            ]
        )
    return "\n".join(lines)


def generate_cda_features_csv(active_project: sqlite3.Row) -> str:
    headers = [
        "project_name",
        "document_id",
        "document_title",
        "segment_id",
        "segment_title",
        "selected_text",
        "feature_id",
        "feature_type",
        "value",
        "interpretation",
        "created_at",
        "updated_at",
    ]
    rows = []
    feature_rows = query_all(
        """
        SELECT documents.id AS document_id, documents.title AS document_title,
               segments.id AS segment_id, COALESCE(segments.name, '') AS segment_title,
               segments.selected_text, discourse_features.id AS feature_id,
               discourse_features.feature_type, discourse_features.value,
               discourse_features.interpretation, discourse_features.created_at,
               discourse_features.updated_at
        FROM discourse_features
        JOIN segments ON segments.id = discourse_features.segment_id
        JOIN documents ON documents.id = segments.document_id
        WHERE documents.project_id = ?
        ORDER BY documents.title COLLATE NOCASE, segments.start_offset,
                 discourse_features.feature_type
        """,
        (active_project["id"],),
    )
    for row in feature_rows:
        data = row_to_dict(row)
        data["project_name"] = active_project["name"]
        rows.append(data)
    return make_csv(headers, rows)


def generate_voice_silence_csv(active_project: sqlite3.Row) -> str:
    headers = [
        "project_name",
        "actor_id",
        "actor_name",
        "actor_type",
        "actor_description",
        *ACTOR_RELATION_TYPES.keys(),
        "total_annotations",
        "segment_count",
        "document_count",
    ]
    actors_by_id = {
        actor["id"]: actor
        for actor in query_all(
            """
            SELECT id, name, actor_type, description
            FROM actors
            WHERE project_id = ?
            ORDER BY name COLLATE NOCASE
            """,
            (active_project["id"],),
        )
    }
    rows = []
    for report_row in get_voice_silence_report(active_project["id"]):
        actor = actors_by_id.get(report_row["id"])
        row = {
            "project_name": active_project["name"],
            "actor_id": report_row["id"],
            "actor_name": report_row["name"],
            "actor_type": report_row["actor_type"],
            "actor_description": actor["description"] if actor else "",
            "total_annotations": report_row["total_annotations"],
            "segment_count": report_row["segment_count"],
            "document_count": report_row["document_count"],
        }
        for relation in ACTOR_RELATION_TYPES:
            row[relation] = report_row[relation]
        rows.append(row)
    return make_csv(headers, rows)


def generate_gt_hierarchy_markdown(active_project: sqlite3.Row) -> str:
    project_id = active_project["id"]
    categories = get_gt_categories(project_id)
    unassigned_axial = [
        code for code in get_gt_axial_codes(project_id) if code.get("parent_id") is None
    ]
    unassigned_open = [
        code for code in get_gt_open_codes(project_id) if code.get("parent_id") is None
    ]
    lines = [
        "# Grounded Theory Hierarchy",
        "",
        f"Project: {active_project['name']}",
        f"Exported: {export_timestamp()}",
        "",
        "## Categories",
        "",
    ]
    if not categories:
        lines.extend(["No categories found.", ""])
    for category in categories:
        axial_children = get_child_codes(category["id"], project_id, "axial")
        lines.extend(
            [
                f"### Category: {category['name']}",
                "",
                f"Description: {category.get('description') or ''}",
                f"Definition: {category.get('definition') or ''}",
                f"Analytical note: {category.get('analytical_note') or ''}",
                f"GT theoretical note: {category.get('gt_theoretical_note') or ''}",
                "",
                "Child axial codes:",
                "",
            ]
        )
        if not axial_children:
            lines.extend(["None.", ""])
        for axial in axial_children:
            axial_full = get_code_for_project(axial["id"], project_id)
            open_children = get_child_codes(axial["id"], project_id, "open")
            lines.extend(
                [
                    f"#### Axial code: {axial_full['name']}",
                    "",
                    f"Description: {axial_full['description'] or ''}",
                    f"Conditions: {axial_full['gt_conditions'] or ''}",
                    f"Context: {axial_full['gt_context'] or ''}",
                    f"Actions/interactions: {axial_full['gt_actions_interactions'] or ''}",
                    f"Consequences: {axial_full['gt_consequences'] or ''}",
                    f"Properties: {axial_full['gt_properties'] or ''}",
                    f"Dimensions: {axial_full['gt_dimensions'] or ''}",
                    "",
                    "Child open codes:",
                ]
            )
            if open_children:
                for open_code in open_children:
                    usage_count = get_code_usage_count(open_code["id"])
                    lines.append(f"- {open_code['name']}")
                    lines.append(f"  - Description: {open_code['description'] or ''}")
                    lines.append(f"  - Assigned segment count: {usage_count}")
            else:
                lines.append("- None")
            lines.extend(["", "Representative segments:"])
            segments = get_hierarchy_segments_for_code(axial_full, project_id)[:5]
            if segments:
                for segment in segments:
                    lines.extend([markdown_blockquote(segment["selected_text"]), ""])
            else:
                lines.extend(["> None", ""])

    lines.extend(["## Unassigned axial codes", ""])
    if unassigned_axial:
        for code in unassigned_axial:
            lines.append(f"- {code['name']}")
    else:
        lines.append("None.")
    lines.extend(["", "## Unassigned open codes", ""])
    if unassigned_open:
        for code in unassigned_open:
            lines.append(f"- {code['name']} ({code['assigned_segment_count']} assigned segments)")
    else:
        lines.append("None.")
    lines.append("")
    return "\n".join(lines)


def generate_project_summary_markdown(active_project: sqlite3.Row) -> str:
    project_id = active_project["id"]
    counts = get_dashboard_counts(project_id)
    methodology_counts = get_methodology_note_counts(project_id)
    documents = get_documents_for_project(project_id)
    research_questions = query_all(
        """
        SELECT question, note, created_at
        FROM research_questions
        WHERE project_id = ?
        ORDER BY datetime(created_at) DESC, id DESC
        """,
        (project_id,),
    )
    top_codes = query_all(
        """
        SELECT codes.name, codes.code_type, COUNT(segment_codes.segment_id) AS count
        FROM codes
        LEFT JOIN segment_codes ON segment_codes.code_id = codes.id
        WHERE codes.project_id = ?
        GROUP BY codes.id
        ORDER BY count DESC, codes.name COLLATE NOCASE
        LIMIT 10
        """,
        (project_id,),
    )
    top_actors = query_all(
        """
        SELECT actors.name, actors.actor_type, COUNT(segment_actors.id) AS count
        FROM actors
        LEFT JOIN segment_actors ON segment_actors.actor_id = actors.id
        WHERE actors.project_id = ?
        GROUP BY actors.id
        ORDER BY count DESC, actors.name COLLATE NOCASE
        LIMIT 10
        """,
        (project_id,),
    )
    top_features = get_discourse_feature_counts(project_id)[:10]
    top_relation_types = query_all(
        """
        SELECT relation_type, COUNT(*) AS count
        FROM relations
        WHERE project_id = ?
        GROUP BY relation_type
        ORDER BY count DESC, relation_type
        LIMIT 10
        """,
        (project_id,),
    )
    important_memos = query_all(
        """
        SELECT title, memo_type, status, created_at
        FROM memos
        WHERE project_id = ? AND status IN ('important', 'use_in_article')
        ORDER BY datetime(created_at) DESC, id DESC
        LIMIT 10
        """,
        (project_id,),
    )
    audit_summary = query_all(
        """
        SELECT action, COUNT(*) AS count, MAX(created_at) AS latest_at
        FROM audit_log
        WHERE project_id = ?
        GROUP BY action
        ORDER BY count DESC, action
        """,
        (project_id,),
    )
    lines = [
        "# Project Summary",
        "",
        f"Project: {active_project['name']}",
        f"Description: {active_project['description'] or ''}",
        f"Methodology mode: {METHODOLOGY_MODES.get(active_project['methodology_mode'], active_project['methodology_mode'])}",
        f"Research goal: {active_project['research_goal'] or ''}",
        f"Principal investigator: {active_project['principal_investigator'] or ''}",
        f"Exported: {export_timestamp()}",
        "",
        "## Counts",
        "",
    ]
    for key in [
        "documents",
        "segments",
        "open_codes",
        "axial_codes",
        "categories",
        "coded_segments",
        "memos",
        "cda_markers",
        "actors",
        "discourse_features",
        "relations",
    ]:
        lines.append(f"- {key.replace('_', ' ').title()}: {counts.get(key, 0)}")
    lines.extend(["", "## Documents", ""])
    if documents:
        for document in documents:
            lines.append(
                f"- {document['title']} ({document['file_type']}, {document['segment_count']} segments)"
            )
    else:
        lines.append("No documents imported.")
    lines.extend(["", "## Research Questions", ""])
    if research_questions:
        for question in research_questions:
            lines.append(f"- {question['question']}")
            if question["note"]:
                lines.append(f"  - {question['note']}")
    else:
        lines.append("No research questions recorded.")
    lines.extend(["", "## Top Codes by Segment Count", ""])
    lines.extend([f"- {row['name']} ({row['code_type']}): {row['count']}" for row in top_codes] or ["None."])
    lines.extend(["", "## Top Actors by Annotation Count", ""])
    lines.extend([f"- {row['name']} ({row['actor_type']}): {row['count']}" for row in top_actors] or ["None."])
    lines.extend(["", "## Top Discourse Feature Types", ""])
    lines.extend([f"- {row['feature_type']}: {row['count']}" for row in top_features] or ["None."])
    lines.extend(["", "## Analytical Relations", ""])
    lines.append(f"- Total relations: {counts.get('relations', 0)}")
    lines.append("- Visual exports are available in the complete research package ZIP.")
    lines.extend(["", "### Most Common Relation Types", ""])
    lines.extend([f"- {row['relation_type']}: {row['count']}" for row in top_relation_types] or ["None."])
    lines.extend(["", "## Methodological Protocol", ""])
    lines.append(f"- Methodology mode: {METHODOLOGY_MODES.get(active_project['methodology_mode'], active_project['methodology_mode'])}")
    lines.append(f"- Research goal: {active_project['research_goal'] or ''}")
    lines.append(f"- Methodology notes: {methodology_counts['total']}")
    lines.append(f"- Active protocol notes: {methodology_counts['active_protocol']}")
    lines.append(f"- Notes needing review: {methodology_counts['needs_review']}")
    lines.extend(["", "## Latest Important Memos", ""])
    lines.extend([f"- {row['title']} ({row['memo_type']}, {row['status']}, {row['created_at']})" for row in important_memos] or ["None."])
    lines.extend(["", "## Audit Log Summary", ""])
    lines.extend([f"- {row['action']}: {row['count']} actions, latest {row['latest_at']}" for row in audit_summary] or ["No audit log entries."])
    lines.append("")
    return "\n".join(lines)


def generate_project_json(active_project: sqlite3.Row) -> str:
    project_id = active_project["id"]
    segment_ids = [
        row["id"]
        for row in query_all(
            """
            SELECT segments.id
            FROM segments
            JOIN documents ON documents.id = segments.document_id
            WHERE documents.project_id = ?
            """,
            (project_id,),
        )
    ]
    document_ids = [
        row["id"]
        for row in query_all("SELECT id FROM documents WHERE project_id = ?", (project_id,))
    ]
    code_ids = [
        row["id"] for row in query_all("SELECT id FROM codes WHERE project_id = ?", (project_id,))
    ]
    actor_ids = [
        row["id"] for row in query_all("SELECT id FROM actors WHERE project_id = ?", (project_id,))
    ]
    marker_ids = [
        row["id"]
        for row in query_all("SELECT id FROM discourse_markers WHERE project_id = ?", (project_id,))
    ]

    def rows_for_ids(sql: str, ids: list[int]) -> list[dict]:
        if not ids:
            return []
        placeholders = ",".join("?" for _ in ids)
        return rows_to_dicts(query_all(sql.format(placeholders=placeholders), tuple(ids)))

    data = {
        "project": row_to_dict(active_project),
        "documents": rows_to_dicts(query_all("SELECT * FROM documents WHERE project_id = ? ORDER BY id", (project_id,))),
        "tags": rows_to_dicts(query_all("SELECT * FROM tags WHERE project_id = ? ORDER BY id", (project_id,))),
        "document_tags": rows_for_ids("SELECT * FROM document_tags WHERE document_id IN ({placeholders}) ORDER BY document_id, tag_id", document_ids),
        "codes": rows_to_dicts(query_all("SELECT * FROM codes WHERE project_id = ? ORDER BY id", (project_id,))),
        "segments": rows_for_ids("SELECT * FROM segments WHERE id IN ({placeholders}) ORDER BY id", segment_ids),
        "segment_codes": rows_for_ids("SELECT * FROM segment_codes WHERE segment_id IN ({placeholders}) ORDER BY segment_id, code_id", segment_ids),
        "memos": rows_to_dicts(query_all("SELECT * FROM memos WHERE project_id = ? ORDER BY id", (project_id,))),
        "methodology_notes": rows_to_dicts(query_all("SELECT * FROM methodology_notes WHERE project_id = ? ORDER BY id", (project_id,))),
        "relations": rows_to_dicts(query_all("SELECT * FROM relations WHERE project_id = ? ORDER BY id", (project_id,))),
        "research_questions": rows_to_dicts(query_all("SELECT * FROM research_questions WHERE project_id = ? ORDER BY id", (project_id,))),
        "discourse_markers": rows_to_dicts(query_all("SELECT * FROM discourse_markers WHERE project_id = ? ORDER BY id", (project_id,))),
        "segment_discourse_markers": rows_for_ids("SELECT * FROM segment_discourse_markers WHERE marker_id IN ({placeholders}) ORDER BY segment_id, marker_id", marker_ids),
        "actors": rows_to_dicts(query_all("SELECT * FROM actors WHERE project_id = ? ORDER BY id", (project_id,))),
        "segment_actors": rows_for_ids("SELECT * FROM segment_actors WHERE actor_id IN ({placeholders}) ORDER BY id", actor_ids),
        "discourse_features": rows_for_ids("SELECT * FROM discourse_features WHERE segment_id IN ({placeholders}) ORDER BY id", segment_ids),
        "audit_log": rows_to_dicts(query_all("SELECT * FROM audit_log WHERE project_id = ? ORDER BY id", (project_id,))),
    }
    return json.dumps(data, indent=2, ensure_ascii=False)


def get_model_export_data(active_project: sqlite3.Row, filters: dict | None = None) -> dict:
    project_id = active_project["id"]
    relations = get_visual_model_relations(project_id, filters) if filters else get_relations_for_project(project_id, {})
    entities = {}
    for relation in relations:
        for side in ("source", "target"):
            key = (relation[f"{side}_type"], relation[f"{side}_id"])
            entities[key] = {
                "type": relation[f"{side}_type"],
                "id": relation[f"{side}_id"],
                "label": relation[f"{side}_label"],
            }
    return {
        "project": row_to_dict(active_project),
        "methodology_mode": active_project["methodology_mode"],
        "generated_at": export_timestamp(),
        "entities": list(entities.values()),
        "relations": [
            {
                "id": relation["id"],
                "title": relation["title"],
                "source_type": relation["source_type"],
                "source_id": relation["source_id"],
                "source_label": relation["source_label"],
                "target_type": relation["target_type"],
                "target_id": relation["target_id"],
                "target_label": relation["target_label"],
                "relation_type": relation["relation_type"],
                "strength": relation["strength"],
                "memo": relation["memo"],
                "evidence_note": relation["evidence_note"],
                "created_at": relation["created_at"],
                "updated_at": relation["updated_at"],
            }
            for relation in relations
        ],
    }


def generate_model_markdown(active_project: sqlite3.Row) -> str:
    data = get_model_export_data(active_project)
    counts = get_model_counts(active_project["id"])
    relations = data["relations"]
    lines = [
        "# Analytical Model",
        "",
        f"Project: {active_project['name']}",
        f"Methodology mode: {METHODOLOGY_MODES.get(active_project['methodology_mode'], active_project['methodology_mode'])}",
        f"Exported: {data['generated_at']}",
        "",
        "## Summary",
        "",
        f"- Total relations: {counts['total']}",
        f"- Strong relations: {counts['strong']}",
        f"- Uncertain relations: {counts['uncertain']}",
        f"- Negative case relations: {counts['negative_cases']}",
        "",
        "## Relations",
        "",
    ]
    if not relations:
        lines.extend(["No analytical relations created yet.", ""])
    for relation in relations:
        lines.extend(
            [
                f"### {relation['title'] or 'Relation #' + str(relation['id'])}",
                "",
                f"Source: {relation['source_label']}",
                f"Relation: {relation['relation_type']}",
                f"Target: {relation['target_label']}",
                f"Strength: {relation['strength']}",
                "",
                "Memo:",
                "",
                relation["memo"] or "",
                "",
                "Evidence:",
                "",
                relation["evidence_note"] or "",
                "",
            ]
        )
    lines.extend(["## Relations by type", ""])
    by_type: dict[str, list] = {}
    for relation in relations:
        by_type.setdefault(relation["relation_type"], []).append(relation)
    if by_type:
        for relation_type, type_relations in sorted(by_type.items()):
            lines.extend([f"### {relation_type}", ""])
            for relation in type_relations:
                lines.append(
                    f"- {relation['source_label']} -> {relation_type} -> {relation['target_label']}"
                )
            lines.append("")
    else:
        lines.extend(["No relation types yet.", ""])
    lines.extend(["## Entity-centered summaries", ""])
    if data["entities"]:
        for entity in data["entities"]:
            outgoing = len(
                [
                    relation for relation in relations
                    if relation["source_type"] == entity["type"] and relation["source_id"] == entity["id"]
                ]
            )
            incoming = len(
                [
                    relation for relation in relations
                    if relation["target_type"] == entity["type"] and relation["target_id"] == entity["id"]
                ]
            )
            lines.append(f"- {entity['label']}: {outgoing} outgoing, {incoming} incoming")
    else:
        lines.append("No related entities yet.")
    lines.append("")
    return "\n".join(lines)


def generate_model_json(active_project: sqlite3.Row) -> str:
    return json.dumps(get_model_export_data(active_project), indent=2, ensure_ascii=False)


def visual_model_nodes(relations: list[dict]) -> list[dict]:
    nodes = {}
    for relation in relations:
        for side in ("source", "target"):
            entity_type = relation[f"{side}_type"]
            entity_id = relation[f"{side}_id"]
            key = (entity_type, entity_id)
            nodes[key] = {
                "type": entity_type,
                "id": entity_id,
                "node_id": graph_node_id(entity_type, entity_id),
                "label": relation[f"{side}_label"],
                "short_label": short_label(relation[f"{side}_label"]),
                "code_type": code_type_from_label(relation[f"{side}_label"]) if entity_type == "code" else "",
            }
    return sorted(nodes.values(), key=lambda node: (node["type"], node["id"]))


def visual_export_header(prefix: str, active_project: sqlite3.Row) -> list[str]:
    return [
        f"{prefix} Generated by discourseLab",
        f"{prefix} Project: {active_project['name']}",
        f"{prefix} Exported: {export_timestamp()}",
    ]


def relation_edge_label(relation: dict) -> str:
    return f"{short_relation_label(relation['relation_type'])} / {relation['strength']}"


def visual_node_style(node: dict) -> dict[str, str]:
    entity_type = node["type"]
    code_type = node.get("code_type", "")
    styles = {
        "document": {"class": "document", "label": "Document", "fill": "#f1f5f9", "stroke": "#64748b", "dot_shape": "tab", "tikz": "documentnode"},
        "segment": {"class": "segment", "label": "Segment", "fill": "#fff8d6", "stroke": "#9a7a00", "dot_shape": "note", "tikz": "segmentnode"},
        "memo": {"class": "memo", "label": "Memo", "fill": "#e8f7e4", "stroke": "#3c7a3c", "dot_shape": "folder", "tikz": "memonode"},
        "research_question": {"class": "rq", "label": "RQ", "fill": "#ffe6e6", "stroke": "#a33", "dot_shape": "diamond", "tikz": "rqnode"},
        "discourse_marker": {"class": "marker", "label": "CDA marker", "fill": "#d9f6f1", "stroke": "#0f766e", "dot_shape": "hexagon", "tikz": "markernode"},
        "actor": {"class": "actor", "label": "Actor", "fill": "#fff2cc", "stroke": "#a66a00", "dot_shape": "ellipse", "tikz": "actornode"},
        "discourse_feature": {"class": "feature", "label": "Feature", "fill": "#e0f7ff", "stroke": "#0369a1", "dot_shape": "component", "tikz": "featurenode"},
    }
    if entity_type == "code":
        if code_type == "category":
            return {"class": "category", "label": "Code: category", "fill": "#ddd0ff", "stroke": "#5b21b6", "dot_shape": "box", "tikz": "categorynode"}
        if code_type == "axial":
            return {"class": "axial", "label": "Code: axial", "fill": "#eee5ff", "stroke": "#6b46a3", "dot_shape": "box", "tikz": "axialnode"}
        return {"class": "code", "label": f"Code: {code_type or 'open'}", "fill": "#e8f1ff", "stroke": "#3566a0", "dot_shape": "box", "tikz": "codenode"}
    return styles.get(entity_type, {"class": "generic", "label": entity_type.replace("_", " ").title(), "fill": "#f8fafc", "stroke": "#64748b", "dot_shape": "box", "tikz": "genericnode"})


def visual_model_mode_label(filters: dict | None) -> str:
    mode = (filters or {}).get("model_mode", "simplified")
    return VISUAL_MODEL_MODES.get(mode, VISUAL_MODEL_MODES["simplified"])


def visual_model_mode_filters(model_mode: str) -> dict:
    full_mode = model_mode == "full"
    return {
        "model_mode": model_mode,
        "relation_type": "",
        "strength": "",
        "involves_type": "",
        "include_uncertain": full_mode,
        "include_weak": full_mode,
        "max_relations": 100 if full_mode else 25,
    }


def generate_model_mermaid(active_project: sqlite3.Row, filters: dict | None = None) -> str:
    data = get_model_export_data(active_project, filters)
    relations = data["relations"]
    lines = visual_export_header("%%", active_project)
    lines.append(f"%% Model mode: {visual_model_mode_label(filters)}")
    if not relations:
        lines.extend(["%% No relations exist yet.", "flowchart TD", '  empty["No analytical relations yet"]'])
        return "\n".join(lines) + "\n"
    lines.extend(
        [
            "flowchart TD",
            "  classDef actor fill:#fff2cc,stroke:#a66a00;",
            "  classDef code fill:#e8f1ff,stroke:#3566a0;",
            "  classDef axial fill:#eee5ff,stroke:#6b46a3;",
            "  classDef category fill:#ddd0ff,stroke:#5b21b6;",
            "  classDef memo fill:#e8f7e4,stroke:#3c7a3c;",
            "  classDef segment fill:#fff8d6,stroke:#9a7a00;",
            "  classDef document fill:#f1f5f9,stroke:#64748b;",
            "  classDef rq fill:#ffe6e6,stroke:#a33;",
            "  classDef marker fill:#d9f6f1,stroke:#0f766e;",
            "  classDef feature fill:#e0f7ff,stroke:#0369a1;",
        ]
    )
    for node in visual_model_nodes(relations):
        style = visual_node_style(node)
        label = escape_mermaid_text(node["short_label"])
        lines.append(f'  {node["node_id"]}["{label}"]')
        lines.append(f'  class {node["node_id"]} {style["class"]};')
    for relation in relations:
        source_id = graph_node_id(relation["source_type"], relation["source_id"])
        target_id = graph_node_id(relation["target_type"], relation["target_id"])
        edge_label = escape_mermaid_text(relation_edge_label(relation))
        arrow = strength_style(relation["strength"])["mermaid"]
        lines.append(f'  {source_id} {arrow}|"{edge_label}"| {target_id}')
    return "\n".join(lines) + "\n"


def generate_model_dot(active_project: sqlite3.Row, filters: dict | None = None) -> str:
    data = get_model_export_data(active_project, filters)
    relations = data["relations"]
    if not relations:
        return "\n".join(
            [
                "digraph discourseLabModel {",
                '  empty [label="No analytical relations yet"];',
                "}",
                "",
            ]
        )
    lines = [
        "digraph discourseLabModel {",
        "  graph [rankdir=LR, splines=true, overlap=false];",
        "  node [shape=box, style=\"rounded,filled\", fontname=\"Arial\"];",
        "  edge [fontname=\"Arial\"];",
        f'  label="discourseLab analytical model — {escape_dot_text(visual_model_mode_label(filters))}";',
        "  labelloc=t;",
        "",
    ]
    for node in visual_model_nodes(relations):
        style = visual_node_style(node)
        node_id = escape_dot_text(node["node_id"])
        label = escape_dot_text(f"{style['label']}\\n{node['short_label']}").replace("\\\\n", "\\n")
        lines.append(
            f'  "{node_id}" [label="{label}", shape={style["dot_shape"]}, '
            f'fillcolor="{style["fill"]}", color="{style["stroke"]}"];'
        )
    for relation in relations:
        source_id = escape_dot_text(graph_node_id(relation["source_type"], relation["source_id"]))
        target_id = escape_dot_text(graph_node_id(relation["target_type"], relation["target_id"]))
        edge_label = escape_dot_text(relation_edge_label(relation))
        family = relation_family(relation["relation_type"])
        family_style = relation_family_styles()[family]
        strength = strength_style(relation["strength"])
        tooltip = escape_dot_text(edge_tooltip(relation))
        lines.append(
            f'  "{source_id}" -> "{target_id}" [label="{edge_label}", color="{family_style["color"]}", '
            f'fontcolor="{family_style["color"]}", penwidth={strength["dot_penwidth"]}, '
            f'style={strength["dot_style"]}, tooltip="{tooltip}"];'
        )
    lines.extend(["}", ""])
    return "\n".join(lines)


def generate_model_tikz(active_project: sqlite3.Row, filters: dict | None = None) -> str:
    data = get_model_export_data(active_project, filters)
    relations = data["relations"]
    lines = visual_export_header("%", active_project)
    lines.append(f"% Model mode: {visual_model_mode_label(filters)}")
    lines.append("% Generated layout is deterministic and may need manual adjustment in publication documents.")
    if not relations:
        lines.extend(
            [
                r"\begin{tikzpicture}",
                r"\node[draw, rounded corners] at (0,0) {No analytical relations yet};",
                r"\end{tikzpicture}",
                "",
            ]
        )
        return "\n".join(lines)
    nodes = visual_model_nodes(relations)
    lines.extend(
        [
            r"\begin{tikzpicture}[",
            r"  node distance=2cm,",
            r"  every node/.style={draw, rounded corners, align=center, font=\small},",
            r"  relation/.style={->, >=stealth},",
            r"  documentnode/.style={fill=gray!12, draw=gray!70},",
            r"  segmentnode/.style={fill=yellow!18, draw=yellow!55!black},",
            r"  codenode/.style={fill=blue!10, draw=blue!55!black},",
            r"  axialnode/.style={fill=purple!10, draw=purple!60!black},",
            r"  categorynode/.style={fill=purple!20, draw=purple!75!black},",
            r"  memonode/.style={fill=green!12, draw=green!50!black},",
            r"  rqnode/.style={fill=red!10, draw=red!60!black},",
            r"  actornode/.style={fill=orange!15, draw=orange!70!black},",
            r"  markernode/.style={fill=teal!12, draw=teal!70!black},",
            r"  featurenode/.style={fill=cyan!12, draw=cyan!70!black},",
            r"  genericnode/.style={fill=gray!8, draw=gray!70}",
            r"]",
        ]
    )
    positions, _, _ = layered_svg_positions(nodes)
    for node in nodes:
        x, y = positions[node["node_id"]]
        style = visual_node_style(node)
        label = f"{escape_latex_text(style['label'])}\\\\ {escape_latex_text(node['short_label'])}"
        lines.append(fr"\node[{style['tikz']}] ({node['node_id']}) at ({x / 90:.2f},{-y / 90:.2f}) {{{label}}};")
    for relation in relations:
        source_id = graph_node_id(relation["source_type"], relation["source_id"])
        target_id = graph_node_id(relation["target_type"], relation["target_id"])
        edge_label = escape_latex_text(relation_edge_label(relation))
        family = relation_family_styles()[relation_family(relation["relation_type"])]
        strength = strength_style(relation["strength"])
        lines.append(
            fr"\draw[relation,{strength['tikz']},{family['tikz']}] ({source_id}) "
            fr"to[bend left=8] node[above, draw=none, fill=white] {{{edge_label}}} ({target_id});"
        )
    lines.extend([r"\end{tikzpicture}", ""])
    return "\n".join(lines)


def node_layer(node: dict) -> int:
    if node["type"] in {"document", "segment", "discourse_feature"}:
        return 0
    if node["type"] in {"actor", "discourse_marker"}:
        return 1
    if node["type"] == "code":
        return {"open": 1, "axial": 2, "category": 3}.get(node.get("code_type", ""), 1)
    if node["type"] == "memo":
        return 2
    if node["type"] == "research_question":
        return 3
    return 4


def layered_svg_positions(nodes: list[dict]) -> tuple[dict[str, tuple[float, float]], int, int]:
    node_width = 230
    node_height = 76
    x_gap = 90
    y_gap = 28
    margin_x = 50
    margin_y = 90
    layers: dict[int, list[dict]] = {}
    for node in nodes:
        layers.setdefault(node_layer(node), []).append(node)
    for layer_nodes in layers.values():
        layer_nodes.sort(key=lambda node: (node["type"], node.get("code_type", ""), node["label"].lower()))
    layer_count = max(layers.keys(), default=0) + 1
    max_layer_size = max((len(layer_nodes) for layer_nodes in layers.values()), default=1)
    graph_width = margin_x * 2 + layer_count * node_width + max(0, layer_count - 1) * x_gap
    graph_height = margin_y * 2 + max_layer_size * node_height + max(0, max_layer_size - 1) * y_gap
    graph_width = max(graph_width, 1100)
    graph_height = max(graph_height, 700)
    positions = {}
    for layer, layer_nodes in layers.items():
        layer_height = len(layer_nodes) * node_height + max(0, len(layer_nodes) - 1) * y_gap
        y_start = margin_y + max(0, (graph_height - margin_y * 2 - layer_height) / 2)
        x = margin_x + layer * (node_width + x_gap)
        for index, node in enumerate(layer_nodes):
            y = y_start + index * (node_height + y_gap)
            positions[node["node_id"]] = (x, y)
    return positions, graph_width, graph_height


def svg_wrapped_text(text: str, x: float, y: float, max_chars: int = 24, max_lines: int = 3, css_class: str = "node-label") -> list[str]:
    words = " ".join((text or "").split()).split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) <= max_chars:
            current = candidate
            continue
        if current:
            lines.append(current)
        current = word
        if len(lines) == max_lines - 1:
            break
    remaining_words = words[len(" ".join(lines + ([current] if current else [])).split()):]
    if current:
        if remaining_words or len(current) > max_chars:
            current = truncate_text(current, max_chars)
        lines.append(current)
    if not lines:
        lines = [""]
    lines = lines[:max_lines]
    output = [f'  <text x="{x:.1f}" y="{y:.1f}" text-anchor="middle" class="{css_class}">']
    for index, line in enumerate(lines):
        dy = 0 if index == 0 else 15
        output.append(f'    <tspan x="{x:.1f}" dy="{dy}">{escape_svg_text(line)}</tspan>')
    output.append("  </text>")
    return output


def edge_tooltip(relation: dict) -> str:
    parts = [
        f"Relation: {relation['relation_type']}",
        f"Strength: {relation['strength']}",
    ]
    if relation.get("title"):
        parts.append(f"Title: {relation['title']}")
    if relation.get("memo"):
        parts.append(f"Memo: {short_label(relation['memo'], 120)}")
    if relation.get("evidence_note"):
        parts.append(f"Evidence: {short_label(relation['evidence_note'], 120)}")
    return "\n".join(parts)


def node_tooltip(node: dict) -> str:
    return f"{node['label']}\nEntity type: {node['type']}"


def svg_legend(x: float, y: float) -> list[str]:
    lines = [
        f'  <g id="legend">',
        f'    <rect x="{x}" y="{y}" width="280" height="675" rx="8" fill="#f8fafc" stroke="#cbd5e1" />',
        f'    <text x="{x + 16}" y="{y + 26}" class="legend-title">Legend</text>',
        f'    <text x="{x + 16}" y="{y + 52}" class="legend-title">Node types</text>',
    ]
    node_examples = [
        {"label": "Document", "fill": "#f1f5f9", "stroke": "#64748b"},
        {"label": "Segment", "fill": "#fff8d6", "stroke": "#9a7a00"},
        {"label": "Open code", "fill": "#e8f1ff", "stroke": "#3566a0"},
        {"label": "Axial code", "fill": "#eee5ff", "stroke": "#6b46a3"},
        {"label": "Category", "fill": "#ddd0ff", "stroke": "#5b21b6"},
        {"label": "Memo", "fill": "#e8f7e4", "stroke": "#3c7a3c"},
        {"label": "Research question", "fill": "#ffe6e6", "stroke": "#a33"},
        {"label": "Actor", "fill": "#fff2cc", "stroke": "#a66a00"},
        {"label": "CDA marker", "fill": "#d9f6f1", "stroke": "#0f766e"},
        {"label": "Discourse feature", "fill": "#e0f7ff", "stroke": "#0369a1"},
    ]
    current_y = y + 74
    for item in node_examples:
        lines.append(f'    <rect x="{x + 18}" y="{current_y - 12}" width="22" height="14" rx="3" fill="{item["fill"]}" stroke="{item["stroke"]}" />')
        lines.append(f'    <text x="{x + 50}" y="{current_y}" class="legend-text">{item["label"]}</text>')
        current_y += 25
    current_y += 12
    lines.append(f'    <text x="{x + 16}" y="{current_y}" class="legend-title">Edge strength</text>')
    current_y += 24
    strength_examples = [
        ("strong", "thick solid"),
        ("moderate", "medium solid"),
        ("weak", "thin dashed"),
        ("uncertain", "thin dotted"),
    ]
    for strength, label in strength_examples:
        style = strength_style(strength)
        dash = f' stroke-dasharray="{style["svg_dash"]}"' if style["svg_dash"] else ""
        lines.append(f'    <line x1="{x + 18}" y1="{current_y - 4}" x2="{x + 78}" y2="{current_y - 4}" stroke="#334155" stroke-width="{style["svg_width"]}" opacity="{style["svg_opacity"]}"{dash} />')
        lines.append(f'    <text x="{x + 90}" y="{current_y}" class="legend-text">{strength} = {label}</text>')
        current_y += 25
    current_y += 12
    lines.append(f'    <text x="{x + 16}" y="{current_y}" class="legend-title">Relation families</text>')
    current_y += 24
    for family in ["support", "contrast", "causal_process", "gt", "cda_power", "cda_representation", "generic"]:
        style = relation_family_styles()[family]
        lines.append(f'    <line x1="{x + 18}" y1="{current_y - 4}" x2="{x + 78}" y2="{current_y - 4}" stroke="{style["color"]}" stroke-width="3" />')
        lines.append(f'    <text x="{x + 90}" y="{current_y}" class="legend-text">{escape_svg_text(style["label"])}</text>')
        current_y += 25
    lines.append("  </g>")
    return lines


def generate_model_svg(active_project: sqlite3.Row, filters: dict | None = None) -> str:
    data = get_model_export_data(active_project, filters)
    relations = data["relations"]
    if not relations:
        return "\n".join(
            [
                '<svg xmlns="http://www.w3.org/2000/svg" width="640" height="160" viewBox="0 0 640 160">',
                "  <title>discourseLab analytical model</title>",
                '  <text x="320" y="82" text-anchor="middle" font-family="Arial, sans-serif" font-size="18">No analytical relations yet</text>',
                "</svg>",
                "",
            ]
        )
    nodes = visual_model_nodes(relations)
    positions, graph_width, graph_height = layered_svg_positions(nodes)
    node_width = 230
    node_height = 76
    legend_width = 300
    width = graph_width + legend_width + 40
    height = max(graph_height, 820)
    family_styles = relation_family_styles()
    lines = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        "  <title>discourseLab analytical model</title>",
        "  <style>",
        "    text { font-family: Arial, sans-serif; fill: #172033; }",
        "    .node-type { font-size: 10px; font-weight: 700; text-transform: uppercase; letter-spacing: .04em; fill: #475569; }",
        "    .node-label { font-size: 12px; }",
        "    .edge-label { font-size: 10.5px; fill: #172033; }",
        "    .legend-title { font-size: 13px; font-weight: 700; }",
        "    .legend-text { font-size: 11px; fill: #334155; }",
        "  </style>",
        "  <defs>",
        '    <marker id="arrow" markerWidth="10" markerHeight="10" refX="9" refY="3" orient="auto" markerUnits="strokeWidth">',
        '      <path d="M0,0 L0,6 L9,3 z" fill="#334155" />',
        "    </marker>",
        "  </defs>",
        '  <rect width="100%" height="100%" fill="#ffffff" />',
        f'  <text x="50" y="34" font-size="18" font-weight="700">discourseLab analytical model — {escape_svg_text(visual_model_mode_label(filters))}</text>',
        f'  <text x="50" y="56" font-size="12" fill="#64748b">Project: {escape_svg_text(active_project["name"])}</text>',
    ]
    for index, relation in enumerate(relations):
        source_id = graph_node_id(relation["source_type"], relation["source_id"])
        target_id = graph_node_id(relation["target_type"], relation["target_id"])
        source_x, source_y = positions[source_id]
        target_x, target_y = positions[target_id]
        source_center_x = source_x + node_width / 2
        target_center_x = target_x + node_width / 2
        source_center_y = source_y + node_height / 2
        target_center_y = target_y + node_height / 2
        if target_center_x >= source_center_x:
            x1 = source_x + node_width
            x2 = target_x
        else:
            x1 = source_x
            x2 = target_x + node_width
        y1 = source_y + node_height / 2
        y2 = target_y + node_height / 2
        curve = max(70, abs(x2 - x1) * 0.38)
        c1x = x1 + curve if x2 >= x1 else x1 - curve
        c2x = x2 - curve if x2 >= x1 else x2 + curve
        path = f"M {x1:.1f},{y1:.1f} C {c1x:.1f},{y1:.1f} {c2x:.1f},{y2:.1f} {x2:.1f},{y2:.1f}"
        label_x = (x1 + x2) / 2
        label_y = (y1 + y2) / 2 - 10 + ((index % 3) - 1) * 12
        edge_label = escape_svg_text(relation_edge_label(relation))
        family = relation_family(relation["relation_type"])
        family_style = family_styles[family]
        strength = strength_style(relation["strength"])
        dash = f' stroke-dasharray="{strength["svg_dash"]}"' if strength["svg_dash"] else ""
        lines.append(
            f'  <path d="{path}" fill="none" stroke="{family_style["color"]}" '
            f'stroke-width="{strength["svg_width"]}" opacity="{strength["svg_opacity"]}" '
            f'marker-end="url(#arrow)"{dash}>'
        )
        lines.append(f"    <title>{escape_svg_text(edge_tooltip(relation))}</title>")
        lines.append("  </path>")
        label_width = min(150, max(70, len(edge_label) * 6))
        lines.append(f'  <rect x="{label_x - label_width / 2:.1f}" y="{label_y - 12:.1f}" width="{label_width}" height="17" rx="3" fill="#ffffff" opacity="0.92" />')
        lines.append(f'  <text x="{label_x:.1f}" y="{label_y:.1f}" text-anchor="middle" class="edge-label">{edge_label}</text>')
    for node in nodes:
        x, y = positions[node["node_id"]]
        style = visual_node_style(node)
        rx = 20 if node["type"] == "actor" else 7
        lines.append(f'  <g id="{node["node_id"]}">')
        lines.append(f"    <title>{escape_svg_text(node_tooltip(node))}</title>")
        lines.append(f'    <rect x="{x}" y="{y}" width="{node_width}" height="{node_height}" rx="{rx}" fill="{style["fill"]}" stroke="{style["stroke"]}" stroke-width="1.4" />')
        lines.append(f'    <text x="{x + node_width / 2:.1f}" y="{y + 18:.1f}" text-anchor="middle" class="node-type">{escape_svg_text(style["label"])}</text>')
        lines.extend(svg_wrapped_text(node["label"], x + node_width / 2, y + 39, max_chars=26, max_lines=2))
        lines.append("  </g>")
    lines.extend(svg_legend(graph_width + 30, 90))
    lines.extend(["</svg>", ""])
    return "\n".join(lines)


def generate_methodology_protocol_markdown(active_project: sqlite3.Row) -> str:
    project_id = active_project["id"]
    notes = get_methodology_notes_for_project(project_id, {})
    libraries = get_relevant_methodology_libraries(active_project)
    counts = get_methodology_note_counts(project_id)
    lines = [
        "# Methodological Protocol",
        "",
        f"Project: {active_project['name']}",
        f"Methodology mode: {METHODOLOGY_MODES.get(active_project['methodology_mode'], active_project['methodology_mode'])}",
        f"Principal investigator: {active_project['principal_investigator'] or ''}",
        f"Research goal: {active_project['research_goal'] or ''}",
        f"Exported: {export_timestamp()}",
        "",
        "## Project methodology",
        "",
        get_methodology_overview(active_project),
        "",
        "## Protocol counts",
        "",
        f"- Total methodology notes: {counts['total']}",
        f"- Active protocol notes: {counts['active_protocol']}",
        f"- Coding rules: {counts['coding_rules']}",
        f"- Sampling rules: {counts['sampling_rules']}",
        f"- Notes needing review: {counts['needs_review']}",
        "",
        "## Active protocol notes",
        "",
    ]
    active_notes = [note for note in notes if note["status"] == "active"]
    if active_notes:
        grouped: dict[tuple[str, str], list[dict]] = {}
        for note in active_notes:
            grouped.setdefault((note["note_type_label"], note["methodology_area_label"]), []).append(note)
        for (note_type, area), group_notes in sorted(grouped.items()):
            lines.extend([f"### {note_type} — {area}", ""])
            for note in group_notes:
                lines.extend(methodology_note_markdown_block(note))
    else:
        lines.extend(["No active protocol notes.", ""])
    review_notes = [note for note in notes if note["status"] == "needs_review"]
    lines.extend(["## Notes needing review", ""])
    if review_notes:
        for note in review_notes:
            lines.extend(methodology_note_markdown_block(note))
    else:
        lines.extend(["None.", ""])
    lines.extend(["## Methodological sources", ""])
    seen_sources = set()
    for library in libraries:
        lines.append(f"### {library['title']}")
        lines.append("")
        for source in library.get("sources", []):
            if source["id"] in seen_sources:
                continue
            seen_sources.add(source["id"])
            lines.append(f"- {source['apa']}")
            if source.get("note"):
                lines.append(f"  - {source['note']}")
        lines.append("")
    lines.extend(["## Helper prompts used in this project", ""])
    for library in libraries:
        lines.append(f"### {library['title']}")
        lines.append("")
        for phase in library.get("phases", []):
            if phase.get("prompts"):
                lines.append(f"- {phase['title']}")
                for prompt in phase["prompts"]:
                    lines.append(f"  - {prompt}")
        for concept in library.get("concepts", []):
            if concept.get("prompts"):
                lines.append(f"- {concept['term']}")
                for prompt in concept["prompts"]:
                    lines.append(f"  - {prompt}")
        lines.append("")
    return "\n".join(lines)


def methodology_note_markdown_block(note: dict) -> list[str]:
    return [
        f"#### {note['title']}",
        "",
        f"Type: {note['note_type_label']}",
        f"Area: {note['methodology_area_label']}",
        f"Status: {note['status_label']}",
        f"Linked entity: {note['linked_entity_label']}",
        f"Created: {note['created_at']}",
        f"Updated: {note['updated_at']}",
        "",
        note["body"],
        "",
    ]


def generate_project_package_zip(active_project: sqlite3.Row) -> bytes:
    timestamp = export_timestamp()
    readme = "\n".join(
        [
            "discourseLab research export package",
            "",
            f"Generated by: discourseLab",
            f"Project: {active_project['name']}",
            f"Methodology mode: {METHODOLOGY_MODES.get(active_project['methodology_mode'], active_project['methodology_mode'])}",
            f"Exported: {timestamp}",
            "",
            "Contents:",
            "- codebook.md",
            "- coded_segments.csv",
            "- coded_segments.md",
            "- memos.md",
            "- methodology_protocol.md",
            "- project_summary.md",
            "- project.json",
            "- analytical_model.md",
            "- analytical_model.json",
            "- analytical_model_[mode].mmd for simplified, argument, evidence, gt, cda, full",
            "- analytical_model_[mode].dot for simplified, argument, evidence, gt, cda, full",
            "- analytical_model_[mode].tikz for simplified, argument, evidence, gt, cda, full",
            "- analytical_model_[mode].svg for simplified, argument, evidence, gt, cda, full",
            "- cooccurrence_network.json",
            "- cooccurrence_edges.csv",
            *(
                ["- gt_hierarchy.md"]
                if project_supports_gt(active_project)
                else []
            ),
            *(
                ["- cda_features.csv", "- voice_silence.csv"]
                if project_supports_cda(active_project)
                else []
            ),
            "",
            "Visual model exports are generated from saved analytical relations.",
            "The co-occurrence network is generated from segment-level assignments. It is not the same as the manually curated analytical model.",
            "Model modes:",
            "- simplified: strong and moderate relations by default, capped for readability.",
            "- argument: research questions, categories, axial codes, memos, strong relations, and argument-building relation types.",
            "- evidence: documents, segments, memos, evidence/example/negative-case/support/contradiction relations.",
            "- gt: code-only grounded theory relation structure.",
            "- cda: actor, marker, feature, and CDA relation structure.",
            "- full: all relations, including weak and uncertain, capped at 100.",
            "",
            "Uploaded source documents are not included in this package in Phase 14.",
            "",
        ]
    )
    files = {
        "codebook.md": build_codebook_markdown(active_project),
        "coded_segments.csv": generate_coded_segments_csv(active_project),
        "coded_segments.md": generate_coded_segments_markdown(active_project),
        "memos.md": generate_memos_markdown(active_project),
        "methodology_protocol.md": generate_methodology_protocol_markdown(active_project),
        "project_summary.md": generate_project_summary_markdown(active_project),
        "project.json": generate_project_json(active_project),
        "analytical_model.md": generate_model_markdown(active_project),
        "analytical_model.json": generate_model_json(active_project),
        "cooccurrence_network.json": generate_cooccurrence_network_json(active_project, default_network_filters()),
        "cooccurrence_edges.csv": generate_cooccurrence_edges_csv(active_project, default_network_filters()),
        "README_EXPORT.txt": readme,
    }
    for mode in VISUAL_MODEL_MODE_ORDER:
        filters = visual_model_mode_filters(mode)
        files[f"analytical_model_{mode}.mmd"] = generate_model_mermaid(active_project, filters)
        files[f"analytical_model_{mode}.dot"] = generate_model_dot(active_project, filters)
        files[f"analytical_model_{mode}.tikz"] = generate_model_tikz(active_project, filters)
        files[f"analytical_model_{mode}.svg"] = generate_model_svg(active_project, filters)
    if project_supports_gt(active_project):
        files["gt_hierarchy.md"] = generate_gt_hierarchy_markdown(active_project)
    if project_supports_cda(active_project):
        files["cda_features.csv"] = generate_cda_features_csv(active_project)
        files["voice_silence.csv"] = generate_voice_silence_csv(active_project)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for filename, content in files.items():
            archive.writestr(filename, content.encode("utf-8"))
    return output.getvalue()


def generate_project_backup_zip(active_project: sqlite3.Row) -> bytes:
    timestamp = export_timestamp()
    files = {
        "project.json": generate_project_json(active_project),
        "codebook.md": build_codebook_markdown(active_project),
        "memos.md": generate_memos_markdown(active_project),
        "analytical_model.md": generate_model_markdown(active_project),
        "methodology_protocol.md": generate_methodology_protocol_markdown(active_project),
    }
    readme_lines = [
        "discourseLab project backup",
        "",
        "Generated by discourseLab",
        f"Project name: {active_project['name']}",
        f"Timestamp: {timestamp}",
        "",
        "Contents:",
        *[f"- {name}" for name in sorted(files.keys())],
        "- README_BACKUP.txt",
        "",
        "Restore from backup is not implemented yet.",
        "Original uploaded source files are not included because this project stores extracted text but not per-document upload paths.",
        "",
    ]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("README_BACKUP.txt", "\n".join(readme_lines).encode("utf-8"))
        for filename, content in files.items():
            archive.writestr(filename, content.encode("utf-8"))
    return output.getvalue()


def normalize_code_color(color: str) -> str:
    color = color.strip()
    if re.fullmatch(r"#[0-9a-fA-F]{6}", color):
        return color.lower()
    return DEFAULT_CODE_COLOR


def normalize_cda_color(color: str) -> str:
    color = color.strip()
    if re.fullmatch(r"#[0-9a-fA-F]{6}", color):
        return color.lower()
    return DEFAULT_CDA_MARKER_COLOR


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower().lstrip(".")


def build_unique_upload_name(filename: str) -> str:
    safe_name = secure_filename(filename)
    if not safe_name:
        safe_name = f"document-{uuid4().hex}"

    path = UPLOAD_DIR / safe_name
    if not path.exists():
        return safe_name

    stem = Path(safe_name).stem
    suffix = Path(safe_name).suffix
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    return f"{stem}-{timestamp}-{uuid4().hex[:8]}{suffix}"


def extract_text(path: Path, extension: str) -> str:
    if extension == "txt":
        return extract_txt(path)
    if extension == "docx":
        return extract_docx(path)
    raise ValueError("Unsupported file type.")


def extract_txt(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="latin-1")


def extract_docx(path: Path) -> str:
    document = DocxDocument(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    return "\n\n".join(paragraphs)


def normalize_line_endings(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


app = create_app()


if __name__ == "__main__":
    app.run(debug=True)
